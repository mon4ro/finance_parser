from pathlib import Path

from openpyxl import Workbook, load_workbook

from finance_parser.utilities.rebuild_workbook_values_only import (
    read_workbook_values_only,
    safe_table_name,
    validate_xlsx_basic,
    write_fresh_value_workbook,
)


PROJECT_ROOT = Path(__file__).resolve().parents[1]


def test_rebuild_workbook_values_only_roundtrip(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "values_only.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "UnifiedTransactions"
    ws.append(["UnifiedID", "Comments", "Category"])
    ws.append(["U1", "", "Old"])
    raw = wb.create_sheet("Raw_OP")
    raw.append(["A", "B"])
    raw.append([1, 2])
    wb.save(source)
    wb.close()

    sheets = read_workbook_values_only(source)
    stats = write_fresh_value_workbook(output, sheets)

    problems = validate_xlsx_basic(output, expected_sheets={"UnifiedTransactions", "Raw_OP"})
    assert problems == []
    assert stats["sheets_written"] == 2

    out_wb = load_workbook(output, read_only=True)
    assert out_wb.sheetnames == ["UnifiedTransactions", "Raw_OP"]
    out_ws = out_wb["UnifiedTransactions"]
    assert out_ws["A2"].value == "U1"
    assert out_ws["C2"].value == "Old"
    out_wb.close()


def test_excel_tables_are_generated_from_scratch(tmp_path):
    output = tmp_path / "tables.xlsx"
    sheets = {
        "UnifiedTransactions": [
            ["UnifiedID", "Date", "Amount", "Category"],
            ["U1", "2026-06-08", 12.34, "Test"],
        ],
        "EmptyHeaderOnly": [["A", "B"]],
    }

    stats = write_fresh_value_workbook(
        output,
        sheets,
        basic_formatting=True,
        excel_tables=True,
    )

    assert stats["tables_added"] == 1

    wb = load_workbook(output)
    ws = wb["UnifiedTransactions"]
    assert len(ws.tables) == 1
    assert ws.auto_filter.ref == "A1:D2"
    assert ws["A1"].font.bold is True
    wb.close()


def test_safe_table_name_is_unique_and_validish():
    existing = set()
    first = safe_table_name("Raw OP", existing)
    second = safe_table_name("Raw OP", existing)
    third = safe_table_name("123 weird sheet!", existing)

    assert first == "tbl_Raw_OP"
    assert second == "tbl_Raw_OP_1"
    assert third.startswith("tbl_T_123_weird_sheet")
