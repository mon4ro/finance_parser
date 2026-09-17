from datetime import date
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from finance_parser import settings as settings_module
from finance_parser.settings import AppSettings
from finance_parser.investments.build_monthly_position_value import (
    build_cash_balance_rows,
    build_monthly_values,
    last_completed_month_end,
    load_cash_balance_source,
)


POSITIONS_HEADERS = [
    "Broker", "Portfolio", "PortfolioOwner", "NormalizedInstrument", "Date", "TransactionType",
    "QuantityDelta", "CumulativeQuantity", "CashDelta", "CumulativeNetInvested",
]
PRICES_HEADERS = ["NormalizedInstrument", "ISIN", "PriceSymbol", "Date", "Close", "Currency", "PriceSource", "FetchedAt"]
FX_HEADERS = ["Currency", "Date", "Rate", "Source", "FetchedAt"]
TRANSACTIONS_HEADERS = [
    "Broker", "Portfolio", "PortfolioOwner", "PortfolioType", "NormalizedInstrument",
    "TransactionType", "TradeDate", "CashAmount",
]
RAW_TRANSACTIONS_HEADERS = ["Broker", "Portfolio", "TradeDate", "CashAmount", "CashBalance", "TransactionTypeRaw"]


def _write_sheet(path, sheet_name, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    wb.save(path)


def _positions(path, rows):
    _write_sheet(path, "PortfolioPositions", POSITIONS_HEADERS, rows)


def _prices(path, rows):
    _write_sheet(path, "InstrumentPrices", PRICES_HEADERS, rows)


def _fx(path, rows):
    _write_sheet(path, "FXRates", FX_HEADERS, rows)


def _transactions(path, rows):
    _write_sheet(path, "InvestmentTransactions", TRANSACTIONS_HEADERS, rows)


def _raw_transactions(path, rows):
    """
    Appends a RawInvestmentTransactions sheet to an existing workbook (as
    written by _transactions()) - openpyxl's default Workbook() would
    otherwise overwrite it, so this loads and re-saves in place.
    """
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb.create_sheet("RawInvestmentTransactions")
    ws.append(RAW_TRANSACTIONS_HEADERS)
    for row in rows:
        ws.append([row.get(h, "") for h in RAW_TRANSACTIONS_HEADERS])
    wb.save(path)


_CASH_SETTINGS_YAML = """
project:
  name: "Test"
  locale: "fi_FI"
  default_currency: "EUR"

paths:
  budgeting_input: "input/budgeting"
  budgeting_output: "output/budgeting/ParsedTransactions.xlsx"
  budgeting_rules: "rules/budgeting/TransactionRules.xlsx"
  investment_input: "input/investments"
  investment_output: "output/investments/ParsedInvestments.xlsx"
  investment_rules: "rules/investments/InstrumentMaster.xlsx"

budgeting:
  default_include: "YES"
  source_bank_aliases: {}
  source_account_inference: {}

investments:
  portfolio_owners:
    NORDNET:
      "1": "PERSON_A"
    EVLI: "PERSON_A"
  portfolio_types:
    NORDNET:
      "1": "OSAKESAASTOTILI"
"""


def _with_cash_settings(tmp_path: Path) -> AppSettings:
    settings_path = tmp_path / "settings.yaml"
    settings_path.write_text(_CASH_SETTINGS_YAML, encoding="utf-8")
    return AppSettings.load(settings_path=settings_path, example_path=Path("does-not-exist.yaml"))


def test_last_completed_month_end():
    assert last_completed_month_end(date(2026, 9, 16)).strftime("%Y-%m-%d") == "2026-08-31"
    assert last_completed_month_end(date(2026, 1, 1)).strftime("%Y-%m-%d") == "2025-12-31"


def test_eur_instrument_market_value_and_gain(tmp_path):
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"
    fx_path = tmp_path / "FXRates.xlsx"

    _positions(positions_path, [
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "SAMPO A", "Date": "2021-03-15", "CumulativeQuantity": 100, "CumulativeNetInvested": -1000},
    ])
    _prices(prices_path, [
        {"NormalizedInstrument": "SAMPO A", "Date": "2021-03-31", "Close": 12.0, "Currency": "EUR"},
    ])
    _fx(fx_path, [])

    result, stats = build_monthly_values(positions_path, prices_path, fx_path, as_of=date(2021, 4, 5))

    row = result.iloc[0]
    assert row["MonthEnd"] == "2021-03-31"
    assert row["Year"] == 2021
    assert row["Month"] == 3
    assert row["MarketValueEUR"] == 1200.0
    # Real bug caught by the user against real data: this used to assert
    # 2200.0 (market value MINUS a negative net_invested, i.e. added
    # together) instead of the real gain (value minus amount actually paid).
    assert row["UnrealizedGainEUR"] == 200.0
    assert row["UnrealizedGainPercent"] == 20.0
    assert row["FXRate"] == 1.0


def test_unrealized_gain_percent_blank_when_net_invested_is_zero(tmp_path):
    """A fully-sold-then-recovered-exactly (or genuinely unknown cost basis)
    position has nothing sensible to divide by - blank, not an error."""
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"
    fx_path = tmp_path / "FXRates.xlsx"

    _positions(positions_path, [
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "SAMPO A", "Date": "2021-03-15", "CumulativeQuantity": 100, "CumulativeNetInvested": 0},
    ])
    _prices(prices_path, [
        {"NormalizedInstrument": "SAMPO A", "Date": "2021-03-31", "Close": 12.0, "Currency": "EUR"},
    ])
    _fx(fx_path, [])

    result, stats = build_monthly_values(positions_path, prices_path, fx_path, as_of=date(2021, 4, 5))

    assert result.iloc[0]["UnrealizedGainPercent"] == ""


def test_portfolio_owner_carried_through_from_positions(tmp_path):
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"
    fx_path = tmp_path / "FXRates.xlsx"

    _positions(positions_path, [
        {"Broker": "NORDNET", "Portfolio": "1", "PortfolioOwner": "PERSON_A", "NormalizedInstrument": "SAMPO A", "Date": "2021-03-15", "CumulativeQuantity": 100, "CumulativeNetInvested": -1000},
    ])
    _prices(prices_path, [
        {"NormalizedInstrument": "SAMPO A", "Date": "2021-03-31", "Close": 12.0, "Currency": "EUR"},
    ])
    _fx(fx_path, [])

    result, stats = build_monthly_values(positions_path, prices_path, fx_path, as_of=date(2021, 4, 5))

    assert result.iloc[0]["PortfolioOwner"] == "PERSON_A"


def test_portfolio_owner_blank_when_positions_sheet_predates_the_column(tmp_path):
    """PortfolioPositions.xlsx files written before this column existed must
    not error - just carry an empty PortfolioOwner through."""
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"
    fx_path = tmp_path / "FXRates.xlsx"

    _write_sheet(
        positions_path,
        "PortfolioPositions",
        ["Broker", "Portfolio", "NormalizedInstrument", "Date", "TransactionType", "QuantityDelta", "CumulativeQuantity", "CashDelta", "CumulativeNetInvested"],
        [{"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "SAMPO A", "Date": "2021-03-15", "CumulativeQuantity": 100, "CumulativeNetInvested": -1000}],
    )
    _prices(prices_path, [
        {"NormalizedInstrument": "SAMPO A", "Date": "2021-03-31", "Close": 12.0, "Currency": "EUR"},
    ])
    _fx(fx_path, [])

    result, stats = build_monthly_values(positions_path, prices_path, fx_path, as_of=date(2021, 4, 5))

    assert result.iloc[0]["PortfolioOwner"] == ""


def test_non_eur_instrument_converts_using_fx_rate(tmp_path):
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"
    fx_path = tmp_path / "FXRates.xlsx"

    _positions(positions_path, [
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "NORWEGIAN AIR SHUTTLE", "Date": "2021-03-15", "CumulativeQuantity": 30, "CumulativeNetInvested": -159.26},
    ])
    _prices(prices_path, [
        {"NormalizedInstrument": "NORWEGIAN AIR SHUTTLE", "Date": "2021-03-31", "Close": 13.0, "Currency": "NOK"},
    ])
    _fx(fx_path, [
        {"Currency": "NOK", "Date": "2021-03-31", "Rate": 0.1},
    ])

    result, stats = build_monthly_values(positions_path, prices_path, fx_path, as_of=date(2021, 4, 5))

    row = result.iloc[0]
    # 30 * 13.0 * 0.1 = 39.0
    assert row["MarketValueEUR"] == 39.0
    assert row["FXRate"] == 0.1


def test_price_and_fx_lookups_backfill_to_most_recent_prior_date(tmp_path):
    """Month-end (e.g. a weekend) may not have an exact price/FX row - use the
    most recent one on or before it, not an exact-date match."""
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"
    fx_path = tmp_path / "FXRates.xlsx"

    _positions(positions_path, [
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "NORWEGIAN AIR SHUTTLE", "Date": "2021-03-15", "CumulativeQuantity": 10, "CumulativeNetInvested": -50},
    ])
    _prices(prices_path, [
        {"NormalizedInstrument": "NORWEGIAN AIR SHUTTLE", "Date": "2021-03-29", "Close": 10.0, "Currency": "NOK"},
    ])
    _fx(fx_path, [
        {"Currency": "NOK", "Date": "2021-03-28", "Rate": 0.1},
    ])

    result, stats = build_monthly_values(positions_path, prices_path, fx_path, as_of=date(2021, 4, 5))

    row = result.iloc[0]
    assert row["PriceDate"] == "2021-03-29"
    assert row["FXDate"] == "2021-03-28"
    assert row["MarketValueEUR"] == 10.0


def test_fully_sold_out_instrument_stops_appearing(tmp_path):
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"
    fx_path = tmp_path / "FXRates.xlsx"

    _positions(positions_path, [
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "SAMPO A", "Date": "2021-01-15", "CumulativeQuantity": 100, "CumulativeNetInvested": -1000},
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "SAMPO A", "Date": "2021-02-15", "CumulativeQuantity": 0, "CumulativeNetInvested": 200},
    ])
    _prices(prices_path, [
        {"NormalizedInstrument": "SAMPO A", "Date": "2021-01-31", "Close": 10.0, "Currency": "EUR"},
        {"NormalizedInstrument": "SAMPO A", "Date": "2021-02-28", "Close": 11.0, "Currency": "EUR"},
        {"NormalizedInstrument": "SAMPO A", "Date": "2021-03-31", "Close": 12.0, "Currency": "EUR"},
    ])
    _fx(fx_path, [])

    result, stats = build_monthly_values(positions_path, prices_path, fx_path, as_of=date(2021, 4, 5))

    assert list(result["MonthEnd"]) == ["2021-01-31"]


def test_only_fully_completed_months_included_not_current_in_progress(tmp_path):
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"
    fx_path = tmp_path / "FXRates.xlsx"

    _positions(positions_path, [
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "SAMPO A", "Date": "2021-03-15", "CumulativeQuantity": 100, "CumulativeNetInvested": -1000},
    ])
    _prices(prices_path, [
        {"NormalizedInstrument": "SAMPO A", "Date": "2021-03-31", "Close": 12.0, "Currency": "EUR"},
        {"NormalizedInstrument": "SAMPO A", "Date": "2021-04-10", "Close": 13.0, "Currency": "EUR"},
    ])
    _fx(fx_path, [])

    result, stats = build_monthly_values(positions_path, prices_path, fx_path, as_of=date(2021, 4, 15))

    assert list(result["MonthEnd"]) == ["2021-03-31"]
    assert stats["last_month_end"] == "2021-03-31"


def test_missing_price_is_reported_and_row_skipped(tmp_path):
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"
    fx_path = tmp_path / "FXRates.xlsx"

    _positions(positions_path, [
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "SAMPO A", "Date": "2021-01-15", "CumulativeQuantity": 100, "CumulativeNetInvested": -1000},
    ])
    _prices(prices_path, [])
    _fx(fx_path, [])

    result, stats = build_monthly_values(positions_path, prices_path, fx_path, as_of=date(2021, 2, 5))

    assert len(result) == 0
    assert stats["missing_price_instrument_months"] == [("SAMPO A", "2021-01")]


def test_missing_fx_rate_is_reported_and_row_skipped(tmp_path):
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"
    fx_path = tmp_path / "FXRates.xlsx"

    _positions(positions_path, [
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "NORWEGIAN AIR SHUTTLE", "Date": "2021-01-15", "CumulativeQuantity": 10, "CumulativeNetInvested": -50},
    ])
    _prices(prices_path, [
        {"NormalizedInstrument": "NORWEGIAN AIR SHUTTLE", "Date": "2021-01-31", "Close": 10.0, "Currency": "NOK"},
    ])
    _fx(fx_path, [])

    result, stats = build_monthly_values(positions_path, prices_path, fx_path, as_of=date(2021, 2, 5))

    assert len(result) == 0
    assert stats["missing_fx_currency_months"] == [("NOK", "2021-01")]


def test_dividend_appears_in_its_own_month_and_carries_into_the_cumulative_total(tmp_path):
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"
    fx_path = tmp_path / "FXRates.xlsx"
    investments_path = tmp_path / "ParsedInvestments.xlsx"

    _positions(positions_path, [
        {"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "TELIA COMPANY AB", "Date": "2021-01-05", "CumulativeQuantity": 100, "CumulativeNetInvested": -1000},
    ])
    _prices(prices_path, [
        {"NormalizedInstrument": "TELIA COMPANY AB", "Date": "2021-01-31", "Close": 10.0, "Currency": "EUR"},
        {"NormalizedInstrument": "TELIA COMPANY AB", "Date": "2021-02-28", "Close": 10.0, "Currency": "EUR"},
        {"NormalizedInstrument": "TELIA COMPANY AB", "Date": "2021-03-31", "Close": 10.0, "Currency": "EUR"},
    ])
    _fx(fx_path, [])
    _transactions(investments_path, [
        {"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "TELIA COMPANY AB", "TransactionType": "DIVIDEND", "TradeDate": "2021-02-15", "CashAmount": 100.0},
        {"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "TELIA COMPANY AB", "TransactionType": "TAX", "TradeDate": "2021-02-15", "CashAmount": -15.0},
    ])

    result, stats = build_monthly_values(positions_path, prices_path, fx_path, investments_path, as_of=date(2021, 4, 5))
    by_month = result.set_index("MonthEnd")

    assert by_month.loc["2021-01-31", "DividendGrossEUR"] == 0.0
    assert by_month.loc["2021-01-31", "CumulativeDividendGrossEUR"] == 0.0

    assert by_month.loc["2021-02-28", "DividendGrossEUR"] == 100.0
    assert by_month.loc["2021-02-28", "DividendTaxEUR"] == 15.0
    assert by_month.loc["2021-02-28", "DividendNetEUR"] == 85.0
    assert by_month.loc["2021-02-28", "CumulativeDividendGrossEUR"] == 100.0
    assert by_month.loc["2021-02-28", "CumulativeDividendNetEUR"] == 85.0

    # March had no dividend of its own, but the cumulative total carries forward.
    assert by_month.loc["2021-03-31", "DividendGrossEUR"] == 0.0
    assert by_month.loc["2021-03-31", "CumulativeDividendGrossEUR"] == 100.0
    assert by_month.loc["2021-03-31", "CumulativeDividendNetEUR"] == 85.0


def test_dividend_for_a_different_instrument_does_not_leak_across_groups(tmp_path):
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"
    fx_path = tmp_path / "FXRates.xlsx"
    investments_path = tmp_path / "ParsedInvestments.xlsx"

    _positions(positions_path, [
        {"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "TELIA COMPANY AB", "Date": "2021-01-05", "CumulativeQuantity": 100, "CumulativeNetInvested": -1000},
        {"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "FORTUM", "Date": "2021-01-05", "CumulativeQuantity": 50, "CumulativeNetInvested": -500},
    ])
    _prices(prices_path, [
        {"NormalizedInstrument": "TELIA COMPANY AB", "Date": "2021-01-31", "Close": 10.0, "Currency": "EUR"},
        {"NormalizedInstrument": "FORTUM", "Date": "2021-01-31", "Close": 10.0, "Currency": "EUR"},
    ])
    _fx(fx_path, [])
    _transactions(investments_path, [
        {"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "FORTUM", "TransactionType": "DIVIDEND", "TradeDate": "2021-01-15", "CashAmount": 120.0},
    ])

    result, stats = build_monthly_values(positions_path, prices_path, fx_path, investments_path, as_of=date(2021, 2, 5))
    by_instrument = result.set_index("NormalizedInstrument")

    assert by_instrument.loc["FORTUM", "DividendGrossEUR"] == 120.0
    assert by_instrument.loc["TELIA COMPANY AB", "DividendGrossEUR"] == 0.0


def test_no_investments_workbook_defaults_dividend_columns_to_zero(tmp_path):
    """Backward compatible: an omitted/missing investments_workbook must not
    error, just report zero dividend activity everywhere."""
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"
    fx_path = tmp_path / "FXRates.xlsx"

    _positions(positions_path, [
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "SAMPO A", "Date": "2021-01-15", "CumulativeQuantity": 100, "CumulativeNetInvested": -1000},
    ])
    _prices(prices_path, [
        {"NormalizedInstrument": "SAMPO A", "Date": "2021-01-31", "Close": 10.0, "Currency": "EUR"},
    ])
    _fx(fx_path, [])

    result, stats = build_monthly_values(positions_path, prices_path, fx_path, as_of=date(2021, 2, 5))

    assert result.iloc[0]["DividendGrossEUR"] == 0.0
    assert result.iloc[0]["CumulativeDividendNetEUR"] == 0.0


def test_portfolio_type_carried_through_from_positions(tmp_path):
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"
    fx_path = tmp_path / "FXRates.xlsx"

    _write_sheet(
        positions_path,
        "PortfolioPositions",
        POSITIONS_HEADERS + ["PortfolioType"],
        [{"Broker": "NORDNET", "Portfolio": "1", "PortfolioOwner": "PERSON_A", "PortfolioType": "OSAKESAASTOTILI", "NormalizedInstrument": "SAMPO A", "Date": "2021-03-15", "CumulativeQuantity": 100, "CumulativeNetInvested": -1000}],
    )
    _prices(prices_path, [
        {"NormalizedInstrument": "SAMPO A", "Date": "2021-03-31", "Close": 12.0, "Currency": "EUR"},
    ])
    _fx(fx_path, [])

    result, stats = build_monthly_values(positions_path, prices_path, fx_path, as_of=date(2021, 4, 5))

    assert result.iloc[0]["PortfolioType"] == "OSAKESAASTOTILI"


def test_portfolio_type_blank_when_positions_sheet_predates_the_column(tmp_path):
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"
    fx_path = tmp_path / "FXRates.xlsx"

    _positions(positions_path, [
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "SAMPO A", "Date": "2021-03-15", "CumulativeQuantity": 100, "CumulativeNetInvested": -1000},
    ])
    _prices(prices_path, [
        {"NormalizedInstrument": "SAMPO A", "Date": "2021-03-31", "Close": 12.0, "Currency": "EUR"},
    ])
    _fx(fx_path, [])

    result, stats = build_monthly_values(positions_path, prices_path, fx_path, as_of=date(2021, 4, 5))

    assert result.iloc[0]["PortfolioType"] == ""


def test_instrument_type_carried_through_from_positions(tmp_path):
    """InstrumentType (STOCK/FUND/CRYPTO/...) enables a stocks/funds/cash
    split directly from MonthlyPositionValue.xlsx, without a separate lookup
    against InstrumentMaster.xlsx."""
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"
    fx_path = tmp_path / "FXRates.xlsx"

    _write_sheet(
        positions_path,
        "PortfolioPositions",
        POSITIONS_HEADERS + ["InstrumentType"],
        [{"Broker": "NORDNET", "Portfolio": "1", "PortfolioOwner": "PERSON_A", "InstrumentType": "STOCK", "NormalizedInstrument": "SAMPO A", "Date": "2021-03-15", "CumulativeQuantity": 100, "CumulativeNetInvested": -1000}],
    )
    _prices(prices_path, [
        {"NormalizedInstrument": "SAMPO A", "Date": "2021-03-31", "Close": 12.0, "Currency": "EUR"},
    ])
    _fx(fx_path, [])

    result, stats = build_monthly_values(positions_path, prices_path, fx_path, as_of=date(2021, 4, 5))

    assert result.iloc[0]["InstrumentType"] == "STOCK"


def test_instrument_type_blank_when_positions_sheet_predates_the_column(tmp_path):
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"
    fx_path = tmp_path / "FXRates.xlsx"

    _positions(positions_path, [
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "SAMPO A", "Date": "2021-03-15", "CumulativeQuantity": 100, "CumulativeNetInvested": -1000},
    ])
    _prices(prices_path, [
        {"NormalizedInstrument": "SAMPO A", "Date": "2021-03-31", "Close": 12.0, "Currency": "EUR"},
    ])
    _fx(fx_path, [])

    result, stats = build_monthly_values(positions_path, prices_path, fx_path, as_of=date(2021, 4, 5))

    assert result.iloc[0]["InstrumentType"] == ""


def test_load_cash_balance_source_filters_to_tracked_brokers_only(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_sheet(path, "RawInvestmentTransactions", RAW_TRANSACTIONS_HEADERS, [
        {"Broker": "NORDNET", "Portfolio": "1", "TradeDate": "2021-01-15", "CashAmount": 100.0, "CashBalance": 100.0},
        {"Broker": "EVLI", "Portfolio": "EVLI", "TradeDate": "2021-01-15", "CashAmount": 50.0, "CashBalance": ""},
        {"Broker": "OP", "Portfolio": "OP", "TradeDate": "2021-01-15", "CashAmount": 20.0, "CashBalance": ""},
        {"Broker": "SELIGSON", "Portfolio": "SELIGSON", "TradeDate": "2021-01-15", "CashAmount": 20.0, "CashBalance": ""},
    ])

    source = load_cash_balance_source(path)

    assert set(source["Broker"]) == {"NORDNET", "EVLI"}


def test_load_cash_balance_source_missing_sheet_returns_empty(tmp_path):
    """A workbook written before RawInvestmentTransactions existed (or a
    test fixture that only sets up InvestmentTransactions) must not error."""
    path = tmp_path / "ParsedInvestments.xlsx"
    _transactions(path, [])

    source = load_cash_balance_source(path)

    assert len(source) == 0


def test_cash_balance_rows_nordnet_uses_raw_cash_balance_field_directly(tmp_path):
    source = pd.DataFrame([
        {"Broker": "NORDNET", "Portfolio": "1", "TradeDate": pd.Timestamp("2021-01-10"), "CashAmount": -100.0, "CashBalance": 400.0},
        {"Broker": "NORDNET", "Portfolio": "1", "TradeDate": pd.Timestamp("2021-02-05"), "CashAmount": 50.0, "CashBalance": 450.0},
    ])

    loaded = _with_cash_settings(tmp_path)
    old_cache = settings_module._SETTINGS_CACHE
    try:
        settings_module._SETTINGS_CACHE = loaded
        rows = build_cash_balance_rows(source, pd.Timestamp("2021-02-28"))
    finally:
        settings_module._SETTINGS_CACHE = old_cache

    by_month = {r["MonthEnd"]: r for r in rows}
    assert by_month["2021-01-31"]["MarketValueEUR"] == 400.0
    assert by_month["2021-02-28"]["MarketValueEUR"] == 450.0
    assert by_month["2021-01-31"]["PortfolioOwner"] == "PERSON_A"
    assert by_month["2021-01-31"]["PortfolioType"] == "OSAKESAASTOTILI"
    assert by_month["2021-01-31"]["NormalizedInstrument"] == "CASH"
    assert by_month["2021-01-31"]["InstrumentType"] == "CASH"


def test_cash_balance_rows_evli_reconstructed_via_cumsum(tmp_path):
    """EVLI's export never populates CashBalance - reconstruct it as a
    running cumsum of every CashAmount seen (excluding sell proceeds, see
    the dedicated test below), since there's no other record."""
    source = pd.DataFrame([
        {"Broker": "EVLI", "Portfolio": "EVLI", "TradeDate": pd.Timestamp("2021-01-10"), "CashAmount": 175.0, "CashBalance": None, "TransactionTypeRaw": "SAVINGS"},
        {"Broker": "EVLI", "Portfolio": "EVLI", "TradeDate": pd.Timestamp("2021-02-10"), "CashAmount": 175.0, "CashBalance": None, "TransactionTypeRaw": "SAVINGS"},
        {"Broker": "EVLI", "Portfolio": "EVLI", "TradeDate": pd.Timestamp("2021-02-20"), "CashAmount": -300.0, "CashBalance": None, "TransactionTypeRaw": "SHARE PURCHASE"},
    ])

    loaded = _with_cash_settings(tmp_path)
    old_cache = settings_module._SETTINGS_CACHE
    try:
        settings_module._SETTINGS_CACHE = loaded
        rows = build_cash_balance_rows(source, pd.Timestamp("2021-02-28"))
    finally:
        settings_module._SETTINGS_CACHE = old_cache

    by_month = {r["MonthEnd"]: r for r in rows}
    assert by_month["2021-01-31"]["MarketValueEUR"] == 175.0
    # 175 + 175 - 300 = 50
    assert by_month["2021-02-28"]["MarketValueEUR"] == 50.0


def test_cash_balance_rows_evli_excludes_sell_proceeds(tmp_path):
    """
    Real bug: EVLI's Sell/Sell of purchased share proceeds are wired
    straight to the linked bank account at settlement - they never sit in
    EVLI as cash. Confirmed on the budgeting side: two real EVLI
    withdrawals landed in the bank within days of a sale, each close to (one
    exactly, minus a flat fee) that sale's total proceeds. Including these
    in the cumsum made the reconstructed balance grow forever with no
    matching outflow - caught by the user as an impossible ever-growing
    figure.
    """
    source = pd.DataFrame([
        {"Broker": "EVLI", "Portfolio": "EVLI", "TradeDate": pd.Timestamp("2021-01-10"), "CashAmount": 200.0, "CashBalance": None, "TransactionTypeRaw": "SAVINGS"},
        {"Broker": "EVLI", "Portfolio": "EVLI", "TradeDate": pd.Timestamp("2021-02-10"), "CashAmount": 5000.0, "CashBalance": None, "TransactionTypeRaw": "SELL"},
        {"Broker": "EVLI", "Portfolio": "EVLI", "TradeDate": pd.Timestamp("2021-02-11"), "CashAmount": 3000.0, "CashBalance": None, "TransactionTypeRaw": "SELL OF PURCHASED SHARE"},
    ])

    loaded = _with_cash_settings(tmp_path)
    old_cache = settings_module._SETTINGS_CACHE
    try:
        settings_module._SETTINGS_CACHE = loaded
        rows = build_cash_balance_rows(source, pd.Timestamp("2021-02-28"))
    finally:
        settings_module._SETTINGS_CACHE = old_cache

    by_month = {r["MonthEnd"]: r for r in rows}
    # Both sell rows are excluded entirely - balance stays at the deposit
    # amount, not 200 + 5000 + 3000.
    assert by_month["2021-01-31"]["MarketValueEUR"] == 200.0
    assert by_month["2021-02-28"]["MarketValueEUR"] == 200.0


def test_cash_balance_rows_evli_excludes_allocated_bonus(tmp_path):
    """
    Real bug: EVLI's Allocated plan-cycle bonus never becomes real,
    spendable cash - it converts directly into free Matching shares.
    Confirmed across four independent real plan cycles: each cycle's
    Savings/Share-purchase pairs net to exactly 0, leaving only the
    Allocated amount as a permanent, never-spent residual - a clear sign
    it was never real liquid cash to begin with.
    """
    source = pd.DataFrame([
        {"Broker": "EVLI", "Portfolio": "EVLI", "TradeDate": pd.Timestamp("2021-01-01"), "CashAmount": 200.0, "CashBalance": None, "TransactionTypeRaw": "ALLOCATED"},
        {"Broker": "EVLI", "Portfolio": "EVLI", "TradeDate": pd.Timestamp("2021-01-10"), "CashAmount": 175.0, "CashBalance": None, "TransactionTypeRaw": "SAVINGS"},
        {"Broker": "EVLI", "Portfolio": "EVLI", "TradeDate": pd.Timestamp("2021-02-10"), "CashAmount": 175.0, "CashBalance": None, "TransactionTypeRaw": "SAVINGS"},
        {"Broker": "EVLI", "Portfolio": "EVLI", "TradeDate": pd.Timestamp("2021-02-20"), "CashAmount": -350.0, "CashBalance": None, "TransactionTypeRaw": "SHARE PURCHASE"},
    ])

    loaded = _with_cash_settings(tmp_path)
    old_cache = settings_module._SETTINGS_CACHE
    try:
        settings_module._SETTINGS_CACHE = loaded
        rows = build_cash_balance_rows(source, pd.Timestamp("2021-02-28"))
    finally:
        settings_module._SETTINGS_CACHE = old_cache

    by_month = {r["MonthEnd"]: r for r in rows}
    # The Allocated 200 is excluded entirely - only Savings/Share purchase
    # net (175 + 175 - 350 = 0), not 200 + 175 + 175 - 350 = 200.
    assert by_month["2021-02-28"]["MarketValueEUR"] == 0.0


def test_cash_balance_rows_zero_balance_kept_not_dropped(tmp_path):
    source = pd.DataFrame([
        {"Broker": "EVLI", "Portfolio": "EVLI", "TradeDate": pd.Timestamp("2021-01-10"), "CashAmount": 100.0, "CashBalance": None, "TransactionTypeRaw": "SAVINGS"},
        {"Broker": "EVLI", "Portfolio": "EVLI", "TradeDate": pd.Timestamp("2021-01-15"), "CashAmount": -100.0, "CashBalance": None, "TransactionTypeRaw": "SHARE PURCHASE"},
    ])

    loaded = _with_cash_settings(tmp_path)
    old_cache = settings_module._SETTINGS_CACHE
    try:
        settings_module._SETTINGS_CACHE = loaded
        rows = build_cash_balance_rows(source, pd.Timestamp("2021-01-31"))
    finally:
        settings_module._SETTINGS_CACHE = old_cache

    assert len(rows) == 1
    assert rows[0]["MarketValueEUR"] == 0.0


def test_cash_balance_rows_excludes_months_before_first_event(tmp_path):
    source = pd.DataFrame([
        {"Broker": "EVLI", "Portfolio": "EVLI", "TradeDate": pd.Timestamp("2021-03-10"), "CashAmount": 100.0, "CashBalance": None, "TransactionTypeRaw": "SAVINGS"},
    ])

    loaded = _with_cash_settings(tmp_path)
    old_cache = settings_module._SETTINGS_CACHE
    try:
        settings_module._SETTINGS_CACHE = loaded
        rows = build_cash_balance_rows(source, pd.Timestamp("2021-04-30"))
    finally:
        settings_module._SETTINGS_CACHE = old_cache

    assert [r["MonthEnd"] for r in rows] == ["2021-03-31", "2021-04-30"]


def test_build_monthly_values_includes_cash_balance_rows_alongside_instruments(tmp_path):
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"
    fx_path = tmp_path / "FXRates.xlsx"
    investments_path = tmp_path / "ParsedInvestments.xlsx"

    _positions(positions_path, [
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "SAMPO A", "Date": "2021-01-15", "CumulativeQuantity": 100, "CumulativeNetInvested": -1000},
    ])
    _prices(prices_path, [
        {"NormalizedInstrument": "SAMPO A", "Date": "2021-01-31", "Close": 10.0, "Currency": "EUR"},
    ])
    _fx(fx_path, [])
    _transactions(investments_path, [])
    _raw_transactions(investments_path, [
        {"Broker": "NORDNET", "Portfolio": "1", "TradeDate": "2021-01-20", "CashAmount": 200.0, "CashBalance": 200.0},
    ])

    loaded = _with_cash_settings(tmp_path)
    old_cache = settings_module._SETTINGS_CACHE
    try:
        settings_module._SETTINGS_CACHE = loaded
        result, stats = build_monthly_values(positions_path, prices_path, fx_path, investments_path, as_of=date(2021, 2, 5))
    finally:
        settings_module._SETTINGS_CACHE = old_cache

    assert stats["cash_balance_rows"] == 1
    cash_rows = result[result["NormalizedInstrument"] == "CASH"]
    assert len(cash_rows) == 1
    assert cash_rows.iloc[0]["MarketValueEUR"] == 200.0
    instrument_rows = result[result["NormalizedInstrument"] == "SAMPO A"]
    assert len(instrument_rows) == 1


def test_no_cash_balance_rows_for_untracked_brokers(tmp_path):
    """OP/Nordea (bank-tracked elsewhere) and Seligson (direct-purchase only,
    no cash account) never produce a CASH row."""
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"
    fx_path = tmp_path / "FXRates.xlsx"
    investments_path = tmp_path / "ParsedInvestments.xlsx"

    _positions(positions_path, [
        {"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "SAMPO A", "Date": "2021-01-15", "CumulativeQuantity": 100, "CumulativeNetInvested": -1000},
    ])
    _prices(prices_path, [
        {"NormalizedInstrument": "SAMPO A", "Date": "2021-01-31", "Close": 10.0, "Currency": "EUR"},
    ])
    _fx(fx_path, [])
    _transactions(investments_path, [])
    _raw_transactions(investments_path, [
        {"Broker": "OP", "Portfolio": "OP", "TradeDate": "2021-01-20", "CashAmount": 200.0, "CashBalance": 200.0},
        {"Broker": "SELIGSON", "Portfolio": "SELIGSON", "TradeDate": "2021-01-20", "CashAmount": 200.0, "CashBalance": 200.0},
    ])

    result, stats = build_monthly_values(positions_path, prices_path, fx_path, investments_path, as_of=date(2021, 2, 5))

    assert stats["cash_balance_rows"] == 0
    assert "CASH" not in set(result["NormalizedInstrument"])
