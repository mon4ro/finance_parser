from openpyxl import Workbook

from finance_parser.investments.build_look_through_exposure import (
    build_look_through_exposure,
    load_position_values,
)


INSTRUMENT_MASTER_HEADERS = [
    "Enabled", "MatchPriority", "Broker", "ISIN", "RawInstrumentName",
    "NormalizedInstrument", "InstrumentType", "AssetClass", "Currency",
    "Ticker", "Exchange", "PriceSource", "PriceSymbol", "Notes",
]

MONTHLY_VALUE_HEADERS = [
    "MonthEnd", "Year", "Month", "Broker", "Portfolio", "PortfolioOwner", "PortfolioType",
    "NormalizedInstrument", "CumulativeQuantity", "PriceLocal", "InstrumentCurrency",
    "PriceDate", "FXRate", "FXDate", "MarketValueEUR",
    "CumulativeNetInvested", "UnrealizedGainEUR", "UnrealizedGainPercent",
    "DividendGrossEUR", "DividendTaxEUR", "DividendNetEUR",
    "CumulativeDividendGrossEUR", "CumulativeDividendNetEUR",
]

HOLDINGS_HEADERS = [
    "SnapshotMonth", "SnapshotDate", "PriceSymbol", "NormalizedInstrument",
    "HoldingRank", "HoldingSymbol", "HoldingName", "HoldingPercent", "SnapshotSource",
]


def _write_master(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "InstrumentMaster"
    ws.append(INSTRUMENT_MASTER_HEADERS)
    for row in rows:
        ws.append([row.get(h, "") for h in INSTRUMENT_MASTER_HEADERS])
    wb.save(path)


def _write_monthly_value(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "MonthlyPositionValue"
    ws.append(MONTHLY_VALUE_HEADERS)
    for row in rows:
        ws.append([row.get(h, "") for h in MONTHLY_VALUE_HEADERS])
    wb.save(path)


def _write_holdings(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "FundHoldingsSnapshot"
    ws.append(HOLDINGS_HEADERS)
    for row in rows:
        ws.append([row.get(h, "") for h in HOLDINGS_HEADERS])
    wb.save(path)


def test_load_position_values_excludes_cash_and_uses_latest_month(tmp_path):
    path = tmp_path / "MonthlyPositionValue.xlsx"
    _write_monthly_value(path, [
        {"MonthEnd": "2026-01-31", "NormalizedInstrument": "SAMPLE STOCK", "MarketValueEUR": 100.0},
        {"MonthEnd": "2026-02-28", "NormalizedInstrument": "SAMPLE STOCK", "MarketValueEUR": 120.0},
        {"MonthEnd": "2026-02-28", "NormalizedInstrument": "CASH", "MarketValueEUR": 50.0},
    ])

    month, positions = load_position_values(path)

    assert month == "2026-02"
    assert list(positions["NormalizedInstrument"]) == ["SAMPLE STOCK"]
    assert positions.iloc[0]["MarketValueEUR"] == 120.0


def test_direct_stock_exposure_uses_price_symbol_as_holding_key(tmp_path):
    master_path = tmp_path / "InstrumentMaster.xlsx"
    monthly_value_path = tmp_path / "MonthlyPositionValue.xlsx"
    holdings_path = tmp_path / "FundHoldingsSnapshot.xlsx"

    _write_master(master_path, [
        {"Enabled": "YES", "InstrumentType": "STOCK", "NormalizedInstrument": "SAMPLE CO", "PriceSymbol": "SC.HE"},
    ])
    _write_monthly_value(monthly_value_path, [
        {"MonthEnd": "2026-01-31", "NormalizedInstrument": "SAMPLE CO", "MarketValueEUR": 1000.0},
    ])
    _write_holdings(holdings_path, [])

    detail, summary, fund_coverage, stats = build_look_through_exposure(master_path, monthly_value_path, holdings_path)

    assert list(detail["HoldingKey"]) == ["SC.HE"]
    assert detail.iloc[0]["ExposureEUR"] == 1000.0
    assert summary.iloc[0]["HoldingKey"] == "SC.HE"
    assert summary.iloc[0]["ExposureEUR"] == 1000.0
    assert summary.iloc[0]["ExposurePercent"] == 100.0


def test_fund_exposure_splits_by_holding_percent(tmp_path):
    master_path = tmp_path / "InstrumentMaster.xlsx"
    monthly_value_path = tmp_path / "MonthlyPositionValue.xlsx"
    holdings_path = tmp_path / "FundHoldingsSnapshot.xlsx"

    _write_master(master_path, [
        {"Enabled": "YES", "InstrumentType": "FUND", "NormalizedInstrument": "SAMPLE FUND", "PriceSymbol": "SF.HE"},
    ])
    _write_monthly_value(monthly_value_path, [
        {"MonthEnd": "2026-01-31", "NormalizedInstrument": "SAMPLE FUND", "MarketValueEUR": 1000.0},
    ])
    _write_holdings(holdings_path, [
        {"SnapshotMonth": "2026-01", "PriceSymbol": "SF.HE", "NormalizedInstrument": "SAMPLE FUND",
         "HoldingRank": 1, "HoldingSymbol": "AAA", "HoldingName": "Company A", "HoldingPercent": 30.0, "SnapshotSource": "LIVE_FETCH"},
        {"SnapshotMonth": "2026-01", "PriceSymbol": "SF.HE", "NormalizedInstrument": "SAMPLE FUND",
         "HoldingRank": 11, "HoldingSymbol": "UNKNOWN", "HoldingName": "Unmapped remainder (beyond top 10)", "HoldingPercent": 70.0, "SnapshotSource": "LIVE_FETCH"},
    ])

    detail, summary, fund_coverage, stats = build_look_through_exposure(master_path, monthly_value_path, holdings_path)

    aaa_row = detail[detail["HoldingKey"] == "AAA"].iloc[0]
    assert aaa_row["ExposureEUR"] == 300.0
    unknown_row = detail[detail["HoldingKey"] == "UNKNOWN:SAMPLE FUND"].iloc[0]
    assert unknown_row["ExposureEUR"] == 700.0
    assert stats["funds_missing_holdings_data"] == []

    coverage_row = fund_coverage.iloc[0]
    assert coverage_row["NormalizedInstrument"] == "SAMPLE FUND"
    assert coverage_row["KnownPercent"] == 30.0
    assert coverage_row["UnknownPercent"] == 70.0


def test_direct_and_indirect_exposure_to_same_company_merge_via_price_symbol(tmp_path):
    """
    The core look-through mechanic: a stock held both directly and inside a
    fund must combine into one total, keyed by the shared Yahoo PriceSymbol
    (InstrumentMaster's PriceSymbol matches exactly what a fund's own
    holdings report the same company as - confirmed against real household
    data before writing this test).
    """
    master_path = tmp_path / "InstrumentMaster.xlsx"
    monthly_value_path = tmp_path / "MonthlyPositionValue.xlsx"
    holdings_path = tmp_path / "FundHoldingsSnapshot.xlsx"

    _write_master(master_path, [
        {"Enabled": "YES", "InstrumentType": "STOCK", "NormalizedInstrument": "SAMPLE CO", "PriceSymbol": "SC.HE"},
        {"Enabled": "YES", "InstrumentType": "FUND", "NormalizedInstrument": "SAMPLE FUND", "PriceSymbol": "SF.HE"},
    ])
    _write_monthly_value(monthly_value_path, [
        {"MonthEnd": "2026-01-31", "NormalizedInstrument": "SAMPLE CO", "MarketValueEUR": 500.0},
        {"MonthEnd": "2026-01-31", "NormalizedInstrument": "SAMPLE FUND", "MarketValueEUR": 1000.0},
    ])
    _write_holdings(holdings_path, [
        {"SnapshotMonth": "2026-01", "PriceSymbol": "SF.HE", "NormalizedInstrument": "SAMPLE FUND",
         "HoldingRank": 1, "HoldingSymbol": "SC.HE", "HoldingName": "Sample Co", "HoldingPercent": 20.0, "SnapshotSource": "LIVE_FETCH"},
        {"SnapshotMonth": "2026-01", "PriceSymbol": "SF.HE", "NormalizedInstrument": "SAMPLE FUND",
         "HoldingRank": 11, "HoldingSymbol": "UNKNOWN", "HoldingName": "Unmapped remainder (beyond top 10)", "HoldingPercent": 80.0, "SnapshotSource": "LIVE_FETCH"},
    ])

    detail, summary, fund_coverage, stats = build_look_through_exposure(master_path, monthly_value_path, holdings_path)

    merged = summary[summary["HoldingKey"] == "SC.HE"].iloc[0]
    # Direct 500.0 + indirect 20% of 1000.0 = 700.0
    assert merged["ExposureEUR"] == 700.0


def test_unknown_buckets_from_different_funds_do_not_merge(tmp_path):
    """
    Real bug caught by the user: a shared global "UNKNOWN" key across every
    fund hides which specific fund is poorly covered - a fund with literally
    zero holdings data would look identical, once merged, to a well-covered
    fund's ordinary top-10 cutoff tail. Each fund's UNKNOWN residual must
    stay a distinct row (and be visible per-fund in FundCoverage too).
    """
    master_path = tmp_path / "InstrumentMaster.xlsx"
    monthly_value_path = tmp_path / "MonthlyPositionValue.xlsx"
    holdings_path = tmp_path / "FundHoldingsSnapshot.xlsx"

    _write_master(master_path, [
        {"Enabled": "YES", "InstrumentType": "FUND", "NormalizedInstrument": "WELL COVERED FUND", "PriceSymbol": "WCF.HE"},
        {"Enabled": "YES", "InstrumentType": "FUND", "NormalizedInstrument": "OPAQUE FUND", "PriceSymbol": "OF.HE"},
    ])
    _write_monthly_value(monthly_value_path, [
        {"MonthEnd": "2026-01-31", "NormalizedInstrument": "WELL COVERED FUND", "MarketValueEUR": 1000.0},
        {"MonthEnd": "2026-01-31", "NormalizedInstrument": "OPAQUE FUND", "MarketValueEUR": 1000.0},
    ])
    _write_holdings(holdings_path, [
        {"SnapshotMonth": "2026-01", "PriceSymbol": "WCF.HE", "NormalizedInstrument": "WELL COVERED FUND",
         "HoldingRank": 1, "HoldingSymbol": "AAA", "HoldingName": "Company A", "HoldingPercent": 90.0, "SnapshotSource": "LIVE_FETCH"},
        {"SnapshotMonth": "2026-01", "PriceSymbol": "WCF.HE", "NormalizedInstrument": "WELL COVERED FUND",
         "HoldingRank": 11, "HoldingSymbol": "UNKNOWN", "HoldingName": "Unmapped remainder (beyond top 10)", "HoldingPercent": 10.0, "SnapshotSource": "LIVE_FETCH"},
        {"SnapshotMonth": "2026-01", "PriceSymbol": "OF.HE", "NormalizedInstrument": "OPAQUE FUND",
         "HoldingRank": 11, "HoldingSymbol": "UNKNOWN", "HoldingName": "No holdings data available", "HoldingPercent": 100.0, "SnapshotSource": "LIVE_FETCH"},
    ])

    detail, summary, fund_coverage, stats = build_look_through_exposure(master_path, monthly_value_path, holdings_path)

    unknown_keys = set(summary[summary["HoldingKey"].str.startswith("UNKNOWN:")]["HoldingKey"])
    assert unknown_keys == {"UNKNOWN:WELL COVERED FUND", "UNKNOWN:OPAQUE FUND"}

    well_covered_unknown = summary[summary["HoldingKey"] == "UNKNOWN:WELL COVERED FUND"].iloc[0]
    assert well_covered_unknown["ExposureEUR"] == 100.0

    opaque_unknown = summary[summary["HoldingKey"] == "UNKNOWN:OPAQUE FUND"].iloc[0]
    assert opaque_unknown["ExposureEUR"] == 1000.0

    coverage_by_fund = fund_coverage.set_index("NormalizedInstrument")["UnknownPercent"]
    assert coverage_by_fund["WELL COVERED FUND"] == 10.0
    assert coverage_by_fund["OPAQUE FUND"] == 100.0


def test_fund_with_no_holdings_snapshot_counted_as_unknown_and_flagged(tmp_path):
    master_path = tmp_path / "InstrumentMaster.xlsx"
    monthly_value_path = tmp_path / "MonthlyPositionValue.xlsx"
    holdings_path = tmp_path / "FundHoldingsSnapshot.xlsx"

    _write_master(master_path, [
        {"Enabled": "YES", "InstrumentType": "FUND", "NormalizedInstrument": "UNTRACKED FUND", "PriceSymbol": "UF.HE"},
    ])
    _write_monthly_value(monthly_value_path, [
        {"MonthEnd": "2026-01-31", "NormalizedInstrument": "UNTRACKED FUND", "MarketValueEUR": 400.0},
    ])
    _write_holdings(holdings_path, [])

    detail, summary, fund_coverage, stats = build_look_through_exposure(master_path, monthly_value_path, holdings_path)

    assert stats["funds_missing_holdings_data"] == ["UNTRACKED FUND"]
    unknown_row = summary[summary["HoldingKey"] == "UNKNOWN:UNTRACKED FUND"].iloc[0]
    assert unknown_row["ExposureEUR"] == 400.0
    assert fund_coverage.iloc[0]["UnknownPercent"] == 100.0


def test_unclassified_instrument_treated_as_direct_and_flagged(tmp_path):
    master_path = tmp_path / "InstrumentMaster.xlsx"
    monthly_value_path = tmp_path / "MonthlyPositionValue.xlsx"
    holdings_path = tmp_path / "FundHoldingsSnapshot.xlsx"

    _write_master(master_path, [])
    _write_monthly_value(monthly_value_path, [
        {"MonthEnd": "2026-01-31", "NormalizedInstrument": "MYSTERY ASSET", "MarketValueEUR": 250.0},
    ])
    _write_holdings(holdings_path, [])

    detail, summary, fund_coverage, stats = build_look_through_exposure(master_path, monthly_value_path, holdings_path)

    assert stats["unclassified_instruments_treated_as_direct"] == ["MYSTERY ASSET"]
    assert detail.iloc[0]["HoldingKey"] == "MYSTERY ASSET"
    assert detail.iloc[0]["ExposureEUR"] == 250.0


def test_month_override_uses_nearest_available_fund_snapshot(tmp_path):
    master_path = tmp_path / "InstrumentMaster.xlsx"
    monthly_value_path = tmp_path / "MonthlyPositionValue.xlsx"
    holdings_path = tmp_path / "FundHoldingsSnapshot.xlsx"

    _write_master(master_path, [
        {"Enabled": "YES", "InstrumentType": "FUND", "NormalizedInstrument": "SAMPLE FUND", "PriceSymbol": "SF.HE"},
    ])
    _write_monthly_value(monthly_value_path, [
        {"MonthEnd": "2026-03-31", "NormalizedInstrument": "SAMPLE FUND", "MarketValueEUR": 1000.0},
    ])
    # Only a January snapshot exists - March should fall back to it rather
    # than report the fund as having no data at all.
    _write_holdings(holdings_path, [
        {"SnapshotMonth": "2026-01", "PriceSymbol": "SF.HE", "NormalizedInstrument": "SAMPLE FUND",
         "HoldingRank": 1, "HoldingSymbol": "AAA", "HoldingName": "Company A", "HoldingPercent": 100.0, "SnapshotSource": "LIVE_FETCH"},
    ])

    detail, summary, fund_coverage, stats = build_look_through_exposure(master_path, monthly_value_path, holdings_path, as_of_month="2026-03")

    assert stats["funds_missing_holdings_data"] == []
    assert detail.iloc[0]["HoldingKey"] == "AAA"
    assert detail.iloc[0]["ExposureEUR"] == 1000.0
