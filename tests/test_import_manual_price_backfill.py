from datetime import date

import openpyxl
import pandas as pd
from openpyxl import Workbook

from finance_parser.common import normalise_text
from finance_parser.investments.fetch_instrument_prices import PRICES_COLUMNS
from finance_parser.investments.import_manual_price_backfill import (
    MANUAL_PRICE_ROWS,
    backfill_gaps,
    read_manual_prices,
)
from finance_parser.investments.investment_common import INSTRUMENT_MASTER_COLUMNS


def test_manual_price_rows_keys_survive_normalise_text():
    """
    Real bug: MANUAL_PRICE_ROWS had a key with a trailing space
    ("Franklin Technology Fund A (Acc) "), but read_manual_prices() looks up
    normalise_text(cell_value) - which strips it - so the lookup silently
    failed and the entire instrument (110 real monthly values) never made it
    into the import, with no error or warning. Every key here must equal its
    own normalise_text() output, or the same silent-skip bug recurs for
    whichever row is mistyped.
    """
    for raw_name in MANUAL_PRICE_ROWS:
        assert raw_name == normalise_text(raw_name), (
            f"MANUAL_PRICE_ROWS key {raw_name!r} would never match after normalise_text() "
            f"(becomes {normalise_text(raw_name)!r}) - this row would be silently skipped"
        )


def _write_source(path, instrument_rows: dict[int, str], monthly_values: dict[int, dict[int, float]]):
    """instrument_rows: {row_idx: raw_name}. monthly_values: {row_idx: {col_idx: value}}."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Investments"

    # Header row (row 2): dates in a few columns.
    ws.cell(row=2, column=5, value=date(2021, 12, 31))
    ws.cell(row=2, column=7, value=date(2022, 1, 1))
    ws.cell(row=2, column=8, value=date(2022, 2, 1))

    for row_idx, raw_name in instrument_rows.items():
        ws.cell(row=row_idx, column=3, value=raw_name)
        for col_idx, value in monthly_values.get(row_idx, {}).items():
            ws.cell(row=row_idx, column=col_idx, value=value)

    wb.save(path)


def _write_prices(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "InstrumentPrices"
    ws.append(PRICES_COLUMNS)
    for row in rows:
        ws.append([row.get(h, "") for h in PRICES_COLUMNS])
    wb.save(path)


def _write_instrument_master(path, instrument, isin, currency, symbol):
    wb = Workbook()
    ws = wb.active
    ws.title = "InstrumentMaster"
    ws.append(INSTRUMENT_MASTER_COLUMNS)
    row = {"Enabled": "YES", "NormalizedInstrument": instrument, "ISIN": isin, "Currency": currency, "PriceSymbol": symbol}
    ws.append([row.get(h, "") for h in INSTRUMENT_MASTER_COLUMNS])
    wb.save(path)


def test_read_manual_prices_maps_raw_names_and_normalises_month_end(tmp_path):
    raw_name = list(MANUAL_PRICE_ROWS)[0]
    instrument = MANUAL_PRICE_ROWS[raw_name]
    source_path = tmp_path / "backfill.xlsm"

    _write_source(
        source_path,
        instrument_rows={12: raw_name},
        monthly_values={12: {5: 100.00, 7: 105.00, 8: 103.50}},
    )

    result = read_manual_prices(source_path)

    assert set(result["NormalizedInstrument"]) == {instrument}
    dates = sorted(result["Date"].dt.strftime("%Y-%m-%d"))
    # 2021-12-31 stays as-is (already month-end); 2022-01-01 -> 2022-01-31; 2022-02-01 -> 2022-02-28
    assert dates == ["2021-12-31", "2022-01-31", "2022-02-28"]


def test_read_manual_prices_ignores_unmapped_rows(tmp_path):
    source_path = tmp_path / "backfill.xlsm"
    _write_source(
        source_path,
        instrument_rows={12: "Some Unrelated Fund Not In Mapping"},
        monthly_values={12: {7: 100.0}},
    )

    result = read_manual_prices(source_path)

    assert len(result) == 0


def test_read_manual_prices_treats_zero_as_blank_not_a_real_price(tmp_path):
    """
    Real bug caught on the first real run: future/not-yet-reached months in
    the source file evaluate to a literal 0 (a formula fallback, not blank).
    A fund price of exactly EUR0.00 is never real and must be excluded like
    a blank cell, not imported as data.
    """
    raw_name = list(MANUAL_PRICE_ROWS)[0]
    source_path = tmp_path / "backfill.xlsm"

    _write_source(
        source_path,
        instrument_rows={12: raw_name},
        monthly_values={12: {5: 100.00, 7: 0, 8: -5.0}},
    )

    result = read_manual_prices(source_path)

    assert len(result) == 1
    assert result.iloc[0]["Close"] == 100.00


def test_backfill_gaps_only_adds_dates_not_already_present(tmp_path):
    raw_name = list(MANUAL_PRICE_ROWS)[0]
    instrument = MANUAL_PRICE_ROWS[raw_name]

    source_path = tmp_path / "backfill.xlsm"
    prices_path = tmp_path / "InstrumentPrices.xlsx"
    master_path = tmp_path / "InstrumentMaster.xlsx"

    _write_source(
        source_path,
        instrument_rows={12: raw_name},
        monthly_values={12: {5: 100.00, 7: 105.00}},
    )
    _write_prices(prices_path, [
        {"NormalizedInstrument": instrument, "Date": "2021-12-31", "Close": 999.0, "Currency": "EUR", "PriceSource": "YAHOO"},
    ])
    _write_instrument_master(master_path, instrument, "FI0000000000", "EUR", "TEST.F")

    combined, stats = backfill_gaps(source_path, prices_path, master_path, verbose=False)

    assert stats["rows_added"] == 1
    assert stats["rows_skipped_already_covered"] == 1

    new_row = combined[(combined["NormalizedInstrument"] == instrument) & (combined["Date"] == "2022-01-31")].iloc[0]
    assert new_row["Close"] == 105.00
    assert new_row["PriceSource"] == "MANUAL"
    assert new_row["Currency"] == "EUR"
    assert new_row["ISIN"] == "FI0000000000"

    existing_row = combined[(combined["NormalizedInstrument"] == instrument) & (combined["Date"] == "2021-12-31")].iloc[0]
    assert existing_row["Close"] == 999.0
    assert existing_row["PriceSource"] == "YAHOO"


def test_backfill_gaps_skips_non_numeric_cells(tmp_path):
    raw_name = list(MANUAL_PRICE_ROWS)[0]
    instrument = MANUAL_PRICE_ROWS[raw_name]

    source_path = tmp_path / "backfill.xlsm"
    prices_path = tmp_path / "InstrumentPrices.xlsx"
    master_path = tmp_path / "InstrumentMaster.xlsx"

    _write_source(
        source_path,
        instrument_rows={12: raw_name},
        monthly_values={12: {5: "n/a", 7: 105.00}},
    )
    _write_prices(prices_path, [])
    _write_instrument_master(master_path, instrument, "", "EUR", "TEST.F")

    combined, stats = backfill_gaps(source_path, prices_path, master_path, verbose=False)

    assert stats["rows_added"] == 1
