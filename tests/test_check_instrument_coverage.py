from datetime import date, timedelta

import pytest
from openpyxl import Workbook

from finance_parser.investments import check_instrument_coverage as check_module
from finance_parser.investments.check_instrument_coverage import find_coverage_gaps, has_gaps
from finance_parser.investments.fetch_instrument_prices import PRICES_COLUMNS
from finance_parser.investments.investment_common import INSTRUMENT_MASTER_COLUMNS


POSITIONS_HEADERS = [
    "Broker", "Portfolio", "NormalizedInstrument", "Date", "TransactionType",
    "QuantityDelta", "CumulativeQuantity", "CashDelta", "CumulativeNetInvested",
]


def _write_positions(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "PortfolioPositions"
    ws.append(POSITIONS_HEADERS)
    for row in rows:
        ws.append([row.get(h, "") for h in POSITIONS_HEADERS])
    wb.save(path)


def _write_master(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "InstrumentMaster"
    ws.append(INSTRUMENT_MASTER_COLUMNS)
    for row in rows:
        ws.append([row.get(h, "") for h in INSTRUMENT_MASTER_COLUMNS])
    wb.save(path)


def _write_prices(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "InstrumentPrices"
    ws.append(PRICES_COLUMNS)
    for row in rows:
        ws.append([row.get(h, "") for h in PRICES_COLUMNS])
    wb.save(path)


def test_fully_covered_instrument_reports_no_gaps(tmp_path):
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    master_path = tmp_path / "InstrumentMaster.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"

    _write_positions(positions_path, [
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "SAMPO A", "Date": "2021-01-15", "CumulativeQuantity": 100},
    ])
    _write_master(master_path, [
        {"Enabled": "YES", "NormalizedInstrument": "SAMPO A", "InstrumentType": "STOCK", "Currency": "EUR", "PriceSymbol": "SAMPO.HE"},
    ])
    _write_prices(prices_path, [
        {"NormalizedInstrument": "SAMPO A", "Date": date.today().isoformat(), "Close": 9.5, "Currency": "EUR"},
    ])

    gaps = find_coverage_gaps(positions_path, master_path, prices_path)

    assert not has_gaps(gaps)
    assert gaps["active_instrument_count"] == 1


def test_defunct_zero_quantity_instrument_never_flagged(tmp_path):
    """A fully-sold-out instrument (e.g. the real OP-Delta A gap) must not
    force this check to fail forever just because it's unclassified."""
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    master_path = tmp_path / "InstrumentMaster.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"

    _write_positions(positions_path, [
        {"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "OP-DELTA A", "Date": "2017-09-22", "CumulativeQuantity": 0},
    ])
    _write_master(master_path, [])
    _write_prices(prices_path, [])

    gaps = find_coverage_gaps(positions_path, master_path, prices_path)

    assert not has_gaps(gaps)
    assert gaps["active_instrument_count"] == 0


def test_unclassified_instrument_is_flagged(tmp_path):
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    master_path = tmp_path / "InstrumentMaster.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"

    _write_positions(positions_path, [
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "NEW STOCK", "Date": "2026-01-01", "CumulativeQuantity": 10},
    ])
    _write_master(master_path, [])
    _write_prices(prices_path, [])

    gaps = find_coverage_gaps(positions_path, master_path, prices_path)

    assert has_gaps(gaps)
    assert gaps["unclassified"] == ["NEW STOCK"]


def test_classified_but_no_price_symbol_is_unpriced_not_unclassified(tmp_path):
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    master_path = tmp_path / "InstrumentMaster.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"

    _write_positions(positions_path, [
        {"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "OP-EUROOPPA PIENYHTIÖT A", "Date": "2013-10-10", "CumulativeQuantity": 0.5},
    ])
    _write_master(master_path, [
        {"Enabled": "YES", "NormalizedInstrument": "OP-EUROOPPA PIENYHTIÖT A", "InstrumentType": "FUND", "Currency": "EUR", "PriceSymbol": ""},
    ])
    _write_prices(prices_path, [])

    gaps = find_coverage_gaps(positions_path, master_path, prices_path)

    assert gaps["unclassified"] == []
    assert gaps["unpriced"] == ["OP-EUROOPPA PIENYHTIÖT A"]


def test_stale_price_is_flagged_separately_from_unpriced(tmp_path):
    """Real case this catches: BGF World Technology had a dead ticker that
    left it "priced" but frozen at an old value for months."""
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    master_path = tmp_path / "InstrumentMaster.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"

    _write_positions(positions_path, [
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "BGF WORLD TECHNOLOGY", "Date": "2018-09-07", "CumulativeQuantity": 68},
    ])
    _write_master(master_path, [
        {"Enabled": "YES", "NormalizedInstrument": "BGF WORLD TECHNOLOGY", "InstrumentType": "FUND", "Currency": "EUR", "PriceSymbol": "0P00000AWU"},
    ])
    old_date = (date.today() - timedelta(days=200)).isoformat()
    _write_prices(prices_path, [
        {"NormalizedInstrument": "BGF WORLD TECHNOLOGY", "Date": old_date, "Close": 82.3, "Currency": "EUR"},
    ])

    gaps = find_coverage_gaps(positions_path, master_path, prices_path, stale_days=7)

    assert gaps["unpriced"] == []
    assert len(gaps["stale"]) == 1
    assert gaps["stale"][0][0] == "BGF WORLD TECHNOLOGY"


def test_duplicate_normalized_instrument_rows_do_not_hide_the_real_classification(tmp_path):
    """
    Real bug caught testing this against real data: InstrumentMaster can
    have multiple rows sharing the same NormalizedInstrument (a real stock's
    row plus related SUBSCRIPTION_RIGHT/CORPORATE_ACTION rows with blank
    PriceSymbol). A naive set_index().loc[] lookup returns a DataFrame slice
    for the duplicated name, and normalise_text() defensively - and
    silently - treats that as blank, making a fully-classified instrument
    look unclassified.
    """
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    master_path = tmp_path / "InstrumentMaster.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"

    _write_positions(positions_path, [
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "FINNAIR", "Date": "2021-02-09", "CumulativeQuantity": 120},
    ])
    _write_master(master_path, [
        {"Enabled": "YES", "NormalizedInstrument": "FINNAIR", "InstrumentType": "SUBSCRIPTION_RIGHT", "Currency": "EUR", "PriceSymbol": ""},
        {"Enabled": "YES", "NormalizedInstrument": "FINNAIR", "InstrumentType": "STOCK", "Currency": "EUR", "PriceSymbol": "FIA1S.HE"},
        {"Enabled": "YES", "NormalizedInstrument": "FINNAIR", "InstrumentType": "TEMPORARY_SHARE", "Currency": "EUR", "PriceSymbol": ""},
    ])
    _write_prices(prices_path, [
        {"NormalizedInstrument": "FINNAIR", "Date": date.today().isoformat(), "Close": 4.4, "Currency": "EUR"},
    ])

    gaps = find_coverage_gaps(positions_path, master_path, prices_path)

    assert not has_gaps(gaps)


def test_cli_exits_nonzero_on_gaps_without_force(tmp_path, monkeypatch, capsys):
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    master_path = tmp_path / "InstrumentMaster.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"

    _write_positions(positions_path, [
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "NEW STOCK", "Date": "2026-01-01", "CumulativeQuantity": 10},
    ])
    _write_master(master_path, [])
    _write_prices(prices_path, [])

    argv = [
        "check_instrument_coverage.py",
        "--positions-workbook", str(positions_path),
        "--instrument-master", str(master_path),
        "--prices-workbook", str(prices_path),
    ]
    monkeypatch.setattr("sys.argv", argv)

    with pytest.raises(SystemExit) as exc_info:
        check_module.main()

    assert exc_info.value.code == 1
    assert "UNCLASSIFIED" in capsys.readouterr().out


def test_cli_force_reports_gaps_but_exits_zero(tmp_path, monkeypatch, capsys):
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    master_path = tmp_path / "InstrumentMaster.xlsx"
    prices_path = tmp_path / "InstrumentPrices.xlsx"

    _write_positions(positions_path, [
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "NEW STOCK", "Date": "2026-01-01", "CumulativeQuantity": 10},
    ])
    _write_master(master_path, [])
    _write_prices(prices_path, [])

    argv = [
        "check_instrument_coverage.py",
        "--positions-workbook", str(positions_path),
        "--instrument-master", str(master_path),
        "--prices-workbook", str(prices_path),
        "--force",
    ]
    monkeypatch.setattr("sys.argv", argv)

    check_module.main()  # must not raise/exit non-zero

    output = capsys.readouterr().out
    assert "UNCLASSIFIED" in output
    assert "proceeding anyway" in output.lower()
