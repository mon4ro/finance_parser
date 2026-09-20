from pathlib import Path

from openpyxl import Workbook, load_workbook

from finance_parser.budgeting.sync_category_taxonomy import (
    build_taxonomy_rows,
    extract_taxonomy,
    find_latest_master_budget_file,
    write_category_taxonomy_sheet,
)


def _write_master_budget(path: Path, taxonomy_rows: list[tuple[str, ...]]):
    """
    Build a synthetic Master Budget-shaped workbook: a KPI overview block
    (bare Supercategory names, no children) followed by the real detailed
    section (columns A/B/C = Supercategory/Category/Subcategory), followed
    by an unrelated balance-sheet section that must not be picked up.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Master Budget"

    ws.append(["MASTER BUDGET"])
    ws.append([])
    ws.append([])
    ws.append(["INCOME"])
    ws.append(["EXPENSES"])
    ws.append(["INVESTMENTS + MORTGAGE"])
    ws.append(["SAVINGS"])

    for row in taxonomy_rows:
        ws.append(list(row))

    ws.append(["TOTAL ASSETS"])
    ws.append([None, "Funds", None])
    ws.append([None, None, "Should not be picked up"])

    wb.save(path)


def test_extract_taxonomy_skips_overview_block_and_stops_at_balance_sheet(tmp_path):
    path = tmp_path / "master_budget.xlsx"
    _write_master_budget(
        path,
        [
            ("INCOME", None, None),
            (None, "Salary", None),
            (None, None, "Base pay"),
            (None, None, "Overtime"),
            ("EXPENSES", None, None),
            (None, "Groceries", None),
            (None, None, "Food items"),
        ],
    )

    triples = extract_taxonomy(path)

    assert triples == [
        ("INCOME", "Salary", "Base pay"),
        ("INCOME", "Salary", "Overtime"),
        ("EXPENSES", "Groceries", "Food items"),
    ]


def test_extract_taxonomy_handles_flat_supercategory_with_no_category_tier(tmp_path):
    path = tmp_path / "master_budget.xlsx"
    _write_master_budget(
        path,
        [
            ("SAVINGS", None, None),
            (None, None, "Fund A"),
            (None, None, "Fund B"),
        ],
    )

    triples = extract_taxonomy(path)

    assert triples == [("SAVINGS", "", "Fund A"), ("SAVINGS", "", "Fund B")]


def test_build_taxonomy_rows_flags_ambiguous_subcategory():
    triples = [
        ("INCOME", "Bonuses", "Extra pay"),
        ("EXPENSES", "Seed", "Tithe"),
        ("INCOME", "Misc", "Tithe"),
        ("EXPENSES", "Groceries", "Snacks"),
    ]

    rows = build_taxonomy_rows(triples)

    by_sub = {(r["Category"], r["Subcategory"]): r["Ambiguous"] for r in rows}
    assert by_sub[("Seed", "Tithe")] == "YES"
    assert by_sub[("Misc", "Tithe")] == "YES"
    assert by_sub[("Groceries", "Snacks")] == "NO"


def test_build_taxonomy_rows_deduplicates_identical_triples():
    triples = [
        ("EXPENSES", "Groceries", "Snacks"),
        ("EXPENSES", "Groceries", "Snacks"),
    ]

    rows = build_taxonomy_rows(triples)

    assert len(rows) == 1


def test_find_latest_master_budget_file_picks_highest_version_and_skips_test_copies(tmp_path):
    (tmp_path / "Master Budget 2025 v10.5.2 20250101T.xlsm").touch()
    (tmp_path / "Master Budget 2025 v11.2.1 20250101T.xlsm").touch()
    (tmp_path / "Master Budget 2025 v11.9.9 20250101T test.xlsm").touch()
    (tmp_path / "not a budget file.xlsm").touch()

    found = find_latest_master_budget_file(tmp_path)

    assert found.name == "Master Budget 2025 v11.2.1 20250101T.xlsm"


def test_find_latest_master_budget_file_raises_clearly_when_none_found(tmp_path):
    (tmp_path / "not a budget file.xlsm").touch()

    try:
        find_latest_master_budget_file(tmp_path)
        assert False, "expected FileNotFoundError"
    except FileNotFoundError as exc:
        assert "Master Budget" in str(exc)


def _write_rules_workbook(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "CategoryRules"
    ws.append(["RuleID", "Pattern"])
    ws.append(["CR0001", "SOME PATTERN"])

    other = wb.create_sheet("OwnershipRules")
    other.append(["RuleID"])
    other.append(["OR0001"])

    wb.save(path)


def test_write_category_taxonomy_sheet_leaves_other_sheets_untouched(tmp_path):
    path = tmp_path / "rules.xlsx"
    _write_rules_workbook(path)

    rows = [
        {"Supercategory": "EXPENSES", "Category": "Groceries", "Subcategory": "Snacks", "Ambiguous": "NO"},
    ]
    backup_path = write_category_taxonomy_sheet(path, rows)

    assert backup_path.exists()

    wb = load_workbook(path)
    assert set(wb.sheetnames) == {"CategoryRules", "OwnershipRules", "CategoryTaxonomy"}

    category_rules = wb["CategoryRules"]
    assert [c.value for c in category_rules[2]] == ["CR0001", "SOME PATTERN"]

    taxonomy = wb["CategoryTaxonomy"]
    assert [c.value for c in taxonomy[1]] == ["Supercategory", "Category", "Subcategory", "Ambiguous"]
    assert [c.value for c in taxonomy[2]] == ["EXPENSES", "Groceries", "Snacks", "NO"]


def test_write_category_taxonomy_sheet_replaces_existing_taxonomy_sheet(tmp_path):
    path = tmp_path / "rules.xlsx"
    _write_rules_workbook(path)

    write_category_taxonomy_sheet(
        path,
        [{"Supercategory": "EXPENSES", "Category": "Groceries", "Subcategory": "Old", "Ambiguous": "NO"}],
    )
    write_category_taxonomy_sheet(
        path,
        [{"Supercategory": "EXPENSES", "Category": "Groceries", "Subcategory": "New", "Ambiguous": "NO"}],
    )

    wb = load_workbook(path)
    taxonomy = wb["CategoryTaxonomy"]
    values = [row[2].value for row in taxonomy.iter_rows(min_row=2)]
    assert values == ["New"]
