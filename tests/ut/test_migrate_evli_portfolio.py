from openpyxl import Workbook

from finance_parser.common import normalise_text
from finance_parser.investments.investment_common import (
    INVESTMENT_RAW_COLUMNS,
    INVESTMENT_TRANSACTIONS_COLUMNS,
    INVESTMENT_IMPORT_LOG_COLUMNS,
    INVESTMENT_RAW_SHEET,
    INVESTMENT_TRANSACTIONS_SHEET,
    INVESTMENT_IMPORT_LOG_SHEET,
    read_sheet,
)
from finance_parser.investments.migrate_evli_portfolio import migrate_workbook


def _write_workbook(path, raw_rows, transaction_rows):
    wb = Workbook()
    ws = wb.active
    ws.title = INVESTMENT_RAW_SHEET
    ws.append(INVESTMENT_RAW_COLUMNS)
    for row in raw_rows:
        ws.append([row.get(h, "") for h in INVESTMENT_RAW_COLUMNS])

    ws2 = wb.create_sheet(INVESTMENT_TRANSACTIONS_SHEET)
    ws2.append(INVESTMENT_TRANSACTIONS_COLUMNS)
    for row in transaction_rows:
        ws2.append([row.get(h, "") for h in INVESTMENT_TRANSACTIONS_COLUMNS])

    ws3 = wb.create_sheet(INVESTMENT_IMPORT_LOG_SHEET)
    ws3.append(INVESTMENT_IMPORT_LOG_COLUMNS)

    extra = wb.create_sheet("EvliRawExport")
    extra.append(["Instrument", "Event type"])
    extra.append(["Plan Cycle 2022", "Vesting"])

    wb.save(path)


def test_migrates_old_evli_rows_moving_portfolio_into_plan_cycle(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_workbook(
        path,
        raw_rows=[
            {"InvestmentRawID": "EVL-1", "Broker": "EVLI", "Portfolio": "Plan Cycle 2022", "PlanCycle": ""},
            {"InvestmentRawID": "EVL-2", "Broker": "EVLI", "Portfolio": "SIS Dividend", "PlanCycle": ""},
            {"InvestmentRawID": "NN-1", "Broker": "NORDNET", "Portfolio": "1", "PlanCycle": ""},
        ],
        transaction_rows=[
            {"InvestmentTransactionID": "T1", "InvestmentRawID": "EVL-1", "Broker": "EVLI", "Portfolio": "Plan Cycle 2022", "PlanCycle": ""},
            {"InvestmentTransactionID": "T2", "InvestmentRawID": "EVL-2", "Broker": "EVLI", "Portfolio": "SIS Dividend", "PlanCycle": ""},
            {"InvestmentTransactionID": "T3", "InvestmentRawID": "NN-1", "Broker": "NORDNET", "Portfolio": "1", "PlanCycle": ""},
        ],
    )

    stats = migrate_workbook(path)

    assert stats["raw_rows_migrated"] == 2
    assert stats["transaction_rows_migrated"] == 2

    raw = read_sheet(path, INVESTMENT_RAW_SHEET, INVESTMENT_RAW_COLUMNS)
    by_id = raw.set_index("InvestmentRawID")
    assert by_id.loc["EVL-1", "Portfolio"] == "EVLI"
    assert by_id.loc["EVL-1", "PlanCycle"] == "Plan Cycle 2022"
    assert by_id.loc["EVL-2", "Portfolio"] == "EVLI"
    assert by_id.loc["EVL-2", "PlanCycle"] == "SIS Dividend"
    # Non-EVLI rows must be left completely untouched.
    assert by_id.loc["NN-1", "Portfolio"] == "1"
    assert normalise_text(by_id.loc["NN-1", "PlanCycle"]) == ""

    transactions = read_sheet(path, INVESTMENT_TRANSACTIONS_SHEET, INVESTMENT_TRANSACTIONS_COLUMNS)
    t_by_id = transactions.set_index("InvestmentTransactionID")
    assert t_by_id.loc["T1", "Portfolio"] == "EVLI"
    assert t_by_id.loc["T1", "PlanCycle"] == "Plan Cycle 2022"


def test_investment_raw_id_and_transaction_id_are_never_modified(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_workbook(
        path,
        raw_rows=[
            {"InvestmentRawID": "EVL-1", "Broker": "EVLI", "Portfolio": "Plan Cycle 2022", "PlanCycle": ""},
        ],
        transaction_rows=[
            {"InvestmentTransactionID": "T1", "InvestmentRawID": "EVL-1", "Broker": "EVLI", "Portfolio": "Plan Cycle 2022", "PlanCycle": ""},
        ],
    )

    migrate_workbook(path)

    raw = read_sheet(path, INVESTMENT_RAW_SHEET, INVESTMENT_RAW_COLUMNS)
    assert list(raw["InvestmentRawID"]) == ["EVL-1"]

    transactions = read_sheet(path, INVESTMENT_TRANSACTIONS_SHEET, INVESTMENT_TRANSACTIONS_COLUMNS)
    assert list(transactions["InvestmentTransactionID"]) == ["T1"]


def test_broker_raw_export_sheets_are_preserved(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_workbook(
        path,
        raw_rows=[{"InvestmentRawID": "EVL-1", "Broker": "EVLI", "Portfolio": "Plan Cycle 2022", "PlanCycle": ""}],
        transaction_rows=[{"InvestmentTransactionID": "T1", "InvestmentRawID": "EVL-1", "Broker": "EVLI", "Portfolio": "Plan Cycle 2022", "PlanCycle": ""}],
    )

    migrate_workbook(path)

    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    assert "EvliRawExport" in wb.sheetnames
    wb.close()


def test_running_twice_is_a_no_op_the_second_time(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_workbook(
        path,
        raw_rows=[{"InvestmentRawID": "EVL-1", "Broker": "EVLI", "Portfolio": "Plan Cycle 2022", "PlanCycle": ""}],
        transaction_rows=[{"InvestmentTransactionID": "T1", "InvestmentRawID": "EVL-1", "Broker": "EVLI", "Portfolio": "Plan Cycle 2022", "PlanCycle": ""}],
    )

    first = migrate_workbook(path)
    second = migrate_workbook(path)

    assert first["raw_rows_migrated"] == 1
    assert second["raw_rows_migrated"] == 0
    assert second["transaction_rows_migrated"] == 0


def test_dry_run_does_not_modify_the_workbook(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_workbook(
        path,
        raw_rows=[{"InvestmentRawID": "EVL-1", "Broker": "EVLI", "Portfolio": "Plan Cycle 2022", "PlanCycle": ""}],
        transaction_rows=[{"InvestmentTransactionID": "T1", "InvestmentRawID": "EVL-1", "Broker": "EVLI", "Portfolio": "Plan Cycle 2022", "PlanCycle": ""}],
    )

    stats = migrate_workbook(path, dry_run=True)
    assert stats["raw_rows_migrated"] == 1

    raw = read_sheet(path, INVESTMENT_RAW_SHEET, INVESTMENT_RAW_COLUMNS)
    assert raw.iloc[0]["Portfolio"] == "Plan Cycle 2022"
    assert normalise_text(raw.iloc[0]["PlanCycle"]) == ""
