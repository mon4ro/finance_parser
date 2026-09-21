from openpyxl import Workbook

from finance_parser.investments.investment_common import (
    INVESTMENT_RAW_COLUMNS,
    INVESTMENT_TRANSACTIONS_COLUMNS,
    INVESTMENT_IMPORT_LOG_COLUMNS,
    INVESTMENT_RAW_SHEET,
    INVESTMENT_TRANSACTIONS_SHEET,
    INVESTMENT_IMPORT_LOG_SHEET,
    read_sheet,
)
from finance_parser.investments.migrate_evli_sell_proceeds import migrate_workbook


def _write_workbook(path, raw_rows, transaction_rows):
    wb = Workbook()
    ws = wb.active
    ws.title = INVESTMENT_RAW_SHEET
    ws.append(INVESTMENT_RAW_COLUMNS)
    for row in raw_rows:
        ws.append([row.get(h, "") for h in INVESTMENT_RAW_COLUMNS])

    ws2 = wb.create_sheet(INVESTMENT_TRANSACTIONS_SHEET)
    ws2.append(INVESTMENT_TRANSACTIONS_COLUMNS)
    for row in transaction_rows:
        ws2.append([row.get(h, "") for h in INVESTMENT_TRANSACTIONS_COLUMNS])

    ws3 = wb.create_sheet(INVESTMENT_IMPORT_LOG_SHEET)
    ws3.append(INVESTMENT_IMPORT_LOG_COLUMNS)

    wb.save(path)


def test_estimates_proceeds_for_zero_cash_amount_sell_rows(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_workbook(
        path,
        raw_rows=[
            {"InvestmentRawID": "EVL-1", "Broker": "EVLI", "TransactionTypeRaw": "Sell of purchased share", "Quantity": -100, "UnitPrice": 5.0, "CashAmount": 0},
            {"InvestmentRawID": "EVL-2", "Broker": "EVLI", "TransactionTypeRaw": "Dividend", "Quantity": 0, "UnitPrice": 0, "CashAmount": 20},
            {"InvestmentRawID": "NN-1", "Broker": "NORDNET", "TransactionTypeRaw": "SELL", "Quantity": -100, "UnitPrice": 5.0, "CashAmount": 0},
        ],
        transaction_rows=[
            {"InvestmentTransactionID": "T1", "InvestmentRawID": "EVL-1", "Broker": "EVLI", "TransactionType": "SELL", "Quantity": -100, "UnitPrice": 5.0, "CashAmount": 0},
            {"InvestmentTransactionID": "T2", "InvestmentRawID": "EVL-2", "Broker": "EVLI", "TransactionType": "DIVIDEND", "Quantity": 0, "UnitPrice": 0, "CashAmount": 20},
            {"InvestmentTransactionID": "T3", "InvestmentRawID": "NN-1", "Broker": "NORDNET", "TransactionType": "SELL", "Quantity": -100, "UnitPrice": 5.0, "CashAmount": 0},
        ],
    )

    stats = migrate_workbook(path)

    assert stats["raw_rows_migrated"] == 1
    assert stats["transaction_rows_migrated"] == 1

    transactions = read_sheet(path, INVESTMENT_TRANSACTIONS_SHEET, INVESTMENT_TRANSACTIONS_COLUMNS)
    by_id = transactions.set_index("InvestmentTransactionID")
    assert by_id.loc["T1", "CashAmount"] == 500.0  # migrated
    assert by_id.loc["T2", "CashAmount"] == 20      # not a sell - untouched
    assert by_id.loc["T3", "CashAmount"] == 0        # not EVLI - untouched


def test_investment_ids_are_never_modified(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_workbook(
        path,
        raw_rows=[{"InvestmentRawID": "EVL-1", "Broker": "EVLI", "TransactionTypeRaw": "Sell of purchased share", "Quantity": -100, "UnitPrice": 5.0, "CashAmount": 0}],
        transaction_rows=[{"InvestmentTransactionID": "T1", "InvestmentRawID": "EVL-1", "Broker": "EVLI", "TransactionType": "SELL", "Quantity": -100, "UnitPrice": 5.0, "CashAmount": 0}],
    )

    migrate_workbook(path)

    raw = read_sheet(path, INVESTMENT_RAW_SHEET, INVESTMENT_RAW_COLUMNS)
    assert list(raw["InvestmentRawID"]) == ["EVL-1"]


def test_running_twice_is_a_no_op_the_second_time(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_workbook(
        path,
        raw_rows=[{"InvestmentRawID": "EVL-1", "Broker": "EVLI", "TransactionTypeRaw": "Sell of purchased share", "Quantity": -100, "UnitPrice": 5.0, "CashAmount": 0}],
        transaction_rows=[{"InvestmentTransactionID": "T1", "InvestmentRawID": "EVL-1", "Broker": "EVLI", "TransactionType": "SELL", "Quantity": -100, "UnitPrice": 5.0, "CashAmount": 0}],
    )

    first = migrate_workbook(path)
    second = migrate_workbook(path)

    assert first["raw_rows_migrated"] == 1
    assert second["raw_rows_migrated"] == 0


def test_dry_run_does_not_modify_the_workbook(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_workbook(
        path,
        raw_rows=[{"InvestmentRawID": "EVL-1", "Broker": "EVLI", "TransactionTypeRaw": "Sell of purchased share", "Quantity": -100, "UnitPrice": 5.0, "CashAmount": 0}],
        transaction_rows=[{"InvestmentTransactionID": "T1", "InvestmentRawID": "EVL-1", "Broker": "EVLI", "TransactionType": "SELL", "Quantity": -100, "UnitPrice": 5.0, "CashAmount": 0}],
    )

    stats = migrate_workbook(path, dry_run=True)
    assert stats["raw_rows_migrated"] == 1

    raw = read_sheet(path, INVESTMENT_RAW_SHEET, INVESTMENT_RAW_COLUMNS)
    assert raw.iloc[0]["CashAmount"] == 0
