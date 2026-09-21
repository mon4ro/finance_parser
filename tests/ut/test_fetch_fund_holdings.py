import urllib.error
from datetime import date

import pytest
from openpyxl import Workbook

from finance_parser.investments.fetch_fund_holdings import (
    BACKFILL_MONTHS,
    HOLDINGS_COLUMNS,
    active_fund_plan,
    build_backfill_rows,
    build_snapshot_rows,
    build_yahoo_opener,
    fetch_all,
    fetch_top_holdings,
    get_yahoo_crumb,
    load_fund_price_config,
    month_str_n_back,
    write_holdings_workbook,
)


INSTRUMENT_MASTER_HEADERS = [
    "Enabled", "MatchPriority", "Broker", "ISIN", "RawInstrumentName",
    "NormalizedInstrument", "InstrumentType", "AssetClass", "Currency",
    "Ticker", "Exchange", "PriceSource", "PriceSymbol", "Notes",
]

POSITIONS_HEADERS = [
    "Broker", "Portfolio", "NormalizedInstrument", "Date", "TransactionType",
    "QuantityDelta", "CumulativeQuantity", "CashDelta", "CumulativeNetInvested",
]


def _write_instrument_master(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "InstrumentMaster"
    ws.append(INSTRUMENT_MASTER_HEADERS)
    for row in rows:
        ws.append([row.get(h, "") for h in INSTRUMENT_MASTER_HEADERS])
    wb.save(path)


def _write_positions(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "PortfolioPositions"
    ws.append(POSITIONS_HEADERS)
    for row in rows:
        ws.append([row.get(h, "") for h in POSITIONS_HEADERS])
    wb.save(path)


def test_load_fund_price_config_filters_to_funds_only(tmp_path):
    path = tmp_path / "InstrumentMaster.xlsx"
    _write_instrument_master(path, [
        {"Enabled": "YES", "InstrumentType": "FUND", "NormalizedInstrument": "SAMPLE FUND A", "PriceSymbol": "SFA.HE"},
        {"Enabled": "YES", "InstrumentType": "STOCK", "NormalizedInstrument": "SAMPLE STOCK A", "PriceSymbol": "SSA.HE"},
        {"Enabled": "NO", "InstrumentType": "FUND", "NormalizedInstrument": "DISABLED FUND", "PriceSymbol": "XYZ"},
        {"Enabled": "YES", "InstrumentType": "FUND", "NormalizedInstrument": "NO SYMBOL FUND", "PriceSymbol": ""},
    ])

    df = load_fund_price_config(path)

    assert list(df["NormalizedInstrument"]) == ["SAMPLE FUND A"]


def test_active_fund_plan_excludes_inactive_and_non_fund_positions(tmp_path):
    master_path = tmp_path / "InstrumentMaster.xlsx"
    positions_path = tmp_path / "PortfolioPositions.xlsx"

    _write_instrument_master(master_path, [
        {"Enabled": "YES", "InstrumentType": "FUND", "NormalizedInstrument": "HELD FUND", "PriceSymbol": "HF.HE"},
        {"Enabled": "YES", "InstrumentType": "FUND", "NormalizedInstrument": "SOLD FUND", "PriceSymbol": "SF.HE"},
        {"Enabled": "YES", "InstrumentType": "STOCK", "NormalizedInstrument": "HELD STOCK", "PriceSymbol": "HS.HE"},
    ])
    _write_positions(positions_path, [
        {"Broker": "TEST", "Portfolio": "TEST", "NormalizedInstrument": "HELD FUND", "Date": "2026-01-01", "CumulativeQuantity": 10},
        {"Broker": "TEST", "Portfolio": "TEST", "NormalizedInstrument": "SOLD FUND", "Date": "2026-01-01", "CumulativeQuantity": 0},
        {"Broker": "TEST", "Portfolio": "TEST", "NormalizedInstrument": "HELD STOCK", "Date": "2026-01-01", "CumulativeQuantity": 5},
    ])

    plan = active_fund_plan(master_path, positions_path)

    assert list(plan["NormalizedInstrument"]) == ["HELD FUND"]


def test_month_str_n_back_wraps_year_boundary():
    assert month_str_n_back(date(2026, 2, 1), 0) == "2026-02"
    assert month_str_n_back(date(2026, 2, 1), 1) == "2026-01"
    assert month_str_n_back(date(2026, 2, 1), 2) == "2025-12"
    assert month_str_n_back(date(2026, 2, 1), 14) == "2024-12"


def test_build_snapshot_rows_always_sums_to_100_with_holdings():
    holdings = [
        {"symbol": "AAA", "name": "Company A", "percent": 30.0},
        {"symbol": "BBB", "name": "Company B", "percent": 20.0},
    ]

    rows = build_snapshot_rows("FUND.HE", "SAMPLE FUND", "2026-01", date(2026, 1, 15), holdings, source="LIVE_FETCH")

    assert len(rows) == 3
    total = sum(float(r["HoldingPercent"]) for r in rows)
    assert total == pytest.approx(100.0)
    unknown_row = rows[-1]
    assert unknown_row["HoldingSymbol"] == "UNKNOWN"
    assert unknown_row["HoldingName"] == "Unmapped remainder (beyond top 10)"
    assert unknown_row["HoldingPercent"] == pytest.approx(50.0)


def test_build_snapshot_rows_full_unknown_when_no_holdings():
    rows = build_snapshot_rows("FUND.HE", "SAMPLE FUND", "2026-01", date(2026, 1, 15), [], source="LIVE_FETCH")

    assert len(rows) == 1
    assert rows[0]["HoldingSymbol"] == "UNKNOWN"
    assert rows[0]["HoldingName"] == "No holdings data available"
    assert rows[0]["HoldingPercent"] == 100.0


def test_build_snapshot_rows_no_residual_when_holdings_already_sum_to_100():
    holdings = [{"symbol": "AAA", "name": "Company A", "percent": 100.0}]

    rows = build_snapshot_rows("FUND.HE", "SAMPLE FUND", "2026-01", date(2026, 1, 15), holdings, source="LIVE_FETCH")

    assert len(rows) == 1
    assert rows[0]["HoldingSymbol"] == "AAA"


def test_build_backfill_rows_covers_full_range_tagged_backfilled():
    holdings = [{"symbol": "AAA", "name": "Company A", "percent": 60.0}]
    today = date(2026, 3, 15)

    rows = build_backfill_rows("FUND.HE", "SAMPLE FUND", holdings, today, months_back=BACKFILL_MONTHS)

    months = {r["SnapshotMonth"] for r in rows}
    assert len(months) == BACKFILL_MONTHS
    assert "2026-02" in months  # one month back
    assert "2026-03" not in months  # current month is the live fetch's job, not backfill's
    assert all(r["SnapshotSource"] == "BACKFILLED_ASSUMED_CONSTANT" for r in rows)


def test_fetch_all_skips_funds_already_snapshotted_this_month(tmp_path, monkeypatch):
    master_path = tmp_path / "InstrumentMaster.xlsx"
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    holdings_path = tmp_path / "FundHoldingsSnapshot.xlsx"

    _write_instrument_master(master_path, [
        {"Enabled": "YES", "InstrumentType": "FUND", "NormalizedInstrument": "SAMPLE FUND", "PriceSymbol": "SF.HE"},
    ])
    _write_positions(positions_path, [
        {"Broker": "TEST", "Portfolio": "TEST", "NormalizedInstrument": "SAMPLE FUND", "Date": "2026-01-01", "CumulativeQuantity": 10},
    ])

    wb = Workbook()
    ws = wb.active
    ws.title = "FundHoldingsSnapshot"
    ws.append(HOLDINGS_COLUMNS)
    current_month = date.today().strftime("%Y-%m")
    ws.append([current_month, date.today().isoformat(), "SF.HE", "SAMPLE FUND", 11, "UNKNOWN", "No holdings data available", 100.0, "LIVE_FETCH"])
    wb.save(holdings_path)

    def fail_if_called(*args, **kwargs):
        raise AssertionError("should not fetch a fund already snapshotted this month")

    monkeypatch.setattr("finance_parser.investments.fetch_fund_holdings.get_yahoo_crumb", fail_if_called)

    combined, stats = fetch_all(master_path, positions_path, holdings_path, verbose=False)

    assert stats["funds_up_to_date"] == 1
    assert stats["funds_fetched"] == 0
    assert len(combined) == 1


def test_fetch_all_backfills_only_on_first_ever_fetch(tmp_path, monkeypatch):
    master_path = tmp_path / "InstrumentMaster.xlsx"
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    holdings_path = tmp_path / "FundHoldingsSnapshot.xlsx"

    _write_instrument_master(master_path, [
        {"Enabled": "YES", "InstrumentType": "FUND", "NormalizedInstrument": "SAMPLE FUND", "PriceSymbol": "SF.HE"},
    ])
    _write_positions(positions_path, [
        {"Broker": "TEST", "Portfolio": "TEST", "NormalizedInstrument": "SAMPLE FUND", "Date": "2026-01-01", "CumulativeQuantity": 10},
    ])

    monkeypatch.setattr("finance_parser.investments.fetch_fund_holdings.get_yahoo_crumb", lambda opener: "fake-crumb")
    monkeypatch.setattr(
        "finance_parser.investments.fetch_fund_holdings.fetch_top_holdings",
        lambda symbol, crumb, opener: [{"symbol": "AAA", "name": "Company A", "percent": 40.0}],
    )

    combined, stats = fetch_all(master_path, positions_path, holdings_path, verbose=False)

    assert stats["funds_backfilled"] == ["SAMPLE FUND"]
    months = set(combined["SnapshotMonth"])
    assert len(months) == BACKFILL_MONTHS + 1  # backfill + the current live month

    # A second run in the same month should not backfill again - requires the
    # first run's result to actually be on disk, same as a real pipeline run.
    write_holdings_workbook(holdings_path, combined)
    combined2, stats2 = fetch_all(master_path, positions_path, holdings_path, verbose=False)
    assert stats2["funds_up_to_date"] == 1
    assert stats2["funds_backfilled"] == []


def test_fetch_all_records_unknown_row_on_fetch_failure(tmp_path, monkeypatch):
    master_path = tmp_path / "InstrumentMaster.xlsx"
    positions_path = tmp_path / "PortfolioPositions.xlsx"
    holdings_path = tmp_path / "FundHoldingsSnapshot.xlsx"

    _write_instrument_master(master_path, [
        {"Enabled": "YES", "InstrumentType": "FUND", "NormalizedInstrument": "SAMPLE FUND", "PriceSymbol": "SF.HE"},
    ])
    _write_positions(positions_path, [
        {"Broker": "TEST", "Portfolio": "TEST", "NormalizedInstrument": "SAMPLE FUND", "Date": "2026-01-01", "CumulativeQuantity": 10},
    ])

    def fake_crumb(opener):
        raise urllib.error.URLError("crumb endpoint unavailable")

    monkeypatch.setattr("finance_parser.investments.fetch_fund_holdings.get_yahoo_crumb", fake_crumb)

    combined, stats = fetch_all(master_path, positions_path, holdings_path, verbose=False)

    assert len(stats["funds_failed"]) == 1
    live_row = combined[combined["SnapshotSource"] == "LIVE_FETCH"].iloc[0]
    assert live_row["HoldingSymbol"] == "UNKNOWN"
    assert live_row["HoldingPercent"] == 100.0


@pytest.mark.parametrize("symbol", ["SPY"])
def test_fetch_top_holdings_real_network_call(symbol):
    """
    Live integration check against the real (unofficial) Yahoo Finance
    quoteSummary endpoint, using a well-known public ETF (not a real
    household holding) so no real portfolio data is involved. Skips
    gracefully if the network/crumb workaround is unavailable, matching this
    project's existing pattern for live price-API tests.
    """
    try:
        opener = build_yahoo_opener()
        crumb = get_yahoo_crumb(opener)
        holdings = fetch_top_holdings(symbol, crumb, opener)
    except Exception as exc:  # network unavailable, crumb workaround broken, etc.
        pytest.skip(f"Live fund holdings API unreachable, skipping: {exc}")

    assert len(holdings) > 0
    assert len(holdings) <= 10
    for holding in holdings:
        assert holding["percent"] >= 0
