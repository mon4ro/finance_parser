import os
import re
import time
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

from finance_parser.utilities.fresh_workbook_writer import (
    append_changelog_row,
    apply_review_status_column,
    find_timestamped_backups,
    prune_backups,
    read_workbook_values_only,
    records_to_sheet_values,
    replace_with_fresh_workbook,
    sheet_values_to_records,
    validate_xlsx_basic,
    write_fresh_workbook,
)


UNIFIED_HEADERS = [
    "UnifiedID", "RawID", "SourceAccount", "SourceBank", "TransactionType",
    "Date", "Month", "Year", "Amount", "RawReceiver", "NormalizedReceiver",
    "Description", "Message", "Include", "Owner", "Supercategory", "Category",
    "Subcategory", "ReviewStatus", "Review/Notes", "Currency", "Rate",
    "ExportDate", "ImportedAt", "SourceFile",
]


def _blank_unified_row(**overrides) -> list:
    row = {h: "" for h in UNIFIED_HEADERS}
    row.update(overrides)
    return [row[h] for h in UNIFIED_HEADERS]


def test_fresh_writer_reads_values_and_writes_fresh_package(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "fresh.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "UnifiedTransactions"
    ws.append(["UnifiedID", "Category"])
    ws.append(["U1", "Old"])
    wb.create_sheet("ChangeLog").append(["Timestamp", "Script", "Action"])
    wb.save(source)
    wb.close()

    sheets = read_workbook_values_only(source)
    headers, records = sheet_values_to_records(sheets["UnifiedTransactions"])
    records[0]["Category"] = "New"
    sheets["UnifiedTransactions"] = records_to_sheet_values(headers, records)
    append_changelog_row(
        sheets,
        script="test",
        action="Fresh write",
        sheet="UnifiedTransactions",
        rows_updated=1,
        status="Completed",
        details="unit test",
    )
    stats = write_fresh_workbook(output, sheets)

    assert stats["sheets_written"] == 2
    assert validate_xlsx_basic(output, expected_sheets={"UnifiedTransactions", "ChangeLog"}) == []

    wb2 = load_workbook(output, read_only=True)
    assert wb2["UnifiedTransactions"]["B2"].value == "New"
    assert wb2["ChangeLog"].max_row == 2
    wb2.close()


def _inject_cached_formula_value(path: Path, sheet_xml: str, cell_ref: str, cached_value: str) -> None:
    """
    Post-process a saved .xlsx to give one formula cell a real cached <v> -
    what a real Excel-authored file actually looks like, and what openpyxl
    itself never produces (it writes formula cells with no cached result).
    """
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        data = {name: zf.read(name) for name in names}

    xml = data[sheet_xml].decode("utf-8")
    pattern = re.compile(rf'(<c r="{cell_ref}"[^>]*><f>[^<]*</f>)<v></v>(</c>)')
    xml, count = pattern.subn(rf'\1<v>{cached_value}</v>\2', xml)
    assert count == 1, f"expected exactly one match for {cell_ref} in {sheet_xml}"
    data[sheet_xml] = xml.encode("utf-8")

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name in names:
            zf.writestr(name, data[name])


def test_read_workbook_values_only_reads_cached_formula_result_not_formula_text(tmp_path):
    """
    Real bug found and fixed: a foreign sheet pasted into the workbook by
    hand (not written by this pipeline) can carry real Excel formulas -
    reading those with data_only=False returned the formula TEXT as a
    string, and write_fresh_workbook()'s plain cell assignment then let
    openpyxl reinterpret the leading "=" as a live formula in the freshly
    built workbook, one with no Table/defined-name behind whatever
    structured reference the original relied on. Excel flagged and stripped
    it as corrupt on next open. Reading the cached result instead (what
    data_only=True gives) and writing that back as a plain value sidesteps
    the whole failure class - never reconstructs a formula from someone
    else's formula text.
    """
    source = tmp_path / "source.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "PastedSheet"
    ws.append(["Label", "Computed"])
    ws["B2"] = "=1+1"
    wb.save(source)
    wb.close()

    _inject_cached_formula_value(source, "xl/worksheets/sheet1.xml", "B2", "2")

    sheets = read_workbook_values_only(source)

    assert sheets["PastedSheet"][1][1] == 2

    output = tmp_path / "fresh.xlsx"
    write_fresh_workbook(output, sheets, excel_tables=False)

    wb2 = load_workbook(output, data_only=False, read_only=True)
    value = wb2["PastedSheet"]["B2"].value
    wb2.close()
    # A live formula would start with "=" - must be the plain cached number.
    assert value == 2


def test_apply_review_status_column_writes_expected_formula_and_colours():
    wb = Workbook()
    ws = wb.active
    ws.title = "UnifiedTransactions"
    ws.append(UNIFIED_HEADERS)
    ws.append(_blank_unified_row(Include="YES"))

    applied = apply_review_status_column(ws)

    assert applied is True
    # Include=N, Owner=O, Supercategory=P, Category=Q, Subcategory=R,
    # NormalizedReceiver=K, Review/Notes=T, ReviewStatus=S - matches the
    # column layout of the real UnifiedTransactions schema.
    assert ws["S2"].value == (
        '=IF(N2="NO","SKIP",'
        'IF(AND(N2="YES",O2<>"",P2<>"",Q2<>"",R2<>"",K2<>"",T2<>""),'
        '"READY","TODO"))'
    )

    rules = list(ws.conditional_formatting["S2:S2"])
    assert len(rules) == 3
    operands = sorted(r.formula[0] for r in rules)
    assert operands == ['"READY"', '"SKIP"', '"TODO"']


def test_apply_review_status_column_is_noop_without_expected_headers():
    wb = Workbook()
    ws = wb.active
    ws.title = "UnifiedTransactions"
    ws.append(["UnifiedID", "Category"])
    ws.append(["U1", "Restaurants"])

    assert apply_review_status_column(ws) is False
    assert ws["C2"].value is None


def test_write_fresh_workbook_produces_live_formula_not_flattened_value():
    import tempfile
    output = Path(tempfile.mkdtemp()) / "fresh.xlsx"

    sheets = {
        "UnifiedTransactions": [
            UNIFIED_HEADERS,
            _blank_unified_row(Include="NO"),
            _blank_unified_row(
                Include="YES", Owner="PERSONAL", Supercategory="EXPENSES",
                Category="Restaurants", Subcategory="Fast food",
                NormalizedReceiver="HESBURGER", **{"Review/Notes": "OK"},
            ),
        ]
    }
    write_fresh_workbook(output, sheets)

    # Reopen with data_only=False: a flattened value would just be the plain
    # string "SKIP"/"READY". A live formula must start with "=".
    wb = load_workbook(output, data_only=False, read_only=True)
    ws = wb["UnifiedTransactions"]
    s2 = ws["S2"].value
    s3 = ws["S3"].value
    wb.close()

    assert isinstance(s2, str) and s2.startswith("=IF(")
    assert isinstance(s3, str) and s3.startswith("=IF(")


def test_write_clean_output_workbook_also_gets_live_review_status_formula(tmp_path):
    from finance_parser.common import UNIFIED_COLUMNS, RAW_COLUMNS, IMPORT_LOG_COLUMNS, write_clean_output_workbook
    import pandas as pd

    output = tmp_path / "ParsedTransactions.xlsx"

    unified_row = {h: "" for h in UNIFIED_COLUMNS}
    unified_row.update({
        "UnifiedID": "U-CASH-1", "RawID": "CASH-1", "SourceAccount": "CASH",
        "SourceBank": "CASH", "TransactionType": "Cash", "Include": "YES",
        "Owner": "SHARED", "Supercategory": "EXPENSES", "Category": "Groceries",
        "Subcategory": "Food", "NormalizedReceiver": "TEST", "Review/Notes": "OK",
    })
    unified_df = pd.DataFrame([unified_row], columns=UNIFIED_COLUMNS)
    raw_df = pd.DataFrame(columns=RAW_COLUMNS)
    log_df = pd.DataFrame(columns=IMPORT_LOG_COLUMNS)

    write_clean_output_workbook(output, raw_df, unified_df, log_df)

    wb = load_workbook(output, data_only=False, read_only=True)
    ws = wb["UnifiedTransactions"]
    value = ws["S2"].value
    wb.close()

    assert isinstance(value, str) and value.startswith("=IF(")


def test_replace_with_fresh_workbook_backs_up_and_replaces(tmp_path):
    path = tmp_path / "book.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "UnifiedTransactions"
    ws.append(["UnifiedID", "Category"])
    ws.append(["U1", "Old"])
    wb.save(path)
    wb.close()

    sheets = {
        "UnifiedTransactions": [
            ["UnifiedID", "Category"],
            ["U1", "New"],
        ]
    }

    backup, new_path, stats = replace_with_fresh_workbook(
        path,
        sheets,
        backup_label="unit",
        basic_formatting=True,
        excel_tables=True,
    )

    assert path.exists()
    assert backup.exists()
    assert not new_path.exists()
    assert stats["tables_added"] == 1

    wb2 = load_workbook(path, read_only=True)
    assert wb2["UnifiedTransactions"]["B2"].value == "New"
    wb2.close()


def _touch_backup(path: Path, name: str, *, age_seconds: int) -> Path:
    backup = path.parent / name
    backup.write_text("stub")
    now = time.time()
    os.utime(backup, (now - age_seconds, now - age_seconds))
    return backup


def test_find_timestamped_backups_ignores_non_timestamped_pattern(tmp_path):
    path = tmp_path / "book.xlsx"
    path.write_text("stub")

    timestamped = _touch_backup(path, "book_backup_unit_20260101_120000.xlsx", age_seconds=0)
    # The separate, non-accumulating pattern used by write_investment_output_workbook()
    # / common.py's budgeting write path - always overwritten in place, never a
    # pruning candidate.
    _touch_backup(path, "book_backup_before_last_run.xlsx", age_seconds=0)

    found = find_timestamped_backups(path)

    assert found == [timestamped]


def test_prune_backups_keeps_only_the_newest_n(tmp_path):
    path = tmp_path / "book.xlsx"
    path.write_text("stub")

    # Oldest to newest.
    old_to_new = [
        _touch_backup(path, f"book_backup_unit_2026010{i}_120000.xlsx", age_seconds=(10 - i))
        for i in range(1, 8)
    ]

    removed = prune_backups(path, keep=5)

    assert set(removed) == set(old_to_new[:2])
    remaining = {p.name for p in find_timestamped_backups(path)}
    assert remaining == {p.name for p in old_to_new[2:]}


def test_replace_with_fresh_workbook_prunes_old_backups_automatically(tmp_path):
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active.title = "UnifiedTransactions"
    wb.active.append(["UnifiedID"])
    wb.save(path)
    wb.close()

    sheets = {"UnifiedTransactions": [["UnifiedID"], ["U1"]]}

    for i in range(7):
        replace_with_fresh_workbook(path, sheets, backup_label=f"run{i}", backup_retention=5)

    remaining = find_timestamped_backups(path)
    assert len(remaining) == 5


def test_replace_with_fresh_workbook_backup_retention_zero_disables_pruning(tmp_path):
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active.title = "UnifiedTransactions"
    wb.active.append(["UnifiedID"])
    wb.save(path)
    wb.close()

    sheets = {"UnifiedTransactions": [["UnifiedID"], ["U1"]]}

    for i in range(7):
        replace_with_fresh_workbook(path, sheets, backup_label=f"run{i}", backup_retention=0)

    remaining = find_timestamped_backups(path)
    assert len(remaining) == 7
