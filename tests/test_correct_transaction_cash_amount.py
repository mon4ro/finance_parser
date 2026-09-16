import pytest
from openpyxl import Workbook

from finance_parser.investments.investment_common import (
    INVESTMENT_RAW_COLUMNS,
    INVESTMENT_TRANSACTIONS_COLUMNS,
    INVESTMENT_IMPORT_LOG_COLUMNS,
    INVESTMENT_RAW_SHEET,
    INVESTMENT_TRANSACTIONS_SHEET,
    INVESTMENT_IMPORT_LOG_SHEET,
    read_sheet,
)
from finance_parser.investments.correct_transaction_cash_amount import correct_cash_amount


def _write_workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = INVESTMENT_RAW_SHEET
    ws.append(INVESTMENT_RAW_COLUMNS)
    ws.append([("RAW-1" if c == "InvestmentRawID" else "") for c in INVESTMENT_RAW_COLUMNS])

    ws2 = wb.create_sheet(INVESTMENT_TRANSACTIONS_SHEET)
    ws2.append(INVESTMENT_TRANSACTIONS_COLUMNS)
    row1 = {c: "" for c in INVESTMENT_TRANSACTIONS_COLUMNS}
    row1.update({"InvestmentTransactionID": "T-1", "InvestmentRawID": "RAW-1", "Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "OP-SUOMI A", "TransactionType": "VAIHTO - JÄTTÖ", "CashAmount": 0})
    row2 = {c: "" for c in INVESTMENT_TRANSACTIONS_COLUMNS}
    row2.update({"InvestmentTransactionID": "T-2", "InvestmentRawID": "RAW-2", "Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "SAMPO A", "TransactionType": "BUY", "CashAmount": -100})
    ws2.append([row1.get(c, "") for c in INVESTMENT_TRANSACTIONS_COLUMNS])
    ws2.append([row2.get(c, "") for c in INVESTMENT_TRANSACTIONS_COLUMNS])

    ws3 = wb.create_sheet(INVESTMENT_IMPORT_LOG_SHEET)
    ws3.append(INVESTMENT_IMPORT_LOG_COLUMNS)

    wb.save(path)


def test_corrects_only_the_matched_row(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_workbook(path)

    stats = correct_cash_amount(path, "T-1", -3141.19)

    assert stats["old_cash_amount"] == 0
    assert stats["new_cash_amount"] == -3141.19
    assert stats["rows_matched"] == 1

    transactions = read_sheet(path, INVESTMENT_TRANSACTIONS_SHEET, INVESTMENT_TRANSACTIONS_COLUMNS)
    by_id = transactions.set_index("InvestmentTransactionID")
    assert by_id.loc["T-1", "CashAmount"] == -3141.19
    # The unrelated row must be completely untouched.
    assert by_id.loc["T-2", "CashAmount"] == -100


def test_raw_sheet_is_never_modified(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_workbook(path)

    correct_cash_amount(path, "T-1", -3141.19)

    raw = read_sheet(path, INVESTMENT_RAW_SHEET, INVESTMENT_RAW_COLUMNS)
    assert raw.iloc[0]["InvestmentRawID"] == "RAW-1"


def test_raises_on_unknown_transaction_id(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_workbook(path)

    with pytest.raises(ValueError):
        correct_cash_amount(path, "DOES-NOT-EXIST", -100.0)


def test_dry_run_does_not_modify_the_workbook(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_workbook(path)

    stats = correct_cash_amount(path, "T-1", -3141.19, dry_run=True)
    assert stats["new_cash_amount"] == -3141.19

    transactions = read_sheet(path, INVESTMENT_TRANSACTIONS_SHEET, INVESTMENT_TRANSACTIONS_COLUMNS)
    assert transactions.set_index("InvestmentTransactionID").loc["T-1", "CashAmount"] == 0
