from openpyxl import Workbook

from finance_parser.investments.investment_common import (
    INVESTMENT_RAW_COLUMNS,
    INVESTMENT_TRANSACTIONS_COLUMNS,
    INVESTMENT_IMPORT_LOG_COLUMNS,
    INVESTMENT_RAW_SHEET,
    INVESTMENT_TRANSACTIONS_SHEET,
    INVESTMENT_IMPORT_LOG_SHEET,
    read_sheet,
)
from finance_parser.investments.migrate_seligson_cash_amount_sign import migrate_workbook


def _write_workbook(path, raw_rows, transaction_rows):
    wb = Workbook()
    ws = wb.active
    ws.title = INVESTMENT_RAW_SHEET
    ws.append(INVESTMENT_RAW_COLUMNS)
    for row in raw_rows:
        ws.append([row.get(h, "") for h in INVESTMENT_RAW_COLUMNS])

    ws2 = wb.create_sheet(INVESTMENT_TRANSACTIONS_SHEET)
    ws2.append(INVESTMENT_TRANSACTIONS_COLUMNS)
    for row in transaction_rows:
        ws2.append([row.get(h, "") for h in INVESTMENT_TRANSACTIONS_COLUMNS])

    ws3 = wb.create_sheet(INVESTMENT_IMPORT_LOG_SHEET)
    ws3.append(INVESTMENT_IMPORT_LOG_COLUMNS)

    wb.save(path)


def test_flips_positive_buy_rows_and_leaves_others_alone(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_workbook(
        path,
        raw_rows=[
            {"InvestmentRawID": "SEL-1", "Broker": "SELIGSON", "TransactionTypeRaw": "MERKINTÄ", "CashAmount": 25},
            {"InvestmentRawID": "SEL-2", "Broker": "SELIGSON", "TransactionTypeRaw": "LUNASTUS", "CashAmount": 25},
            {"InvestmentRawID": "NN-1", "Broker": "NORDNET", "TransactionTypeRaw": "BUY", "CashAmount": 25},
        ],
        transaction_rows=[
            {"InvestmentTransactionID": "T1", "InvestmentRawID": "SEL-1", "Broker": "SELIGSON", "TransactionType": "BUY", "CashAmount": 25},
            {"InvestmentTransactionID": "T2", "InvestmentRawID": "SEL-2", "Broker": "SELIGSON", "TransactionType": "SELL", "CashAmount": 25},
            {"InvestmentTransactionID": "T3", "InvestmentRawID": "NN-1", "Broker": "NORDNET", "TransactionType": "BUY", "CashAmount": 25},
        ],
    )

    stats = migrate_workbook(path)

    assert stats["raw_rows_migrated"] == 1
    assert stats["transaction_rows_migrated"] == 1

    raw = read_sheet(path, INVESTMENT_RAW_SHEET, INVESTMENT_RAW_COLUMNS)
    by_id = raw.set_index("InvestmentRawID")
    assert by_id.loc["SEL-1", "CashAmount"] == -25  # MERKINTÄ, was positive -> flipped
    assert by_id.loc["SEL-2", "CashAmount"] == 25   # LUNASTUS, already correctly positive -> untouched
    assert by_id.loc["NN-1", "CashAmount"] == 25    # Non-Seligson -> untouched

    transactions = read_sheet(path, INVESTMENT_TRANSACTIONS_SHEET, INVESTMENT_TRANSACTIONS_COLUMNS)
    t_by_id = transactions.set_index("InvestmentTransactionID")
    assert t_by_id.loc["T1", "CashAmount"] == -25
    assert t_by_id.loc["T2", "CashAmount"] == 25
    assert t_by_id.loc["T3", "CashAmount"] == 25


def test_investment_ids_are_never_modified(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_workbook(
        path,
        raw_rows=[{"InvestmentRawID": "SEL-1", "Broker": "SELIGSON", "TransactionTypeRaw": "MERKINTÄ", "CashAmount": 25}],
        transaction_rows=[{"InvestmentTransactionID": "T1", "InvestmentRawID": "SEL-1", "Broker": "SELIGSON", "TransactionType": "BUY", "CashAmount": 25}],
    )

    migrate_workbook(path)

    raw = read_sheet(path, INVESTMENT_RAW_SHEET, INVESTMENT_RAW_COLUMNS)
    assert list(raw["InvestmentRawID"]) == ["SEL-1"]


def test_running_twice_is_a_no_op_the_second_time(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_workbook(
        path,
        raw_rows=[{"InvestmentRawID": "SEL-1", "Broker": "SELIGSON", "TransactionTypeRaw": "MERKINTÄ", "CashAmount": 25}],
        transaction_rows=[{"InvestmentTransactionID": "T1", "InvestmentRawID": "SEL-1", "Broker": "SELIGSON", "TransactionType": "BUY", "CashAmount": 25}],
    )

    first = migrate_workbook(path)
    second = migrate_workbook(path)

    assert first["raw_rows_migrated"] == 1
    assert second["raw_rows_migrated"] == 0


def test_dry_run_does_not_modify_the_workbook(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_workbook(
        path,
        raw_rows=[{"InvestmentRawID": "SEL-1", "Broker": "SELIGSON", "TransactionTypeRaw": "MERKINTÄ", "CashAmount": 25}],
        transaction_rows=[{"InvestmentTransactionID": "T1", "InvestmentRawID": "SEL-1", "Broker": "SELIGSON", "TransactionType": "BUY", "CashAmount": 25}],
    )

    stats = migrate_workbook(path, dry_run=True)
    assert stats["raw_rows_migrated"] == 1

    raw = read_sheet(path, INVESTMENT_RAW_SHEET, INVESTMENT_RAW_COLUMNS)
    assert raw.iloc[0]["CashAmount"] == 25
