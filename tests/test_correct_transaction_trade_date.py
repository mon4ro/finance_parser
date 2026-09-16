import pytest
from openpyxl import Workbook

from finance_parser.investments.investment_common import (
    INVESTMENT_RAW_COLUMNS,
    INVESTMENT_TRANSACTIONS_COLUMNS,
    INVESTMENT_IMPORT_LOG_COLUMNS,
    INVESTMENT_RAW_SHEET,
    INVESTMENT_TRANSACTIONS_SHEET,
    INVESTMENT_IMPORT_LOG_SHEET,
    read_sheet,
)
from finance_parser.investments.correct_transaction_trade_date import correct_trade_date


def _write_workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = INVESTMENT_RAW_SHEET
    ws.append(INVESTMENT_RAW_COLUMNS)
    ws.append([("RAW-1" if c == "InvestmentRawID" else "") for c in INVESTMENT_RAW_COLUMNS])

    ws2 = wb.create_sheet(INVESTMENT_TRANSACTIONS_SHEET)
    ws2.append(INVESTMENT_TRANSACTIONS_COLUMNS)
    row1 = {c: "" for c in INVESTMENT_TRANSACTIONS_COLUMNS}
    row1.update({
        "InvestmentTransactionID": "T-1", "InvestmentRawID": "RAW-1", "Broker": "OP", "Portfolio": "OP",
        "NormalizedInstrument": "SAMPO A", "TransactionType": "BUY",
        "TradeDate": "2020-01-15", "SettlementDate": "2020-01-15", "CashAmount": -1000,
    })
    row2 = {c: "" for c in INVESTMENT_TRANSACTIONS_COLUMNS}
    row2.update({
        "InvestmentTransactionID": "T-2", "InvestmentRawID": "RAW-2", "Broker": "NORDNET", "Portfolio": "1",
        "NormalizedInstrument": "FORTUM", "TransactionType": "BUY",
        "TradeDate": "2020-01-15", "SettlementDate": "2020-01-15", "CashAmount": -100,
    })
    ws2.append([row1.get(c, "") for c in INVESTMENT_TRANSACTIONS_COLUMNS])
    ws2.append([row2.get(c, "") for c in INVESTMENT_TRANSACTIONS_COLUMNS])

    ws3 = wb.create_sheet(INVESTMENT_IMPORT_LOG_SHEET)
    ws3.append(INVESTMENT_IMPORT_LOG_COLUMNS)

    wb.save(path)


def test_corrects_only_the_matched_row(tmp_path):
    """
    Real bug this fixes: an inherited security's broker-reported date is the
    deceased's date of death (correct for the tax cost-basis/holding-period
    start), not when it actually reached the heir's account and started
    counting toward net worth - which, per the user, can be years later
    depending on how long probate takes.
    """
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_workbook(path)

    stats = correct_trade_date(path, "T-1", "2021-03-01")

    assert stats["old_trade_date"] == "2020-01-15"
    assert stats["new_trade_date"] == "2021-03-01"
    assert stats["old_settlement_date"] == "2020-01-15"
    assert stats["new_settlement_date"] == "2021-03-01"
    assert stats["rows_matched"] == 1

    transactions = read_sheet(path, INVESTMENT_TRANSACTIONS_SHEET, INVESTMENT_TRANSACTIONS_COLUMNS)
    by_id = transactions.set_index("InvestmentTransactionID")
    assert by_id.loc["T-1", "TradeDate"] == "2021-03-01"
    assert by_id.loc["T-1", "SettlementDate"] == "2021-03-01"
    assert by_id.loc["T-1", "CashAmount"] == -1000  # untouched
    # The unrelated row must be completely untouched.
    assert by_id.loc["T-2", "TradeDate"] == "2020-01-15"


def test_settlement_date_can_differ_from_trade_date(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_workbook(path)

    stats = correct_trade_date(path, "T-1", "2021-03-01", new_settlement_date="2021-03-02")

    assert stats["new_trade_date"] == "2021-03-01"
    assert stats["new_settlement_date"] == "2021-03-02"


def test_raw_sheet_is_never_modified(tmp_path):
    """The raw sheet keeps the true tax-basis date - still correct and
    needed for real future capital-gains-tax calculations."""
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_workbook(path)

    correct_trade_date(path, "T-1", "2021-03-01")

    raw = read_sheet(path, INVESTMENT_RAW_SHEET, INVESTMENT_RAW_COLUMNS)
    assert raw.iloc[0]["InvestmentRawID"] == "RAW-1"


def test_raises_on_unknown_transaction_id(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_workbook(path)

    with pytest.raises(ValueError):
        correct_trade_date(path, "DOES-NOT-EXIST", "2021-03-01")


def test_dry_run_does_not_modify_the_workbook(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_workbook(path)

    stats = correct_trade_date(path, "T-1", "2021-03-01", dry_run=True)
    assert stats["new_trade_date"] == "2021-03-01"

    transactions = read_sheet(path, INVESTMENT_TRANSACTIONS_SHEET, INVESTMENT_TRANSACTIONS_COLUMNS)
    assert transactions.set_index("InvestmentTransactionID").loc["T-1", "TradeDate"] == "2020-01-15"
