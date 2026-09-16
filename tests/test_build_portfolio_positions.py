from pathlib import Path

import pytest
from openpyxl import Workbook

from finance_parser import settings as settings_module
from finance_parser.settings import AppSettings
from finance_parser.investments.build_portfolio_positions import build_positions


TRANSACTIONS_HEADERS = [
    "Broker", "Portfolio", "PortfolioOwner", "NormalizedInstrument", "TransactionType",
    "TradeDate", "Quantity", "CashAmount",
]


def _with_settings(yaml_text: str, tmp_path: Path):
    settings_path = tmp_path / "settings.yaml"
    settings_path.write_text(yaml_text, encoding="utf-8")
    return AppSettings.load(settings_path=settings_path, example_path=Path("does-not-exist.yaml"))


_BASE_SETTINGS_YAML = """
project:
  name: "Test"
  locale: "fi_FI"
  default_currency: "EUR"

paths:
  budgeting_input: "input/budgeting"
  budgeting_output: "output/budgeting/ParsedTransactions.xlsx"
  budgeting_rules: "rules/budgeting/TransactionRules.xlsx"
  investment_input: "input/investments"
  investment_output: "output/investments/ParsedInvestments.xlsx"
  investment_rules: "rules/investments/InstrumentMaster.xlsx"

budgeting:
  default_include: "YES"
  source_bank_aliases: {{}}
  source_account_inference: {{}}

investments:
  portfolio_owners:
    {portfolio_owners}
"""


def _write_transactions(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "InvestmentTransactions"
    ws.append(TRANSACTIONS_HEADERS)
    for row in rows:
        ws.append([row.get(h, "") for h in TRANSACTIONS_HEADERS])
    wb.save(path)


OPENING_POSITIONS_HEADERS = ["Broker", "Portfolio", "NormalizedInstrument", "Date", "Quantity", "CashAmount", "Notes"]


def _write_instrument_master(path, opening_position_rows, master_rows=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "InstrumentMaster"
    master_headers = ["NormalizedInstrument", "CostBasisMethod"]
    ws.append(master_headers)
    for row in (master_rows or []):
        ws.append([row.get(h, "") for h in master_headers])
    opening = wb.create_sheet("OpeningPositions")
    opening.append(OPENING_POSITIONS_HEADERS)
    for row in opening_position_rows:
        opening.append([row.get(h, "") for h in OPENING_POSITIONS_HEADERS])
    wb.save(path)


def _no_instrument_master(tmp_path):
    """A path that deliberately doesn't exist, so build_positions() gets no
    opening-position seeding unless a test explicitly writes one - keeps
    tests isolated from the real rules/investments/InstrumentMaster.xlsx,
    which build_positions() would otherwise read by default."""
    return tmp_path / "InstrumentMaster.xlsx"


def test_build_positions_simple_buy_sell_tracks_cumulative_totals(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_transactions(path, [
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "SAMPO A", "TransactionType": "BUY", "TradeDate": "2021-01-01", "Quantity": 100, "CashAmount": -1000},
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "SAMPO A", "TransactionType": "SELL", "TradeDate": "2021-06-01", "Quantity": 40, "CashAmount": 450},
    ])

    positions, stats = build_positions(path, _no_instrument_master(tmp_path))

    rows = positions[positions["NormalizedInstrument"] == "SAMPO A"]
    assert list(rows["CumulativeQuantity"]) == [100.0, 60.0]
    assert list(rows["CumulativeNetInvested"]) == [-1000.0, -550.0]
    assert stats["negative_quantity_instruments"] == {}


def test_build_positions_handles_inconsistent_sell_quantity_sign(tmp_path):
    """
    Real data: NORDNET reports SELL as a positive magnitude, EVLI reports it
    already negative. build_positions() must always derive sign from
    TransactionType via abs(Quantity), never trust the raw sign.
    """
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_transactions(path, [
        {"Broker": "EVLI", "Portfolio": "1", "NormalizedInstrument": "NOKIA", "TransactionType": "BUY", "TradeDate": "2021-01-01", "Quantity": 50, "CashAmount": -500},
        {"Broker": "EVLI", "Portfolio": "1", "NormalizedInstrument": "NOKIA", "TransactionType": "SELL", "TradeDate": "2021-06-01", "Quantity": -20, "CashAmount": 220},
    ])

    positions, _ = build_positions(path, _no_instrument_master(tmp_path))

    rows = positions[positions["NormalizedInstrument"] == "NOKIA"]
    assert list(rows["CumulativeQuantity"]) == [50.0, 30.0]


def test_build_positions_stock_split_resets_to_absolute_quantity(tmp_path):
    """
    A SPLIT AP JÄTTÖ row's Quantity is the real resulting absolute share
    count, not an incremental delta - confirmed against real Finnair/
    Norwegian Air Shuttle data. build_positions() must land the cumulative
    total exactly on that value, not add it on top of the running total.
    """
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_transactions(path, [
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "FINNAIR", "TransactionType": "BUY", "TradeDate": "2021-01-01", "Quantity": 1200, "CashAmount": -750},
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "FINNAIR", "TransactionType": "SELL", "TradeDate": "2023-11-03", "Quantity": 400, "CashAmount": 27},
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "FINNAIR", "TransactionType": "SPLIT AP JÄTTÖ", "TradeDate": "2024-03-21", "Quantity": 120, "CashAmount": 0},
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "FINNAIR", "TransactionType": "SPLIT AP OTTO", "TradeDate": "2024-03-21", "Quantity": 800, "CashAmount": 0},
    ])

    positions, stats = build_positions(path, _no_instrument_master(tmp_path))

    rows = positions[positions["NormalizedInstrument"] == "FINNAIR"]
    assert list(rows["CumulativeQuantity"]) == [1200.0, 800.0, 120.0]
    assert stats["negative_quantity_instruments"] == {}


def test_build_positions_flags_negative_quantity_as_warning_not_silent(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_transactions(path, [
        {"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "OP-DELTA A", "TransactionType": "BUY", "TradeDate": "2020-01-01", "Quantity": 10, "CashAmount": -100},
        {"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "OP-DELTA A", "TransactionType": "VAIHTO AP-OTTO", "TradeDate": "2021-01-01", "Quantity": 64, "CashAmount": 0},
    ])

    positions, stats = build_positions(path, _no_instrument_master(tmp_path))

    assert stats["negative_quantity_instruments"] == {("OP", "OP", "OP-DELTA A"): -54.0}


def test_build_positions_excludes_known_neutral_types_from_unhandled_report(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_transactions(path, [
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "SAMPO A", "TransactionType": "DIVIDEND", "TradeDate": "2021-03-01", "Quantity": 0, "CashAmount": 12.5},
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "SAMPO A", "TransactionType": "SOME UNKNOWN TYPE", "TradeDate": "2021-04-01", "Quantity": 0, "CashAmount": 0},
    ])

    _, stats = build_positions(path, _no_instrument_master(tmp_path))

    assert stats["unhandled_types"] == {"SOME UNKNOWN TYPE": 1}


def test_build_positions_flags_positive_net_invested_as_warning_not_silent(tmp_path):
    """
    Real case (numbers here invented, not the real ones): OP-Suomi A
    arrived via a VAIHTO - JÄTTÖ transfer with no recorded cost
    (CashAmount=0 - the real cost is unknown, not zero), then a later SELL
    added real proceeds on top of that unknown zero, making
    CumulativeNetInvested go positive - implying a "gain" from investing
    nothing. Must be flagged, not silently reported as if it were correct.
    """
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_transactions(path, [
        {"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "OP-SUOMI A", "TransactionType": "VAIHTO - JÄTTÖ", "TradeDate": "2017-09-22", "Quantity": 100, "CashAmount": 0},
        {"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "OP-SUOMI A", "TransactionType": "SELL", "TradeDate": "2025-03-05", "Quantity": 60, "CashAmount": 3000},
    ])

    positions, stats = build_positions(path, _no_instrument_master(tmp_path))

    rows = positions[positions["NormalizedInstrument"] == "OP-SUOMI A"]
    assert list(rows["CumulativeNetInvested"]) == [0.0, 3000.0]
    assert stats["positive_net_invested_instruments"] == {("OP", "OP", "OP-SUOMI A"): 3000.0}


def test_build_positions_excludes_blank_normalized_instrument_rows(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_transactions(path, [
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "", "TransactionType": "FEE", "TradeDate": "2021-01-01", "Quantity": 0, "CashAmount": -5},
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "SAMPO A", "TransactionType": "BUY", "TradeDate": "2021-01-01", "Quantity": 10, "CashAmount": -100},
    ])

    positions, stats = build_positions(path, _no_instrument_master(tmp_path))

    assert stats["transactions_scanned"] == 1
    assert len(positions) == 1


def test_build_positions_seeds_opening_balance_before_real_transaction_history(tmp_path):
    """
    A gift/inherited holding with no purchase row in any broker export (real
    case: OP-Eurooppa/OP-Suomi Pienyhtiöt A, received as gifts before
    tracking began) is seeded from InstrumentMaster's OpeningPositions sheet
    as the earliest event for that instrument, so a later SELL doesn't go
    negative.
    """
    investments_path = tmp_path / "ParsedInvestments.xlsx"
    _write_transactions(investments_path, [
        {"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "OP-EUROOPPA PIENYHTIÖT A", "TransactionType": "SELL", "TradeDate": "2024-12-27", "Quantity": 0.7761, "CashAmount": 30.0},
    ])

    instrument_master_path = tmp_path / "InstrumentMaster.xlsx"
    _write_instrument_master(instrument_master_path, [
        # Deliberately mixed-case, matching how a human would actually type it
        # into the sheet - NormalizedInstrument matching must be
        # case-insensitive (real bug: it silently failed to join at first).
        {"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "OP-Eurooppa Pienyhtiöt A", "Date": "2013-10-10", "Quantity": 0.78},
    ])

    positions, stats = build_positions(investments_path, instrument_master_path)

    rows = positions[positions["NormalizedInstrument"] == "OP-EUROOPPA PIENYHTIÖT A"]
    assert list(rows["TransactionType"]) == ["OPENING BALANCE", "SELL"]
    assert list(rows["CumulativeQuantity"]) == pytest.approx([0.78, 0.0039])
    assert stats["negative_quantity_instruments"] == {}


def test_build_positions_honors_manually_supplied_opening_cost_basis(tmp_path):
    """
    A gift's real cost basis (e.g. its tax-assessed value at the time -
    "I still assume I purchased them") can be manually supplied in
    OpeningPositions' own CashAmount column, treated the same as if that
    amount had actually been paid in cash. Without one, cost basis defaults
    to 0.0 (unknown) - this must not silently stay 0 once a real value is
    given, and must not make CumulativeNetInvested go positive after a
    later profitable-looking sell (real bug this replaced).
    """
    investments_path = tmp_path / "ParsedInvestments.xlsx"
    _write_transactions(investments_path, [
        {"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "OP-EUROOPPA PIENYHTIÖT A", "TransactionType": "SELL", "TradeDate": "2024-12-27", "Quantity": 0.7761, "CashAmount": 30.0},
    ])

    instrument_master_path = tmp_path / "InstrumentMaster.xlsx"
    _write_instrument_master(instrument_master_path, [
        {"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "OP-Eurooppa Pienyhtiöt A", "Date": "2013-10-10", "Quantity": 0.78, "CashAmount": -100.0},
    ])

    positions, stats = build_positions(investments_path, instrument_master_path)

    rows = positions[positions["NormalizedInstrument"] == "OP-EUROOPPA PIENYHTIÖT A"]
    assert list(rows["CumulativeNetInvested"]) == [-100.0, -70.0]
    assert stats["positive_net_invested_instruments"] == {}


def test_build_positions_without_opening_positions_sheet_is_a_no_op(tmp_path):
    """InstrumentMaster.xlsx existing but lacking an OpeningPositions sheet
    (the common case - most InstrumentMaster.xlsx files predate this
    feature) must not error, just contribute nothing."""
    investments_path = tmp_path / "ParsedInvestments.xlsx"
    _write_transactions(investments_path, [
        {"Broker": "NORDNET", "Portfolio": "1", "NormalizedInstrument": "SAMPO A", "TransactionType": "BUY", "TradeDate": "2021-01-01", "Quantity": 100, "CashAmount": -1000},
    ])

    instrument_master_path = tmp_path / "InstrumentMaster.xlsx"
    wb = Workbook()
    wb.active.title = "InstrumentMaster"
    wb.save(instrument_master_path)

    positions, stats = build_positions(investments_path, instrument_master_path)

    assert stats["transactions_scanned"] == 1
    assert len(positions) == 1


def test_build_positions_carries_portfolio_owner_directly_from_transactions(tmp_path):
    path = tmp_path / "ParsedInvestments.xlsx"
    _write_transactions(path, [
        {"Broker": "NORDNET", "Portfolio": "1", "PortfolioOwner": "PERSON_A", "NormalizedInstrument": "SAMPO A", "TransactionType": "BUY", "TradeDate": "2021-01-01", "Quantity": 100, "CashAmount": -1000},
    ])

    positions, _ = build_positions(path, _no_instrument_master(tmp_path))

    assert list(positions["PortfolioOwner"]) == ["PERSON_A"]


def test_build_positions_backfills_blank_portfolio_owner_from_settings(tmp_path):
    """
    A row parsed before portfolio_owners existed in settings.yaml (or from a
    broker export that never carries owner info at all, e.g. OP) has a blank
    PortfolioOwner. build_positions() should backfill it from settings, same
    as apply_portfolio_ownership() already does on the transactions side.
    """
    loaded = _with_settings(_BASE_SETTINGS_YAML.format(portfolio_owners="NORDNET: PERSON_A"), tmp_path)
    old_cache = settings_module._SETTINGS_CACHE
    try:
        settings_module._SETTINGS_CACHE = loaded

        path = tmp_path / "ParsedInvestments.xlsx"
        _write_transactions(path, [
            {"Broker": "NORDNET", "Portfolio": "1", "PortfolioOwner": "", "NormalizedInstrument": "SAMPO A", "TransactionType": "BUY", "TradeDate": "2021-01-01", "Quantity": 100, "CashAmount": -1000},
        ])

        positions, _ = build_positions(path, _no_instrument_master(tmp_path))

        assert list(positions["PortfolioOwner"]) == ["PERSON_A"]
    finally:
        settings_module._SETTINGS_CACHE = old_cache


def test_build_positions_opening_balance_row_gets_owner_from_settings(tmp_path):
    """OpeningPositions rows are seeded straight from InstrumentMaster, not a
    parsed broker export, so they have no PortfolioOwner of their own -
    settings.yaml is the only place they can get one."""
    loaded = _with_settings(_BASE_SETTINGS_YAML.format(portfolio_owners="OP: PERSON_A"), tmp_path)
    old_cache = settings_module._SETTINGS_CACHE
    try:
        settings_module._SETTINGS_CACHE = loaded

        investments_path = tmp_path / "ParsedInvestments.xlsx"
        _write_transactions(investments_path, [
            {"Broker": "OP", "Portfolio": "OP", "PortfolioOwner": "", "NormalizedInstrument": "OP-EUROOPPA PIENYHTIÖT A", "TransactionType": "SELL", "TradeDate": "2024-12-27", "Quantity": 0.5, "CashAmount": 30.0},
        ])

        instrument_master_path = tmp_path / "InstrumentMaster.xlsx"
        _write_instrument_master(instrument_master_path, [
            {"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "OP-Eurooppa Pienyhtiöt A", "Date": "2013-10-10", "Quantity": 0.6},
        ])

        positions, _ = build_positions(investments_path, instrument_master_path)

        rows = positions[positions["NormalizedInstrument"] == "OP-EUROOPPA PIENYHTIÖT A"]
        assert list(rows["PortfolioOwner"]) == ["PERSON_A", "PERSON_A"]
    finally:
        settings_module._SETTINGS_CACHE = old_cache


def test_average_cost_basis_reduces_proportionally_on_partial_sell(tmp_path):
    """
    Real case this fixes (numbers here invented, not the real ones): a fund
    arrives as one lot, then a partial sell for proceeds that exceed the
    original cost. Under the default cash-flow model, CumulativeNetInvested
    goes positive (-1000 + 3000 = 2000), overstating the remaining
    position's unrealized gain. Opted into AVERAGE cost basis, the
    remaining cost basis must instead be reduced proportionally to the
    fraction of units sold: -1000 * (40/100) = -400.
    """
    investments_path = tmp_path / "ParsedInvestments.xlsx"
    _write_transactions(investments_path, [
        {"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "OP-SUOMI A", "TransactionType": "VAIHTO - JÄTTÖ", "TradeDate": "2017-09-22", "Quantity": 100, "CashAmount": -1000},
        {"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "OP-SUOMI A", "TransactionType": "SELL", "TradeDate": "2025-03-05", "Quantity": 60, "CashAmount": 3000},
    ])

    instrument_master_path = tmp_path / "InstrumentMaster.xlsx"
    _write_instrument_master(
        instrument_master_path, [],
        master_rows=[{"NormalizedInstrument": "OP-SUOMI A", "CostBasisMethod": "AVERAGE"}],
    )

    positions, stats = build_positions(investments_path, instrument_master_path)

    rows = positions[positions["NormalizedInstrument"] == "OP-SUOMI A"]
    assert list(rows["CumulativeQuantity"]) == pytest.approx([100, 40])
    assert rows["CumulativeNetInvested"].iloc[-1] == pytest.approx(-400.0, abs=0.01)
    assert stats["positive_net_invested_instruments"] == {}


def test_average_cost_basis_not_opted_in_keeps_cash_flow_behavior(tmp_path):
    """Same transactions as above but with no CostBasisMethod set - must
    keep today's existing (cash-flow) behavior unchanged."""
    investments_path = tmp_path / "ParsedInvestments.xlsx"
    _write_transactions(investments_path, [
        {"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "OP-SUOMI A", "TransactionType": "VAIHTO - JÄTTÖ", "TradeDate": "2017-09-22", "Quantity": 100, "CashAmount": -1000},
        {"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "OP-SUOMI A", "TransactionType": "SELL", "TradeDate": "2025-03-05", "Quantity": 60, "CashAmount": 3000},
    ])

    positions, _ = build_positions(investments_path, _no_instrument_master(tmp_path))

    rows = positions[positions["NormalizedInstrument"] == "OP-SUOMI A"]
    assert rows["CumulativeNetInvested"].iloc[-1] == pytest.approx(2000.0, abs=0.01)


def test_average_cost_basis_reaches_exactly_zero_after_full_sell(tmp_path):
    """A full exit under AVERAGE cost basis must land the remaining cost
    basis exactly at 0, not a residual positive/negative rounding value."""
    investments_path = tmp_path / "ParsedInvestments.xlsx"
    _write_transactions(investments_path, [
        {"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "OP-AMERIKKA A", "TransactionType": "BUY", "TradeDate": "2020-01-01", "Quantity": 10, "CashAmount": -500},
        {"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "OP-AMERIKKA A", "TransactionType": "SELL", "TradeDate": "2021-01-01", "Quantity": 10, "CashAmount": 800},
    ])

    instrument_master_path = tmp_path / "InstrumentMaster.xlsx"
    _write_instrument_master(
        instrument_master_path, [],
        master_rows=[{"NormalizedInstrument": "OP-AMERIKKA A", "CostBasisMethod": "AVERAGE"}],
    )

    positions, _ = build_positions(investments_path, instrument_master_path)

    rows = positions[positions["NormalizedInstrument"] == "OP-AMERIKKA A"]
    assert rows["CumulativeNetInvested"].iloc[-1] == pytest.approx(0.0, abs=1e-9)


def test_average_cost_basis_handles_multiple_buys_then_sell_dca_style(tmp_path):
    """
    Real case: OP-Maltillinen A has 63 monthly BUY rows (dollar-cost
    averaging) then a full sell - average cost basis must correctly sum
    every contribution and reduce it fully to 0 on the full exit, same as a
    single-lot purchase.
    """
    investments_path = tmp_path / "ParsedInvestments.xlsx"
    rows = [
        {"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "OP-MALTILLINEN A", "TransactionType": "BUY", "TradeDate": f"2020-0{m}-10", "Quantity": 1, "CashAmount": -50}
        for m in range(1, 4)
    ]
    rows.append({"Broker": "OP", "Portfolio": "OP", "NormalizedInstrument": "OP-MALTILLINEN A", "TransactionType": "SELL", "TradeDate": "2021-01-01", "Quantity": 3, "CashAmount": 180})
    _write_transactions(investments_path, rows)

    instrument_master_path = tmp_path / "InstrumentMaster.xlsx"
    _write_instrument_master(
        instrument_master_path, [],
        master_rows=[{"NormalizedInstrument": "OP-MALTILLINEN A", "CostBasisMethod": "AVERAGE"}],
    )

    positions, _ = build_positions(investments_path, instrument_master_path)

    result_rows = positions[positions["NormalizedInstrument"] == "OP-MALTILLINEN A"]
    assert result_rows["CumulativeNetInvested"].iloc[-1] == pytest.approx(0.0, abs=1e-9)
