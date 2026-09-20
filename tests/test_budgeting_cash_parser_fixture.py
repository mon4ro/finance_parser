from pathlib import Path

from finance_parser.budgeting.parsers import cash
from finance_parser.common import transaction_type_for_source_bank


FIXTURES = Path(__file__).resolve().parent / "fixtures" / "budgeting"


def test_cash_fixture_parses_expected_fields():
    path = FIXTURES / "cash_sample.xlsx"
    ok, reason = cash.can_parse(path)
    assert ok, reason

    df = cash.parse_file(path, "2026-09-13 12:00:00")
    row = df.iloc[0]

    assert row["SourceBank"] == "CASH"
    assert row["SourceAccount"] == "CASH"
    assert row["Amount"] == -5
    assert row["RawReceiver"] == "Yksityishenkilö"
    assert row["Currency"] == "EUR"
    assert row["ValueDate"] == "2026-05-09"
    assert row["BookingDate"] == "2026-05-09"
    assert row["ExportDate"] == "2026-09-13"
    assert row["RawID"].startswith("CASH-")

    # "Source" (whose cash) is deliberately not part of the canonical raw
    # schema - it must not leak into RawTransactions/UnifiedTransactions.
    assert "Source" not in df.columns


def test_cash_fixture_raw_export_preserves_source_column():
    path = FIXTURES / "cash_sample.xlsx"
    raw = cash.parse_bank_raw_rows(path, "2026-09-13 12:00:00")
    row = raw.iloc[0]

    assert row["BankRawExportID"].startswith("CASHRAW-")
    assert row["SourceBank"] == "CASH"
    # Unlike RawTransactions, the raw-export sheet keeps "Source" for audit.
    assert row["Source"] == "PERSONAL"
    assert row["EntryID"] == "CASH-0001"


def test_cash_parses_native_excel_date_cells(tmp_path):
    """
    Manually typed rows in CashEntries.xlsx get auto-converted by Excel into
    real date-typed cells (not text), since Excel recognises the user's
    locale date format on entry. This must parse the same as plain ISO text.
    """
    from datetime import date
    from openpyxl import Workbook

    path = tmp_path / "native_date_cash.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "CashEntries"
    ws.append(["EntryID", "Date", "Source", "Amount", "Currency", "Description", "RawReceiver", "ExportDate"])
    ws.append(["CASH-0004", date(2026, 5, 10), "PERSONAL", -3, "EUR", "Kioski snack", "Kioski", date(2026, 9, 13)])
    wb.save(path)

    df = cash.parse_file(path, "2026-09-13 12:00:00")
    row = df.iloc[0]

    assert row["ValueDate"] == "2026-05-10"
    assert row["BookingDate"] == "2026-05-10"
    assert row["ExportDate"] == "2026-09-13"


def test_bank_raw_export_id_stable_across_date_cell_type(tmp_path):
    """
    A cell's date can be typed as plain text or converted by Excel/us into a
    native date - the identity of the audit row must not change just because
    the cell's type changed (this caused a real duplicate row in production:
    fixing CASH-0001's date cell type from text to a real date silently
    produced a second CashRawExport row for the same entry).
    """
    from datetime import date
    from openpyxl import Workbook

    text_path = tmp_path / "text_date.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "CashEntries"
    ws.append(["EntryID", "Date", "Source", "Amount", "Currency", "Description", "RawReceiver", "ExportDate"])
    ws.append(["CASH-0001", "2026-05-09", "PERSONAL", -5, "EUR", "Test item", "Yksityishenkilö", "2026-09-13"])
    wb.save(text_path)

    native_path = tmp_path / "native_date.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "CashEntries"
    ws.append(["EntryID", "Date", "Source", "Amount", "Currency", "Description", "RawReceiver", "ExportDate"])
    ws.append(["CASH-0001", date(2026, 5, 9), "PERSONAL", -5, "EUR", "Test item", "Yksityishenkilö", date(2026, 9, 13)])
    wb.save(native_path)

    raw_text = cash.parse_bank_raw_rows(text_path, "2026-09-13 12:00:00")
    raw_native = cash.parse_bank_raw_rows(native_path, "2026-09-13 12:00:00")

    assert raw_text.iloc[0]["BankRawExportID"] == raw_native.iloc[0]["BankRawExportID"]


def test_vinted_prefixed_entry_id_gets_vinted_source_bank_and_raw_id(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "cash_with_vinted_rows.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "CashEntries"
    ws.append(["EntryID", "Date", "Source", "Amount", "Currency", "Description", "RawReceiver", "ExportDate"])
    ws.append(["CASH-0001", "2026-05-09", "PERSON_A", -5, "EUR", "Kioski snack", "Kioski", "2026-09-13"])
    ws.append(["VINTED-0001", "2026-03-12", "PERSON_B", 20, "EUR", "Vinted sale: Test jacket", "VINTED", "2026-09-20"])
    ws.append(["VINTED-0002", "2026-03-16", "PERSON_B", -8, "EUR", "Vinted purchase: Test onesie", "VINTED", "2026-09-20"])
    wb.save(path)

    df = cash.parse_file(path, "2026-09-20 12:00:00")
    # EntryID is internal-only (used for hashing) and not part of RAW_COLUMNS
    # - see test_cash_fixture_parses_expected_fields's "Source" check above
    # for the same pattern - so rows are identified by Description here.

    physical = df[df["Description"] == "Kioski snack"].iloc[0]
    assert physical["SourceBank"] == "CASH"
    assert physical["SourceAccount"] == "CASH"
    assert physical["RawID"].startswith("CASH-")

    sale = df[df["Description"] == "Vinted sale: Test jacket"].iloc[0]
    assert sale["SourceBank"] == "VINTED"
    assert sale["SourceAccount"] == "CASH"
    assert sale["RawID"].startswith("VINTED-")

    purchase = df[df["Description"] == "Vinted purchase: Test onesie"].iloc[0]
    assert purchase["SourceBank"] == "VINTED"
    assert purchase["Amount"] == -8


def test_vinted_prefixed_entry_id_reflected_in_bank_raw_export(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "cash_with_vinted_rows.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "CashEntries"
    ws.append(["EntryID", "Date", "Source", "Amount", "Currency", "Description", "RawReceiver", "ExportDate"])
    ws.append(["CASH-0001", "2026-05-09", "PERSON_A", -5, "EUR", "Kioski snack", "Kioski", "2026-09-13"])
    ws.append(["VINTED-0001", "2026-03-12", "PERSON_B", 20, "EUR", "Vinted sale: Test jacket", "VINTED", "2026-09-20"])
    wb.save(path)

    raw = cash.parse_bank_raw_rows(path, "2026-09-20 12:00:00")

    physical = raw[raw["EntryID"] == "CASH-0001"].iloc[0]
    assert physical["SourceBank"] == "CASH"

    sale = raw[raw["EntryID"] == "VINTED-0001"].iloc[0]
    assert sale["SourceBank"] == "VINTED"
    assert sale["BankRawExportID"].startswith("CASHRAW-")


def test_transaction_type_for_vinted_source_bank_is_virtual():
    assert transaction_type_for_source_bank("VINTED") == "Virtual"
    assert transaction_type_for_source_bank("vinted") == "Virtual"
    assert transaction_type_for_source_bank("CASH") == "Cash"
    assert transaction_type_for_source_bank("OP") == "Bank"


def test_cash_raw_ids_differ_for_same_content_different_entry_id(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "two_identical_cash_purchases.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "CashEntries"
    ws.append(["EntryID", "Date", "Source", "Amount", "Currency", "Description", "RawReceiver", "ExportDate"])
    ws.append(["CASH-0002", "2026-05-10", "PERSONAL", -2, "EUR", "Coffee", "Kioski", "2026-09-13"])
    ws.append(["CASH-0003", "2026-05-10", "PERSONAL", -2, "EUR", "Coffee", "Kioski", "2026-09-13"])
    wb.save(path)

    df = cash.parse_file(path, "2026-09-13 12:00:00")

    # Two genuinely separate transactions with identical date/amount/receiver
    # text must not collapse into a single RawID.
    assert len(df) == 2
    assert df.iloc[0]["RawID"] != df.iloc[1]["RawID"]
