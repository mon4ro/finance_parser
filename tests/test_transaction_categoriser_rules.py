from pathlib import Path

from openpyxl import Workbook

from finance_parser.budgeting.transaction_categoriser import (
    Rule,
    categorise_workbook,
    rule_matches,
    value_matches,
)


def make_rules_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "CategoryRules"
    ws.append([
        "RuleID", "Enabled", "Priority", "RuleName", "RuleGroup",
        "MatchField", "MatchType", "Pattern", "CaseSensitive",
        "ConditionField1", "ConditionMatchType1", "ConditionPattern1",
        "ConditionField2", "ConditionMatchType2", "ConditionPattern2",
        "SetSupercategory", "SetCategory", "SetSubcategory", "SetOwner", "SetTag", "SetComments",
        "StopIfMatched", "OverwriteMode", "Confidence", "Notes",
    ])
    ws.append([
        "CR0001", "YES", 100, "Hesburger", "RESTAURANT",
        "NormalizedReceiver", "EXACT", "HESBURGER", "NO",
        None, None, None,
        None, None, None,
        "EXPENSES", "Restaurants", "Fast food", None, None, None,
        "YES", "BLANK_ONLY", "HIGH", None,
    ])
    ws.append([
        "CR0002", "YES", 100, "Manual review comment", "MANUAL_REVIEW",
        "Message", "CONTAINS", "needs review", "NO",
        None, None, None,
        None, None, None,
        None, None, None, None, None, "CHECK MANUALLY",
        "YES", "BLANK_ONLY", "LOW", None,
    ])

    os = wb.create_sheet("OwnershipRules")
    os.append([
        "RuleID", "Enabled", "Priority", "RuleName", "MatchField", "MatchType", "Pattern", "CaseSensitive",
        "SetOwner", "ClearOwner", "StopIfMatched", "OverwriteMode", "Notes",
    ])
    os.append(["OR0001", "YES", 100, "PERSON_1 source", "SourceAccount", "EXACT", "PERSON_1", "NO", "PERSON_1", "NO", "YES", "BLANK_ONLY", None])

    wb.save(path)


def make_transactions_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "UnifiedTransactions"
    ws.append([
        "UnifiedID", "RawID", "SourceBank", "SourceAccount", "Include", "Amount",
        "RawReceiver", "NormalizedReceiver", "Description", "Message",
        "Supercategory", "Category", "Subcategory", "Owner", "Review/Notes",
    ])
    ws.append(["U1", "R1", "OP", "PERSON_1", "YES", -10, "Hesburger", "HESBURGER", "", "", "", "", "", "", ""])
    ws.append(["U2", "R2", "OP", "PERSON_1", "YES", -20, "Hesburger", "HESBURGER", "", "", "", "", "", "", "MANUAL BYPASS"])
    wb.save(path)


def test_value_matches_common_match_types():
    assert value_matches("HESBURGER", "EXACT", "hesburger", "NO")
    assert value_matches("ABC LATAUS", "CONTAINS", "lataus", "NO")
    assert value_matches("MOB.PAY*TEST", "STARTS_WITH", "MOB.PAY", "NO")
    assert value_matches("abc123", "REGEX", r"abc\d+", "NO")
    assert value_matches("HUS TAMMI", "ALL_CONTAINS", "HUS, TAMMI", "NO")


def test_rule_matches_with_conditions():
    rule = Rule(
        rule_id="X",
        enabled=True,
        priority=1,
        source_sheet="CategoryRules",
        raw={
            "MatchField": "NormalizedReceiver",
            "MatchType": "EXACT",
            "Pattern": "HESBURGER",
            "CaseSensitive": "NO",
            "ConditionField1": "SourceAccount",
            "ConditionMatchType1": "EXACT",
            "ConditionPattern1": "PERSON_1",
        },
    )
    assert rule_matches(rule, {"NormalizedReceiver": "HESBURGER", "SourceAccount": "PERSON_1"}, category_rule=True)
    assert not rule_matches(rule, {"NormalizedReceiver": "HESBURGER", "SourceAccount": "PERSON_2"}, category_rule=True)


def test_categorise_only_changes_unified_rows_with_blank_comments(tmp_path):
    workbook = tmp_path / "ParsedTransactions.xlsx"
    rules = tmp_path / "TransactionRules.xlsx"
    make_transactions_workbook(workbook)
    make_rules_workbook(rules)

    changes = categorise_workbook(workbook, rules, dry_run=True)
    assert len(changes) == 1
    assert changes[0].unified_id == "U1"
    assert changes[0].changes["Category"] == ("", "Restaurants")

    # Dry run did not write.
    from openpyxl import load_workbook
    wb = load_workbook(workbook)
    ws = wb["UnifiedTransactions"]
    assert ws["L2"].value in (None, "")
    wb.close()

    changes = categorise_workbook(workbook, rules, dry_run=False)
    assert len(changes) == 1

    wb = load_workbook(workbook)
    ws = wb["UnifiedTransactions"]
    assert ws["K2"].value == "EXPENSES"
    assert ws["L2"].value == "Restaurants"
    assert ws["M2"].value == "Fast food"
    assert ws["N2"].value == "PERSON_1"

    # Row with non-blank Comments is intentionally bypassed.
    assert ws["K3"].value in (None, "")
    assert ws["L3"].value in (None, "")
    assert ws["O3"].value == "MANUAL BYPASS"

    assert "ChangeLog" in wb.sheetnames
    wb.close()
