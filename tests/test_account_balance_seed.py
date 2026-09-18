from openpyxl import Workbook

from finance_parser.budgeting.account_balance_seed import (
    accounts_needing_seed,
    collect_account_balance_seed,
    latest_imported_date,
    load_unified_transactions,
)


UNIFIED_COLUMNS = [
    "UnifiedID", "RawID", "SourceAccount", "SourceBank", "TransactionType",
    "Date", "Month", "Year", "Amount", "RawReceiver", "NormalizedReceiver",
    "Description", "Message", "Include", "Owner", "Supercategory",
    "Category", "Subcategory", "ReviewStatus", "Review/Notes",
]


def _write_unified(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "UnifiedTransactions"
    ws.append(UNIFIED_COLUMNS)
    for row in rows:
        ws.append([row.get(h, "") for h in UNIFIED_COLUMNS])
    wb.save(path)


def _feed(monkeypatch, answers):
    queue = list(answers)

    def fake_input(prompt=""):
        if not queue:
            raise AssertionError(f"Ran out of scripted answers at prompt: {prompt!r}")
        return queue.pop(0)

    monkeypatch.setattr("builtins.input", fake_input)


def test_accounts_needing_seed_excludes_cash_and_nordea_only_accounts(tmp_path):
    path = tmp_path / "ParsedTransactions.xlsx"
    _write_unified(path, [
        {"SourceAccount": "HOUSEHOLD", "SourceBank": "OP", "Date": "2026-01-01", "Amount": -10},
        {"SourceAccount": "PERSONAL", "SourceBank": "NORDEA", "Date": "2026-01-01", "Amount": -10},
        {"SourceAccount": "CASH", "SourceBank": "CASH", "Date": "2026-01-01", "Amount": -10},
        {"SourceAccount": "SPANKKI", "SourceBank": "SPANKKI", "Date": "2026-01-01", "Amount": -10},
    ])

    accounts = accounts_needing_seed(path)

    assert accounts == ["HOUSEHOLD", "SPANKKI"]


def test_accounts_needing_seed_keeps_account_mixing_nordea_with_another_bank(tmp_path):
    """
    An account isn't excluded just because SOME of its transactions are
    Nordea - only if ALL of them are - a partial Nordea presence still
    needs a manual seed since Nordea's own real balance field doesn't cover
    the whole account.
    """
    path = tmp_path / "ParsedTransactions.xlsx"
    _write_unified(path, [
        {"SourceAccount": "MIXED", "SourceBank": "OP", "Date": "2026-01-01", "Amount": -10},
        {"SourceAccount": "MIXED", "SourceBank": "NORDEA", "Date": "2026-01-02", "Amount": -10},
    ])

    assert accounts_needing_seed(path) == ["MIXED"]


def test_accounts_needing_seed_empty_workbook_returns_empty(tmp_path):
    path = tmp_path / "does_not_exist.xlsx"
    assert accounts_needing_seed(path) == []


def test_latest_imported_date_returns_max_date_for_account(tmp_path):
    path = tmp_path / "ParsedTransactions.xlsx"
    _write_unified(path, [
        {"SourceAccount": "HOUSEHOLD", "SourceBank": "OP", "Date": "2026-01-01", "Amount": -10},
        {"SourceAccount": "HOUSEHOLD", "SourceBank": "OP", "Date": "2026-03-15", "Amount": -20},
        {"SourceAccount": "OTHER", "SourceBank": "OP", "Date": "2026-06-01", "Amount": -30},
    ])

    assert latest_imported_date("HOUSEHOLD", path).isoformat() == "2026-03-15"


def test_latest_imported_date_no_transactions_returns_none(tmp_path):
    path = tmp_path / "ParsedTransactions.xlsx"
    _write_unified(path, [])
    assert latest_imported_date("HOUSEHOLD", path) is None


def test_collect_account_balance_seed_writes_to_overrides(tmp_path, monkeypatch):
    path = tmp_path / "ParsedTransactions.xlsx"
    _write_unified(path, [
        {"SourceAccount": "HOUSEHOLD", "SourceBank": "OP", "Date": "2026-03-15", "Amount": -10},
    ])
    _feed(monkeypatch, ["", "1234.56"])  # accept default date, then balance
    overrides = {}

    added = collect_account_balance_seed(overrides, "HOUSEHOLD", budgeting_workbook=path)

    assert added is True
    assert overrides["budgeting"]["account_balance_seeds"]["HOUSEHOLD"]["2026-03-15"] == 1234.56


def test_collect_account_balance_seed_future_date_requires_explicit_confirmation(tmp_path, monkeypatch):
    path = tmp_path / "ParsedTransactions.xlsx"
    _write_unified(path, [
        {"SourceAccount": "HOUSEHOLD", "SourceBank": "OP", "Date": "2026-03-15", "Amount": -10},
    ])
    _feed(monkeypatch, ["2026-04-01", "n"])  # later date, decline the warning
    overrides = {}

    added = collect_account_balance_seed(overrides, "HOUSEHOLD", budgeting_workbook=path)

    assert added is False
    assert overrides == {}


def test_collect_account_balance_seed_future_date_accepted_when_confirmed(tmp_path, monkeypatch):
    path = tmp_path / "ParsedTransactions.xlsx"
    _write_unified(path, [
        {"SourceAccount": "HOUSEHOLD", "SourceBank": "OP", "Date": "2026-03-15", "Amount": -10},
    ])
    _feed(monkeypatch, ["2026-04-01", "y", "500.00"])
    overrides = {}

    added = collect_account_balance_seed(overrides, "HOUSEHOLD", budgeting_workbook=path)

    assert added is True
    assert overrides["budgeting"]["account_balance_seeds"]["HOUSEHOLD"]["2026-04-01"] == 500.00


def test_collect_account_balance_seed_blank_date_skips(tmp_path, monkeypatch):
    path = tmp_path / "ParsedTransactions.xlsx"
    _write_unified(path, [])
    _feed(monkeypatch, [""])
    overrides = {}

    added = collect_account_balance_seed(overrides, "HOUSEHOLD", budgeting_workbook=path)

    assert added is False


def test_collect_account_balance_seed_non_numeric_balance_skips(tmp_path, monkeypatch):
    path = tmp_path / "ParsedTransactions.xlsx"
    _write_unified(path, [
        {"SourceAccount": "HOUSEHOLD", "SourceBank": "OP", "Date": "2026-03-15", "Amount": -10},
    ])
    _feed(monkeypatch, ["", "not a number"])
    overrides = {}

    added = collect_account_balance_seed(overrides, "HOUSEHOLD", budgeting_workbook=path)

    assert added is False
    assert overrides == {}


def test_load_unified_transactions_fills_balance_from_raw_transactions(tmp_path):
    """
    Real bug found and fixed: a transaction unified before Balance was added
    to UnifiedTransactions' own schema stays blank there forever (this
    project doesn't re-unify retroactively) - but RawTransactions always has
    the real value (Nordea's own "Saldo" field, never re-derived), so a gap
    in UnifiedTransactions' own Balance column must be filled from there,
    joined by RawID, not left blank.
    """
    path = tmp_path / "ParsedTransactions.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "UnifiedTransactions"
    ws.append(UNIFIED_COLUMNS)
    ws.append(["U-1", "RAW-1", "PERSONAL", "NORDEA", "PAYMENT", "2026-03-15", 3, 2026, -10, "", "", "", "", "YES", "PERSONAL", "", "", "", "", ""])

    raw_ws = wb.create_sheet("RawTransactions")
    raw_ws.append(["RawID", "Balance"])
    raw_ws.append(["RAW-1", 456.78])
    wb.save(path)

    df = load_unified_transactions(path)

    assert float(df.iloc[0]["Balance"]) == 456.78


def test_load_unified_transactions_prefers_own_balance_over_raw_fallback(tmp_path):
    """A row that already has its own real Balance (unified after the fix)
    must not get overwritten by the raw fallback."""
    path = tmp_path / "ParsedTransactions.xlsx"
    columns_with_balance = UNIFIED_COLUMNS + ["Balance"]
    wb = Workbook()
    ws = wb.active
    ws.title = "UnifiedTransactions"
    ws.append(columns_with_balance)
    ws.append(["U-1", "RAW-1", "PERSONAL", "NORDEA", "PAYMENT", "2026-03-15", 3, 2026, -10, "", "", "", "", "YES", "PERSONAL", "", "", "", "", "", 999.99])

    raw_ws = wb.create_sheet("RawTransactions")
    raw_ws.append(["RawID", "Balance"])
    raw_ws.append(["RAW-1", 456.78])
    wb.save(path)

    df = load_unified_transactions(path)

    assert float(df.iloc[0]["Balance"]) == 999.99
