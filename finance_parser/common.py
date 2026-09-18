from __future__ import annotations

import csv
import hashlib
import io
import os
import re
import uuid
import shutil
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path
from typing import Iterable

import pandas as pd
from finance_parser.settings import get_settings
from openpyxl import Workbook, load_workbook


RAW_SHEET = "RawTransactions"
UNIFIED_SHEET = "UnifiedTransactions"
IMPORT_LOG_SHEET = "ImportLog"
CHANGE_LOG_SHEET = "ChangeLog"

CHANGE_LOG_COLUMNS = [
    "ChangeID",
    "ChangedAt",
    "Script",
    "Action",
    "Workbook",
    "Sheet",
    "RowsBefore",
    "RowsAfter",
    "RowsAdded",
    "RowsUpdated",
    "RowsRemoved",
    "ReviewFile",
    "BackupFile",
    "Status",
    "Details",
]


# 12 hex chars = 48 bits. Much shorter than 16, still very safe for this use case.
DEFAULT_HASH_LENGTH = 12

RAW_COLUMNS = [
    "RawID",
    "SourceAccount",
    "SourceBank",
    "ExportDate",
    "ImportedAt",
    "BookingDate",
    "ValueDate",
    "Amount",
    "TransactionTypeRaw",
    "Description",
    "RawReceiver",
    "ReceiverAccount",
    "ReceiverBankBIC",
    "Reference",
    "Message",
    "ArchiveID",
    "Balance",
    "CurrencyAmount",
    "Currency",
    "Rate",
    "MerchantArea",
    "MerchantCategory",
    "SourceFile",
]

UNIFIED_COLUMNS = [
    "UnifiedID",
    "RawID",
    "SourceAccount",
    "SourceBank",
    "TransactionType",
    "Date",
    "Month",
    "Year",
    "Amount",
    "RawReceiver",
    "NormalizedReceiver",
    "Description",
    "Message",
    "Include",
    "Owner",
    "Supercategory",
    "Category",
    "Subcategory",
    "ReviewStatus",
    "Review/Notes",
    "Currency",
    "Rate",
    "ExportDate",
    "ImportedAt",
    "SourceFile",
    # Appended, not inserted mid-list - see the same column-order lesson
    # applied on the investment side (build_monthly_position_value.py).
    # Real per-transaction running balance - only Nordea's raw export
    # actually populates this (its own "Saldo" field); every other bank's
    # parser leaves it blank. Was already part of RAW_COLUMNS and correctly
    # populated there, but never carried through to UnifiedTransactions
    # until build_monthly_account_balance.py needed it.
    "Balance",
]

IMPORT_LOG_COLUMNS = [
    "ImportRunID",
    "ImportedAt",
    "SourceBank",
    "SourceFile",
    "RowsRead",
    "RowsNew",
    "RowsDuplicate",
    "Status",
    "Error",
]

SOURCE_TO_DEFAULT_OWNER = {
    "joint": "Shared",
    "shared": "Shared",
    "yhteinen": "Shared",
    "person a": "Person A",
    "personal": "Person A",
    "oma": "Person A",
}

_ILLEGAL_XML_RE = re.compile(
    r"[\x00-\x08\x0B\x0C\x0E-\x1F\uFFFE\uFFFF]"
)


def imported_at_now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def make_import_run_id(imported_at: str) -> str:
    return "RUN-" + stable_hash([imported_at], length=12)


def normalise_header(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\ufeff", "").strip()
    text = re.sub(r"\s+", " ", text)
    return text


def normalise_text(value: object) -> str:
    """
    Convert scalar values to clean strings.

    Defensive behaviour:
    - None/NaN -> ""
    - pandas Series/DataFrame -> "" rather than raising ambiguous truth-value errors

    Series/DataFrame should not normally be passed here; returning blank avoids
    crashing and lets the caller continue, while the real fix is to call this
    function element-wise with .map(...) or list comprehension.
    """
    if value is None:
        return ""

    if isinstance(value, (pd.Series, pd.DataFrame)):
        return ""

    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        # pd.isna can return array-like results for non-scalar objects.
        return ""

    return str(value).strip()

def clean_for_excel(value: object) -> object:
    if value is None or pd.isna(value):
        return ""
    if isinstance(value, str):
        return _ILLEGAL_XML_RE.sub("", value)
    return value


def clean_dataframe_for_excel(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for col in out.columns:
        if out[col].dtype == object:
            out[col] = out[col].map(clean_for_excel)
    return out


def normalise_amount(value: object) -> float:
    if pd.isna(value) or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip()
    text = text.replace("\u00a0", "")
    text = text.replace(" ", "")
    text = text.replace("€", "")

    if "," in text and "." not in text:
        text = text.replace(",", ".")
    elif "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")

    return float(text)


def parse_date(value: object) -> pd.Timestamp | pd.NaT:
    """
    Parse dates safely.

    Important detail:
    - Values already in ISO form yyyy-mm-dd MUST be parsed year-month-day.
      Pandas with dayfirst=True can misread "2026-05-12" as 2026-12-05.
    - Finnish-looking dates such as 12.5.2026 are parsed day-month-year.
    - Excel serial date numbers are supported as a fallback.
    """
    if pd.isna(value) or value == "":
        return pd.NaT

    if isinstance(value, pd.Timestamp):
        return value.normalize()

    text = str(value).strip()
    if not text:
        return pd.NaT

    # ISO date, optionally followed by a time-of-day/timezone suffix (e.g. a
    # full ISO 8601 timestamp like "2026-02-07T11:17:44+02:00" from Coinmotion's
    # export). Match only the leading yyyy-mm-dd and parse that explicitly -
    # anything after it is ignored. Do this before generic pandas parsing to
    # avoid day/month reversal: a full-fullmatch-only check here previously let
    # timestamp strings fall through to dayfirst=True parsing below, which
    # misread "2026-02-07T11:17:44+02:00" as 2 July instead of 7 February.
    iso_match = re.match(r"(\d{4})-(\d{1,2})-(\d{1,2})", text)
    if iso_match:
        year, month, day = iso_match.groups()
        parsed = pd.to_datetime(f"{year}-{month}-{day}", format="%Y-%m-%d", errors="coerce")
        return pd.NaT if pd.isna(parsed) else parsed.normalize()

    # Finnish / European dotted or slashed dates: d.m.yyyy or d/m/yyyy.
    euro_match = re.fullmatch(r"(\d{1,2})[./](\d{1,2})[./](\d{2,4})", text)
    if euro_match:
        day, month, year = euro_match.groups()
        if len(year) == 2:
            year = "20" + year
        parsed = pd.to_datetime(f"{year}-{month}-{day}", format="%Y-%m-%d", errors="coerce")
        return pd.NaT if pd.isna(parsed) else parsed.normalize()

    # Excel serial date fallback. Excel's 1899-12-30 origin matches pandas/openpyxl
    # behaviour for ordinary modern workbook dates.
    if re.fullmatch(r"\d+(\.0)?", text):
        try:
            serial = float(text)
            if 20000 <= serial <= 80000:  # roughly 1954-2119
                parsed = pd.to_datetime(serial, unit="D", origin="1899-12-30", errors="coerce")
                return pd.NaT if pd.isna(parsed) else parsed.normalize()
        except Exception:
            pass

    parsed = pd.to_datetime(value, dayfirst=True, errors="coerce")
    if pd.isna(parsed):
        return pd.NaT
    return parsed.normalize()


def format_date(value: object) -> str:
    parsed = parse_date(value)
    if pd.isna(parsed):
        return ""
    # Keep output as ISO text. This sorts chronologically even as text in Excel.
    return parsed.strftime("%Y-%m-%d")


def file_modified_date(path: Path) -> str:
    """Return file modified date as yyyy-mm-dd."""
    return datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d")


def infer_source_account_from_filename(path: Path, source_bank: str) -> str:
    """
    Infer workbook SourceAccount from filename/settings.

    Source-account inference is user/workbook-specific. The generic parser asks
    settings.py first and only falls back to conservative filename cleanup.
    """
    settings = get_settings()
    bank = settings.canonical_source_bank(source_bank)

    fixed_account = settings.fixed_source_account(bank)
    if fixed_account:
        return fixed_account

    filename_account = settings.filename_source_account(bank, path)
    if filename_account:
        return filename_account

    stem = normalise_text(path.stem).upper()

    # Conservative fallback for unknown OP-like files. Shared users should
    # configure filename_prefixes in config/settings.yaml instead of relying on
    # this.
    cleaned = re.sub(r"TAPAHTUMAT.*$", "", stem).strip()
    cleaned = re.sub(r"\d{6,8}.*$", "", cleaned).strip()
    cleaned = cleaned.replace("_", " ").replace("-", " ").strip()
    cleaned = re.sub(r"\s+", " ", cleaned)

    return cleaned or bank

def stable_hash(parts: Iterable[object], length: int = DEFAULT_HASH_LENGTH) -> str:
    joined = "|".join(normalise_text(p).upper() for p in parts)
    return hashlib.sha256(joined.encode("utf-8")).hexdigest()[:length]


def default_owner_from_source_account(source_account: object) -> str:
    source_text = normalise_text(source_account).lower()
    for key, owner in SOURCE_TO_DEFAULT_OWNER.items():
        if key in source_text:
            return owner
    return ""


def make_unified_id(raw_id: str) -> str:
    return "U-" + raw_id


def canonical_source_bank(row) -> str:
    """
    Return stable SourceBank for identity/deduplication.

    This is intentionally tolerant of older output rows where SourceBank was
    blank or used different capitalisation.
    """
    source_bank = normalise_text(row.get("SourceBank", "")).upper()
    raw_id = normalise_text(row.get("RawID", "")).upper()
    source_file = normalise_text(row.get("SourceFile", "")).upper()

    if source_bank:
        if source_bank in {"NORDEA", "NOD"}:
            return "NORDEA"
        if source_bank in {"NORWEGIAN", "BANK NORWEGIAN", "NWG"}:
            return "NORWEGIAN"
        if source_bank == "OP":
            return "OP"
        if source_bank in {"SPANKKI", "S-PANKKI", "SPK"}:
            return "SPANKKI"
        return source_bank

    if raw_id.startswith("OP-") or source_file.startswith("TILI_"):
        return "OP"
    if raw_id.startswith("NWG-") or raw_id.startswith("NORWEGIAN-") or "NORWEGIAN" in source_file:
        return "NORWEGIAN"
    if raw_id.startswith("NOD-") or "NORDEA" in source_file:
        return "NORDEA"
    if raw_id.startswith("SPK-") or "SPANKKI" in source_file or "S-PANKKI" in source_file:
        return "SPANKKI"

    return ""


def canonical_source_account(value: object, source_bank: object = "") -> str:
    """
    Canonicalise SourceAccount labels using settings first.

    User/workbook-specific mappings such as NORDEA -> a person's account label,
    OP filename prefixes, or SPANKKI fixed account names belong in
    config/settings.yaml, not in generic parser logic.
    """
    if isinstance(value, pd.Series):
        bank_value = "" if isinstance(source_bank, pd.Series) else source_bank
        for item in value.dropna().tolist():
            result = canonical_source_account(item, bank_value)
            if result:
                return result
        return ""

    if isinstance(source_bank, pd.Series):
        source_bank = ""

    settings = get_settings()
    bank = settings.canonical_source_bank(source_bank)
    account = normalise_text(value).upper()

    fixed_account = settings.fixed_source_account(bank)
    if fixed_account:
        return fixed_account

    configured_prefixes = (
        settings.get("budgeting", "source_account_inference", bank, "filename_prefixes", default={})
        or {}
    )
    if configured_prefixes:
        tokens = [
            token
            for token in re.split(r"[^A-ZÅÄÖ0-9]+", account)
            if token
        ]

        for prefix, mapped_account in configured_prefixes.items():
            prefix_norm = normalise_text(prefix).upper()
            mapped_norm = normalise_text(mapped_account).upper()

            if prefix_norm in tokens or account.startswith(prefix_norm):
                return mapped_norm

    return account

def canonical_transaction_key(row) -> str:
    """
    Stable duplicate-detection key independent of RawID.

    This protects us when a parser mapping change alters RawID even though the
    underlying bank transaction is the same. We deliberately exclude Description
    and Message because those mappings may change over time.
    """
    bank = canonical_source_bank(row)
    account = canonical_source_account(row)
    date = format_date(row.get("ValueDate", row.get("Date", "")))
    amount = normalise_amount(row.get("Amount", 0))
    receiver = normalise_text(row.get("RawReceiver", "")).upper()
    reference = normalise_text(row.get("Reference", "")).upper()
    currency = normalise_text(row.get("Currency", "")).upper()

    return stable_hash([
        bank,
        account,
        date,
        f"{amount:.2f}",
        receiver,
        reference,
        currency,
    ], length=20)


def add_canonical_transaction_key(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out = normalize_source_metadata(out)
    if len(out) == 0:
        out["_CanonicalTransactionKey"] = []
        return out
    out["_CanonicalTransactionKey"] = out.apply(canonical_transaction_key, axis=1)
    return out


def normalize_source_metadata(df: pd.DataFrame) -> pd.DataFrame:
    """
    Normalise SourceBank/SourceAccount labels after import and duplicate merging.

    Source-bank aliases and fixed source accounts are read from settings.py.
    """
    out = df.copy()
    settings = get_settings()

    if "SourceBank" in out.columns:
        out["SourceBank"] = out["SourceBank"].map(normalise_text).map(settings.canonical_source_bank)

    if "SourceAccount" in out.columns:
        if "SourceBank" in out.columns:
            out["SourceAccount"] = [
                canonical_source_account(account, bank)
                for account, bank in zip(out["SourceAccount"].tolist(), out["SourceBank"].tolist())
            ]
        else:
            out["SourceAccount"] = out["SourceAccount"].map(lambda account: canonical_source_account(account))

    return out

def merge_blank_metadata_before_dedup(df: pd.DataFrame) -> pd.DataFrame:
    """
    When parser mappings change, the same bank transaction may appear twice:
    an older preserved row and a newer imported row with better metadata.

    We still keep the first row to protect manual category/comment edits, but
    before dropping duplicates we fill blank technical metadata from any
    duplicate rows with the same canonical transaction key.
    """
    out = df.copy()

    if len(out) == 0:
        return out

    if "_CanonicalTransactionKey" not in out.columns:
        out = add_canonical_transaction_key(out)

    metadata_cols = [
        "SourceBank",
        "SourceAccount",
        "ExportDate",
        "ImportedAt",
        "SourceFile",
        "Currency",
        "Rate",
    ]

    for key, idxs in out.groupby("_CanonicalTransactionKey").groups.items():
        idxs = list(idxs)
        if len(idxs) <= 1:
            continue

        # For each metadata column, find the first nonblank value in the duplicate group.
        for col in metadata_cols:
            if col not in out.columns:
                continue

            nonblank_values = []
            for idx in idxs:
                value = out.at[idx, col]
                if pd.notna(value) and str(value).strip() != "":
                    nonblank_values.append(value)

            if not nonblank_values:
                continue

            fill_value = nonblank_values[0]

            # Fill only blanks. Do not overwrite manual edits or existing metadata.
            for idx in idxs:
                value = out.at[idx, col]
                if pd.isna(value) or str(value).strip() == "":
                    out.at[idx, col] = fill_value

    return out


def default_include_for_row(row) -> str:
    """
    Default Include value for new UnifiedTransactions rows.

    This is user/workbook-specific policy and is read from settings instead of
    being hardcoded in parser logic.
    """
    source_bank = row.get("SourceBank", "") if hasattr(row, "get") else ""
    return default_include_for_source_bank(source_bank)


def default_include_for_source_bank(source_bank: object) -> str:
    """
    Return default Include value for a source bank from settings.

    Example settings:
      budgeting.default_include = YES
      budgeting.source_account_inference.SPANKKI.default_include = NO
    """
    return get_settings().default_include_for_source_bank(source_bank)


def transaction_type_for_source_bank(source_bank: object) -> str:
    """
    Return the UnifiedTransactions TransactionType for a given SourceBank.

    Bank-sourced imports are "Bank"; manually entered cash-ledger rows are
    "Cash". Manual SPLIT rows are created afterward, directly in Excel, and
    are not covered here.
    """
    if normalise_text(source_bank).upper() == "CASH":
        return "Cash"
    return "Bank"


def raw_to_unified_rows(raw_new: pd.DataFrame) -> pd.DataFrame:
    rows = []

    for _, row in raw_new.iterrows():
        raw_receiver = row["RawReceiver"]
        description = row["Description"]
        message = row["Message"]

        # Parser-created unified rows should not pretend that RawReceiver is
        # already normalised. These fields are filled later by
        # transaction_normaliser.py or manually in Excel.
        normalized_receiver = ""
        supercategory = ""
        category = ""
        subcategory = ""

        date = row["ValueDate"]
        parsed_date = parse_date(date)

        rows.append({
            "UnifiedID": make_unified_id(row["RawID"]),
            "RawID": row["RawID"],
            "SourceAccount": row["SourceAccount"],
            "SourceBank": row["SourceBank"],
            "TransactionType": transaction_type_for_source_bank(row["SourceBank"]),
            "Date": format_date(date),
            "Month": "" if pd.isna(parsed_date) else int(parsed_date.month),
            "Year": "" if pd.isna(parsed_date) else int(parsed_date.year),
            "Amount": row["Amount"],
            "RawReceiver": raw_receiver,
            "NormalizedReceiver": normalized_receiver,
            "Description": description,
            "Message": message,
            "Include": default_include_for_row(row),
            "Owner": default_owner_from_source_account(row["SourceAccount"]),
            "Supercategory": supercategory,
            "Category": category,
            "Subcategory": subcategory,
            "ReviewStatus": "",
            "Review/Notes": "",
            "Currency": row.get("Currency", "EUR"),
            "Rate": row.get("Rate", 1),
            "ExportDate": row["ExportDate"],
            "ImportedAt": row["ImportedAt"],
            "SourceFile": row["SourceFile"],
            "Balance": row.get("Balance", ""),
        })

    return pd.DataFrame(rows, columns=UNIFIED_COLUMNS)


def sheet_to_dataframe(path: Path, sheet_name: str, columns: list[str]) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame(columns=columns)

    try:
        df = pd.read_excel(path, sheet_name=sheet_name, dtype=object, engine="openpyxl")
    except ValueError:
        return pd.DataFrame(columns=columns)

    df.columns = [normalise_header(c) for c in df.columns]

    # Migration from older single-bank versions.
    if "SourceAccount" in columns and "SourceAccount" not in df.columns and "Source" in df.columns:
        df["SourceAccount"] = df["Source"]

    if "SourceBank" in columns and "SourceBank" not in df.columns:
        df["SourceBank"] = df.apply(infer_source_bank_from_existing_row, axis=1)

    if "SourceBank" in columns and "SourceBank" in df.columns:
        blank_source_bank = df["SourceBank"].isna() | (df["SourceBank"].astype(str).str.strip() == "")
        if blank_source_bank.any():
            inferred = df.loc[blank_source_bank].apply(infer_source_bank_from_existing_row, axis=1)
            df.loc[blank_source_bank, "SourceBank"] = inferred

    # v11 Nordea migration: normalise existing Nordea rows created by older
    # parser versions so they do not duplicate when the parser mapping changes.
    # S-Pankki migration: normalise S-PANKKI/SPK variants to SPANKKI and keep
    # SourceAccount fixed as SPANKKI.
    if "SourceBank" in df.columns:
        nordea_mask = df.apply(lambda r: canonical_source_bank(r) == "NORDEA", axis=1)
        if nordea_mask.any():
            df.loc[nordea_mask, "SourceBank"] = "NORDEA"
            if "SourceAccount" in df.columns:
                nordea_account = get_settings().fixed_source_account("NORDEA") or "NORDEA"
                df.loc[nordea_mask, "SourceAccount"] = nordea_account

        spankki_mask = df.apply(lambda r: canonical_source_bank(r) == "SPANKKI", axis=1)
        if spankki_mask.any():
            df.loc[spankki_mask, "SourceBank"] = "SPANKKI"
            if "SourceAccount" in df.columns:
                df.loc[spankki_mask, "SourceAccount"] = "SPANKKI"

    if "Description" in df.columns and "Message" in df.columns and "SourceBank" in df.columns:
        nordea_mask = df.apply(lambda r: canonical_source_bank(r) == "NORDEA", axis=1)
        if nordea_mask.any():
            blank_or_oldish = df["Description"].isna() | (df["Description"].astype(str).str.strip() == "")
            has_message = df["Message"].notna() & (df["Message"].astype(str).str.strip() != "")
            # Only auto-fill blank descriptions from Message. We avoid overwriting
            # nonblank descriptions because users may have manually edited them.
            fill_mask = nordea_mask & blank_or_oldish & has_message
            df.loc[fill_mask, "Description"] = df.loc[fill_mask, "Message"]

    if "Review/Notes" in columns and "Review/Notes" not in df.columns:
        if "Notes" in df.columns:
            df["Review/Notes"] = df["Notes"]
        else:
            df["Review/Notes"] = ""

    # Migration from old OP-specific raw column names to standard raw names.
    raw_migrations = {
        "Kirjauspäivä": "BookingDate",
        "Arvopäivä": "ValueDate",
        "Määrä EUROA": "Amount",
        "Laji": "TransactionTypeRaw",
        "Selitys": "Description",
        "Saaja/Maksaja": "RawReceiver",
        "Saajan tilinumero": "ReceiverAccount",
        "Saajan pankin BIC": "ReceiverBankBIC",
        "Viite": "Reference",
        "Viesti": "Message",
        "Arkistointitunnus": "ArchiveID",
    }
    for old, new in raw_migrations.items():
        if new in columns and new not in df.columns and old in df.columns:
            df[new] = df[old]

    if "Currency" in columns and "Currency" not in df.columns:
        df["Currency"] = "EUR"

    if "Rate" in columns and "Rate" not in df.columns:
        df["Rate"] = 1

    if "Balance" in columns and "Balance" not in df.columns:
        df["Balance"] = ""

    if "CurrencyAmount" in columns and "CurrencyAmount" not in df.columns:
        df["CurrencyAmount"] = ""

    if "MerchantArea" in columns and "MerchantArea" not in df.columns:
        df["MerchantArea"] = ""

    if "MerchantCategory" in columns and "MerchantCategory" not in df.columns:
        df["MerchantCategory"] = ""

    for col in columns:
        if col not in df.columns:
            df[col] = ""

    return df[columns]


def append_dataframe_to_worksheet(wb: Workbook, sheet_name: str, df: pd.DataFrame) -> None:
    df = clean_dataframe_for_excel(df)

    if sheet_name in wb.sheetnames:
        del wb[sheet_name]

    ws = wb.create_sheet(sheet_name)
    ws.append([clean_for_excel(c) for c in df.columns])

    for row in df.itertuples(index=False, name=None):
        ws.append([clean_for_excel(v) for v in row])

    ws.freeze_panes = "A2"

    if len(df) > 0 and len(df.columns) > 0:
        end_col = ws.cell(row=1, column=len(df.columns)).column_letter
        end_row = len(df) + 1
        ws.auto_filter.ref = f"A1:{end_col}{end_row}"


def write_clean_output_workbook(
    workbook_path: Path,
    raw_df: pd.DataFrame,
    unified_df: pd.DataFrame,
    import_log_df: pd.DataFrame,
) -> None:
    tmp_path = workbook_path.with_name(workbook_path.stem + "_tmp_save" + workbook_path.suffix)
    backup_path = workbook_path.with_name(workbook_path.stem + "_backup_before_last_run" + workbook_path.suffix)

    # 1. Read the existing ChangeLog BEFORE we overwrite the workbook
    change_log_df = pd.DataFrame(columns=CHANGE_LOG_COLUMNS)
    if workbook_path.exists():
        shutil.copy2(workbook_path, backup_path)
        # Extract existing ChangeLog history
        change_log_df = sheet_to_dataframe(workbook_path, CHANGE_LOG_SHEET, CHANGE_LOG_COLUMNS)

    # 2. Create the new workbook
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    # 3. Append all sheets, including the preserved ChangeLog
    append_dataframe_to_worksheet(wb, RAW_SHEET, raw_df)
    append_dataframe_to_worksheet(wb, UNIFIED_SHEET, unified_df)

    from finance_parser.utilities.fresh_workbook_writer import apply_review_status_column
    apply_review_status_column(wb[UNIFIED_SHEET])

    append_dataframe_to_worksheet(wb, IMPORT_LOG_SHEET, import_log_df)
    
    # Only append the ChangeLog sheet if it has data or we want to initialize it
    if not change_log_df.empty or workbook_path.exists():
        append_dataframe_to_worksheet(wb, CHANGE_LOG_SHEET, change_log_df)

    wb.save(tmp_path)

    test_wb = load_workbook(tmp_path, read_only=True)
    test_wb.close()

    os.replace(tmp_path, workbook_path)


# ---------------------------------------------------------------------------
# Robust file reading helpers retained from the OP parser
# ---------------------------------------------------------------------------

def read_delimited_text_any_extension(path: Path, required_columns: list[str] | None = None) -> pd.DataFrame:
    raw = path.read_bytes()

    encodings = [
        "utf-8-sig",
        "utf-8",
        "cp1252",
        "iso-8859-15",
        "latin1",
        "utf-16",
        "utf-16-le",
        "utf-16-be",
    ]
    delimiters = [";", "\t", ",", "|"]

    attempts = []
    best_candidate = None
    best_score = -1
    required_columns = required_columns or []

    for encoding in encodings:
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError as exc:
            attempts.append(f"encoding={encoding} failed decode: {exc}")
            continue

        text = text.replace("\x00", "")
        text = text.replace("\r\n", "\n").replace("\r", "\n")

        for sep in delimiters:
            parse_modes = [
                {"quoting": csv.QUOTE_MINIMAL, "quotechar": '"'},
                {"quoting": csv.QUOTE_NONE, "quotechar": None},
            ]

            for mode in parse_modes:
                try:
                    buffer = io.StringIO(text)
                    kwargs = dict(
                        sep=sep,
                        dtype=object,
                        engine="python",
                        keep_default_na=False,
                        na_values=[],
                        on_bad_lines="skip",
                    )
                    kwargs.update(mode)
                    candidate = pd.read_csv(buffer, **kwargs)
                    candidate.columns = [normalise_header(c) for c in candidate.columns]

                    score = len(set(required_columns).intersection(set(candidate.columns)))
                    total_score = score * 10 + min(len(candidate.columns), 20)

                    attempts.append(
                        f"encoding={encoding}, sep={repr(sep)}, quoting={mode['quoting']} -> "
                        f"{len(candidate.columns)} columns, required={score}, rows={len(candidate)}"
                    )

                    if total_score > best_score:
                        best_score = total_score
                        best_candidate = candidate

                    if required_columns and score >= min(3, len(required_columns)):
                        print(
                            f"Read {path.name} as delimited text: "
                            f"encoding={encoding}, delimiter={repr(sep)}, quoting={mode['quoting']}"
                        )
                        return candidate

                except Exception as exc:
                    attempts.append(
                        f"encoding={encoding}, sep={repr(sep)}, quoting={mode['quoting']} failed: {exc}"
                    )

    if best_candidate is not None:
        detected = list(best_candidate.columns)[:30]
        required_preview = required_columns[:10]
        raise ValueError(
            "Delimited text parsing ran, but it did not contain enough expected columns. "
            "This may still be a valid XLSX file, so the parser will try Excel/XML fallbacks next.\n"
            f"Required preview: {required_preview}\n"
            f"Detected columns preview: {detected}\n"
            "Recent attempts:\n"
            + "\n".join(f"  - {a}" for a in attempts[-20:])
        )

    raise ValueError(
        "Delimited text parse failed completely. Attempts:\n"
        + "\n".join(f"  - {a}" for a in attempts[-30:])
    )


def read_html_guess(path: Path, required_columns: list[str] | None = None) -> pd.DataFrame:
    tables = pd.read_html(path, flavor="lxml")
    if not tables:
        raise ValueError("No HTML tables found")

    best = None
    best_score = -1
    required_columns = required_columns or []

    for table in tables:
        table = table.copy()
        table.columns = [normalise_header(c) for c in table.columns]
        score = len(set(required_columns).intersection(set(table.columns)))
        if score > best_score:
            best = table
            best_score = score

    if best is None:
        raise ValueError("No usable HTML table found")

    return best


def read_xlsx_xml_direct(path: Path, required_columns: list[str] | None = None) -> pd.DataFrame:
    if not zipfile.is_zipfile(path):
        raise ValueError("Not a zipped XLSX file")

    required_columns = required_columns or []

    def lname(tag: str) -> str:
        if "}" in tag:
            return tag.rsplit("}", 1)[1]
        return tag

    def children_named(el, name: str):
        return [child for child in list(el) if lname(child.tag) == name]

    def first_child_named(el, name: str):
        for child in list(el):
            if lname(child.tag) == name:
                return child
        return None

    def iter_named(el, name: str):
        for item in el.iter():
            if lname(item.tag) == name:
                yield item

    with zipfile.ZipFile(path) as z:
        names = z.namelist()

        worksheet_names = sorted(
            n for n in names
            if n.startswith("xl/worksheets/") and n.endswith(".xml")
        )

        if not worksheet_names:
            worksheet_names = sorted(
                n for n in names
                if "worksheets" in n.lower() and n.endswith(".xml")
            )

        if not worksheet_names:
            raise ValueError(
                "No worksheet XML files found inside XLSX. Zip entries preview: "
                + ", ".join(names[:50])
            )

        shared_strings = []
        if "xl/sharedStrings.xml" in names:
            root = ET.fromstring(z.read("xl/sharedStrings.xml"))
            for si in iter_named(root, "si"):
                parts = []
                for t in iter_named(si, "t"):
                    parts.append(t.text or "")
                shared_strings.append("".join(parts))

        def cell_col_index(cell_ref: str) -> int:
            letters = re.sub(r"[^A-Z]", "", cell_ref.upper())
            if not letters:
                return 0
            total = 0
            for ch in letters:
                total = total * 26 + (ord(ch) - ord("A") + 1)
            return total - 1

        def cell_value(cell) -> str:
            cell_type = cell.attrib.get("t")
            if cell_type == "s":
                v = first_child_named(cell, "v")
                if v is not None and v.text is not None:
                    try:
                        idx = int(v.text)
                        return shared_strings[idx] if idx < len(shared_strings) else ""
                    except ValueError:
                        return ""
                return ""

            if cell_type == "inlineStr":
                parts = []
                is_el = first_child_named(cell, "is")
                if is_el is not None:
                    for t in iter_named(is_el, "t"):
                        parts.append(t.text or "")
                return "".join(parts)

            v = first_child_named(cell, "v")
            if v is not None and v.text is not None:
                return v.text

            parts = []
            for t in iter_named(cell, "t"):
                parts.append(t.text or "")
            return "".join(parts)

        diagnostics = []
        best_df = None
        best_score = -1
        best_sheet = None

        for sheet_name in worksheet_names:
            try:
                root = ET.fromstring(z.read(sheet_name))
            except Exception as exc:
                diagnostics.append(f"{sheet_name}: XML parse failed: {exc}")
                continue

            rows_out = []
            row_count = 0
            cell_count = 0

            for row in iter_named(root, "row"):
                row_count += 1
                row_values = []
                cells = children_named(row, "c")
                cell_count += len(cells)

                for cell in cells:
                    ref = cell.attrib.get("r", "")
                    col_idx = cell_col_index(ref) if ref else len(row_values)

                    while len(row_values) <= col_idx:
                        row_values.append("")

                    row_values[col_idx] = cell_value(cell)

                if any(str(v).strip() for v in row_values):
                    rows_out.append(row_values)

            diagnostics.append(
                f"{sheet_name}: XML rows={row_count}, cells={cell_count}, nonblank_rows={len(rows_out)}"
            )

            if not rows_out:
                continue

            width = max(len(r) for r in rows_out)
            rows_out = [r + [""] * (width - len(r)) for r in rows_out]

            header_idx = 0
            header_score = -1
            for i, r in enumerate(rows_out[:100]):
                headers = [normalise_header(x) for x in r]
                score = len(set(required_columns).intersection(set(headers)))
                if score > header_score:
                    header_score = score
                    header_idx = i

            headers = [normalise_header(x) for x in rows_out[header_idx]]
            data_rows = rows_out[header_idx + 1:]

            last_nonblank = 0
            for i, h in enumerate(headers):
                if h:
                    last_nonblank = max(last_nonblank, i)
            for r in data_rows[:50]:
                for i, v in enumerate(r):
                    if str(v).strip():
                        last_nonblank = max(last_nonblank, i)

            headers = headers[:last_nonblank + 1]
            data_rows = [r[:last_nonblank + 1] for r in data_rows]

            fixed_headers = []
            seen = {}
            for i, h in enumerate(headers):
                h = h or f"BlankColumn{i+1}"
                if h in seen:
                    seen[h] += 1
                    h = f"{h}_{seen[h]}"
                else:
                    seen[h] = 1
                fixed_headers.append(h)

            candidate = pd.DataFrame(data_rows, columns=fixed_headers)

            nonblank_cols = [
                c for c in candidate.columns
                if str(c).strip() and not candidate[c].astype(str).str.strip().eq("").all()
            ]
            candidate = candidate[nonblank_cols]

            score = len(set(required_columns).intersection(set(candidate.columns)))
            diagnostics.append(
                f"{sheet_name}: header_idx={header_idx}, header_score={header_score}, "
                f"candidate_score={score}, columns={list(candidate.columns)[:30]}"
            )

            if score > best_score:
                best_score = score
                best_df = candidate
                best_sheet = sheet_name

        if best_df is None:
            raise ValueError(
                "Worksheet XML existed, but no rows could be extracted. Diagnostics:\n"
                + "\n".join("  - " + d for d in diagnostics)
            )

        if required_columns and best_score < min(3, len(required_columns)):
            raise ValueError(
                f"Direct XLSX XML read extracted rows but did not find expected headers. "
                f"Best sheet={best_sheet}, score={best_score}, columns={list(best_df.columns)[:50]}\n"
                "Diagnostics:\n"
                + "\n".join("  - " + d for d in diagnostics)
            )

        print(f"Read {path.name} by namespace-agnostic direct XLSX XML fallback from {best_sheet}")
        return best_df


def read_excel_guess(path: Path) -> pd.DataFrame:
    suffix = path.suffix.lower()

    if suffix in {".xlsx", ".xlsm"}:
        if not zipfile.is_zipfile(path):
            raise ValueError("File is not a real zipped Excel workbook")
        xls = pd.ExcelFile(path, engine="openpyxl")
        if not xls.sheet_names:
            raise ValueError("No worksheets found")
        return pd.read_excel(xls, sheet_name=xls.sheet_names[0], dtype=object)

    xls = pd.ExcelFile(path)
    if not xls.sheet_names:
        raise ValueError("No worksheets found")
    return pd.read_excel(xls, sheet_name=xls.sheet_names[0], dtype=object)


def read_input_file(path: Path, required_columns: list[str] | None = None) -> pd.DataFrame:
    suffix = path.suffix.lower()
    errors = []

    if suffix not in {".xlsx", ".xlsm", ".xls", ".csv", ".txt", ".html", ".htm"}:
        raise ValueError(f"Unsupported file type: {path.name}")

    attempts = [
        ("Delimited text first", lambda p: read_delimited_text_any_extension(p, required_columns)),
        ("Excel fallback", read_excel_guess),
        ("Direct XLSX XML fallback", lambda p: read_xlsx_xml_direct(p, required_columns)),
        ("HTML fallback", lambda p: read_html_guess(p, required_columns)),
    ]

    for label, reader in attempts:
        try:
            df = reader(path)
            df.columns = [normalise_header(c) for c in df.columns]
            print(f"Successfully read {path.name} using: {label}")
            return df
        except Exception as exc:
            errors.append(f"{label}: {exc}")

    raise ValueError(
        f"Could not read {path.name} in any supported format.\n"
        + "\n".join(f"  - {e}" for e in errors)
    )


def change_log_timestamp() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def make_change_id() -> str:
    return "CHG-" + datetime.now().strftime("%Y%m%d-%H%M%S") + "-" + uuid.uuid4().hex[:6].upper()


def append_change_log_entry(
    workbook_path: Path,
    *,
    script: str,
    action: str,
    sheet: str = "",
    rows_before: object = "",
    rows_after: object = "",
    rows_added: object = "",
    rows_updated: object = "",
    rows_removed: object = "",
    review_file: object = "",
    backup_file: object = "",
    status: str = "Completed",
    details: str = "",
) -> None:
    """
    Append a workbook mutation audit row to ChangeLog.

    This is intentionally separate from ImportLog:
    - ImportLog records parsed source files and import counts.
    - ChangeLog records scripts that changed ParsedTransactions.xlsx.

    The workbook must already exist. This function opens and saves the workbook.
    """
    workbook_path = Path(workbook_path)

    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found for ChangeLog append: {workbook_path}")

    from openpyxl import load_workbook

    wb = load_workbook(workbook_path)

    if CHANGE_LOG_SHEET not in wb.sheetnames:
        ws = wb.create_sheet(CHANGE_LOG_SHEET)
        ws.append(CHANGE_LOG_COLUMNS)
        ws.freeze_panes = "A2"
    else:
        ws = wb[CHANGE_LOG_SHEET]
        existing_headers = [cell.value for cell in ws[1]]

        # If the sheet exists but is empty or malformed, rewrite header row.
        if existing_headers != CHANGE_LOG_COLUMNS:
            if ws.max_row == 1 and all(value in (None, "") for value in existing_headers):
                ws.delete_rows(1)
                ws.append(CHANGE_LOG_COLUMNS)
            else:
                # Append missing columns without destroying existing entries.
                for col in CHANGE_LOG_COLUMNS:
                    if col not in existing_headers:
                        ws.cell(row=1, column=ws.max_column + 1).value = col
                existing_headers = [cell.value for cell in ws[1]]

    header_to_col = {cell.value: idx for idx, cell in enumerate(ws[1], start=1)}
    row_values = {
        "ChangeID": make_change_id(),
        "ChangedAt": change_log_timestamp(),
        "Script": script,
        "Action": action,
        "Workbook": workbook_path.name,
        "Sheet": sheet,
        "RowsBefore": rows_before,
        "RowsAfter": rows_after,
        "RowsAdded": rows_added,
        "RowsUpdated": rows_updated,
        "RowsRemoved": rows_removed,
        "ReviewFile": str(review_file) if review_file else "",
        "BackupFile": str(backup_file) if backup_file else "",
        "Status": status,
        "Details": details,
    }

    next_row = ws.max_row + 1
    for col_name in CHANGE_LOG_COLUMNS:
        col_idx = header_to_col.get(col_name)
        if col_idx is not None:
            ws.cell(row=next_row, column=col_idx).value = row_values.get(col_name, "")

    if ws.max_column > 0:
        end_col = ws.cell(row=1, column=ws.max_column).column_letter
        ws.auto_filter.ref = f"A1:{end_col}{ws.max_row}"

    wb.save(workbook_path)
    wb.close()

