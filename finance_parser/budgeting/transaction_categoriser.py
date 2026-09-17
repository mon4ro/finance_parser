from __future__ import annotations

import argparse
import re
import shutil
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from finance_parser.common import append_change_log_entry, clean_for_excel
from finance_parser.utilities.fresh_workbook_writer import (
    append_changelog_row,
    read_workbook_values_only,
    records_to_sheet_values,
    replace_with_fresh_workbook,
    sheet_values_to_records,
)


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_WORKBOOK = PROJECT_ROOT / "output" / "budgeting" / "ParsedTransactions.xlsx"
DEFAULT_RULES = PROJECT_ROOT / "rules" / "budgeting" / "TransactionRules.xlsx"
UNIFIED_SHEET = "UnifiedTransactions"
CATEGORY_RULES_SHEET = "CategoryRules"
OWNERSHIP_RULES_SHEET = "OwnershipRules"
SCRIPT_NAME = "transaction_categoriser.py"

CATEGORY_SET_FIELDS = {
    "SetSupercategory": "Supercategory",
    "SetCategory": "Category",
    "SetSubcategory": "Subcategory",
    "SetOwner": "Owner",
    "SetComments": "Review/Notes",
    "SetReviewNotes": "Review/Notes",
    "SetTag": "Tag",
}
OWNERSHIP_SET_FIELDS = {
    "SetOwner": "Owner",
}

RECATEGORISE_FIELDS = [
    "Supercategory",
    "Category",
    "Subcategory",
    "Owner",
    "Tag",
]
RECATEGORISE_SCOPES = {"all", "only_automatic"}
WRITE_MODES = {"openpyxl-mutating", "fresh-rebuild"}


@dataclass(frozen=True)
class Rule:
    rule_id: str
    enabled: bool
    priority: int
    source_sheet: str
    raw: dict[str, Any]


@dataclass
class RowChange:
    excel_row: int
    unified_id: str
    raw_receiver: str
    normalized_receiver: str
    rule_ids: list[str]
    changes: dict[str, tuple[Any, Any]]

    def describe(self) -> str:
        fields = ", ".join(
            f"{field}: {old!r} -> {new!r}" for field, (old, new) in self.changes.items()
        )
        return (
            f"row={self.excel_row} UnifiedID={self.unified_id!r} "
            f"NormalizedReceiver={self.normalized_receiver!r} Rules={'+'.join(self.rule_ids)} {fields}"
        )


@dataclass
class CategoriseStats:
    rows_scanned: int = 0
    rows_protected_by_comments: int = 0
    rows_eligible: int = 0
    rows_with_cleared_fields: int = 0
    fields_cleared: int = 0
    rows_changed_by_rules: int = 0
    fields_changed_by_rules: int = 0
    rows_changed_total: int = 0


def is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and pd.isna(value):
        return True
    return str(value).strip() == ""


def display_blank(value: object) -> str:
    """Return blank Excel values as empty string for readable change reports."""
    return "" if is_blank(value) else str(value)


def as_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and pd.isna(value):
        return ""
    return str(value)


def as_yes(value: Any, default: bool = False) -> bool:
    if is_blank(value):
        return default
    return str(value).strip().upper() in {"YES", "Y", "TRUE", "1"}


def normalise_header(value: Any) -> str:
    return as_text(value).strip()


def load_rules(rules_path: Path, sheet_name: str) -> list[Rule]:
    if not rules_path.exists():
        raise FileNotFoundError(f"Rules workbook not found: {rules_path}")

    try:
        df = pd.read_excel(rules_path, sheet_name=sheet_name, dtype=object, engine="openpyxl")
    except ValueError:
        # CategoryRules/OwnershipRules (the only two sheets this is called
        # for) are deliberately absent from TransactionRules.template.xlsx -
        # category taxonomy is a personal choice the template intentionally
        # doesn't prescribe, and OwnershipRules only gets built up as
        # accounts are configured. A brand-new user's first pipeline run
        # should just apply zero rules from a not-yet-created sheet, not
        # crash - this is not a misconfiguration to fail loudly on.
        return []

    df.columns = [normalise_header(c) for c in df.columns]
    rules: list[Rule] = []

    for _, row in df.iterrows():
        raw = {col: row.get(col) for col in df.columns}
        rule_id = as_text(raw.get("RuleID")).strip()
        if not rule_id:
            continue

        enabled = as_yes(raw.get("Enabled"), default=True)
        try:
            priority = int(float(raw.get("Priority"))) if not is_blank(raw.get("Priority")) else 1000
        except Exception:
            priority = 1000

        rules.append(Rule(rule_id=rule_id, enabled=enabled, priority=priority, source_sheet=sheet_name, raw=raw))

    return sorted([r for r in rules if r.enabled], key=lambda r: (r.priority, r.rule_id))


def value_matches(value: Any, match_type: Any, pattern: Any, case_sensitive: Any = None) -> bool:
    match_type_text = as_text(match_type).strip().upper()
    pattern_text = as_text(pattern)
    value_text = as_text(value)

    if not pattern_text and match_type_text not in {"BLANK", "ISBLANK", "NOT_BLANK", "NOTBLANK"}:
        return False

    if not as_yes(case_sensitive, default=False):
        value_cmp = value_text.upper()
        pattern_cmp = pattern_text.upper()
    else:
        value_cmp = value_text
        pattern_cmp = pattern_text

    if match_type_text in {"", "CONTAINS"}:
        return pattern_cmp in value_cmp
    if match_type_text in {"EXACT", "EQUALS", "="}:
        return value_cmp == pattern_cmp
    if match_type_text == "STARTS_WITH":
        return value_cmp.startswith(pattern_cmp)
    if match_type_text == "ENDS_WITH":
        return value_cmp.endswith(pattern_cmp)
    if match_type_text == "REGEX":
        flags = 0 if as_yes(case_sensitive, default=False) else re.IGNORECASE
        return re.search(pattern_text, value_text, flags=flags) is not None
    if match_type_text == "ALL_CONTAINS":
        parts = [part.strip() for part in re.split(r"[,;]", pattern_cmp) if part.strip()]
        return bool(parts) and all(part in value_cmp for part in parts)
    if match_type_text in {"BLANK", "ISBLANK"}:
        return is_blank(value)
    if match_type_text in {"NOT_BLANK", "NOTBLANK"}:
        return not is_blank(value)

    raise ValueError(f"Unsupported match type: {match_type}")


def row_value(row: dict[str, Any], field_name: Any) -> Any:
    field = as_text(field_name).strip()
    if not field:
        return ""
    return row.get(field, "")


def condition_matches(rule: Rule, row: dict[str, Any], index: int) -> bool:
    field = rule.raw.get(f"ConditionField{index}")
    match_type = rule.raw.get(f"ConditionMatchType{index}")
    pattern = rule.raw.get(f"ConditionPattern{index}")

    if is_blank(field) and is_blank(match_type) and is_blank(pattern):
        return True

    return value_matches(row_value(row, field), match_type, pattern, rule.raw.get("CaseSensitive"))


def rule_matches(rule: Rule, row: dict[str, Any], *, category_rule: bool) -> bool:
    field = rule.raw.get("MatchField")
    match_type = rule.raw.get("MatchType")
    pattern = rule.raw.get("Pattern")

    if category_rule:
        if not value_matches(row_value(row, field), match_type, pattern, rule.raw.get("CaseSensitive")):
            return False
        return condition_matches(rule, row, 1) and condition_matches(rule, row, 2)

    return value_matches(row_value(row, field), match_type, pattern, rule.raw.get("CaseSensitive"))


def should_set_value(current_value: Any, new_value: Any, overwrite_mode: Any) -> bool:
    if is_blank(new_value):
        return False

    mode = as_text(overwrite_mode).strip().upper() or "BLANK_ONLY"
    if mode == "FORCE":
        return True
    if mode == "RULE_ONLY":
        return True

    return is_blank(current_value)


def apply_rule_to_row(
    rule: Rule,
    row: dict[str, Any],
    set_fields: dict[str, str],
) -> dict[str, tuple[Any, Any]]:
    changes: dict[str, tuple[Any, Any]] = {}
    overwrite_mode = rule.raw.get("OverwriteMode")

    for rule_field, target_field in set_fields.items():
        if target_field not in row:
            continue
        new_value = rule.raw.get(rule_field)
        old_value = row.get(target_field)

        if should_set_value(old_value, new_value, overwrite_mode):
            row[target_field] = new_value
            changes[target_field] = (display_blank(old_value), display_blank(new_value))

    return changes


def find_comment_column(headers: list[str]) -> str:
    # Added "Comments" back into the list for backwards compatibility
    for candidate in ["Review/Notes", "Comments", "Comment"]: 
        if candidate in headers:
            return candidate
    raise ValueError("UnifiedTransactions must contain a Comments, Comment, or Review/Notes column. This blank column is the categorisation trigger.")


def worksheet_headers(ws) -> dict[str, int]:
    headers: dict[str, int] = {}
    for col_idx, cell in enumerate(ws[1], start=1):
        name = normalise_header(cell.value)
        if name:
            headers[name] = col_idx
    return headers


def read_row(ws, row_idx: int, headers: dict[str, int]) -> dict[str, Any]:
    return {header: ws.cell(row=row_idx, column=col_idx).value for header, col_idx in headers.items()}


def write_row_changes(ws, row_idx: int, headers: dict[str, int], changes: dict[str, tuple[Any, Any]]) -> None:
    for field, (_, new_value) in changes.items():
        if field in headers:
            ws.cell(row=row_idx, column=headers[field]).value = clean_for_excel(new_value)


def merge_changes(
    base_changes: dict[str, tuple[Any, Any]],
    new_changes: dict[str, tuple[Any, Any]],
) -> None:
    for field, (old_value, new_value) in new_changes.items():
        if field in base_changes:
            base_changes[field] = (base_changes[field][0], new_value)
        else:
            base_changes[field] = (old_value, new_value)


def clear_recategorise_fields(
    row: dict[str, Any],
    headers: dict[str, int],
) -> dict[str, tuple[Any, Any]]:
    changes: dict[str, tuple[Any, Any]] = {}

    for field in RECATEGORISE_FIELDS:
        if field not in headers or field not in row:
            continue

        old_value = row.get(field)
        if is_blank(old_value):
            continue

        row[field] = ""
        changes[field] = (display_blank(old_value), "")

    return changes


def make_backup(workbook_path: Path, *, recategorise_scope: str | None) -> Path:
    suffix = "_backup_before_recategorise" if recategorise_scope else "_backup_before_categorise"
    backup_path = workbook_path.with_name(workbook_path.stem + suffix + workbook_path.suffix)
    shutil.copy2(workbook_path, backup_path)
    return backup_path


def save_workbook_safely(wb, workbook_path: Path, *, recategorise_scope: str | None) -> Path:
    backup_path = make_backup(workbook_path, recategorise_scope=recategorise_scope)
    tmp_suffix = "_tmp_recategorise" if recategorise_scope else "_tmp_categorise"
    tmp_path = workbook_path.with_name(workbook_path.stem + tmp_suffix + workbook_path.suffix)

    wb.save(tmp_path)
    wb.close()

    test_wb = load_workbook(tmp_path, read_only=True)
    test_wb.close()

    tmp_path.replace(workbook_path)
    return backup_path


def supports_ansi_colour() -> bool:
    return sys.stdout.isatty()


def colour_text(value: str, code: str) -> str:
    if not supports_ansi_colour():
        return value
    return f"\033[{code}m{value}\033[0m"


def print_recategorise_warning(scope: str) -> None:
    if scope == "all":
        scope_text = (
            "--recategorise --all clears existing Supercategory, Category, "
            "Subcategory, Owner and Tag values for every row, including rows "
            "with non-blank Review/Notes."
        )
    else:
        scope_text = (
            "--recategorise --only-automatic clears existing Supercategory, "
            "Category, Subcategory, Owner and Tag values only where "
            "Review/Notes/Comment is blank. Rows with non-blank Review/Notes are protected."
        )

    lines = [
        "⚠  DESTRUCTIVE RECATEGORISATION",
        "",
        scope_text,
        "Manual categorisation values can be overwritten.",
        "A backup will be created before a wet run writes changes.",
    ]

    width = max(len(line) for line in lines)
    border = "!" + "=" * (width + 2) + "!"

    print("")
    print(colour_text(border, "1;31"))
    for line in lines:
        padded = f"! {line.ljust(width)} !"
        print(colour_text(padded, "1;31"))
    print(colour_text(border, "1;31"))
    print("")


def confirm_recategorise(scope: str) -> bool:
    first = input('Type "yes" to continue: ').strip()
    if first != "yes":
        print("Recategorise cancelled.")
        return False

    second = input('Are you really sure? Type "yes" again: ').strip()
    if second != "yes":
        print("Recategorise cancelled.")
        return False

    return True


def update_row_with_rules(
    original_row: dict[str, Any],
    *,
    excel_row: int,
    headers: dict[str, int],
    comment_column: str,
    category_rules: list[Rule],
    ownership_rules: list[Rule],
    recategorise_scope: str | None,
    stats: CategoriseStats,
) -> tuple[dict[str, Any] | None, RowChange | None]:
    stats.rows_scanned += 1
    comment_is_blank = is_blank(original_row.get(comment_column))

    if recategorise_scope is None and not comment_is_blank:
        stats.rows_protected_by_comments += 1
        return None, None
    if recategorise_scope == "only_automatic" and not comment_is_blank:
        stats.rows_protected_by_comments += 1
        return None, None

    stats.rows_eligible += 1

    working_row = dict(original_row)
    row_changes: dict[str, tuple[Any, Any]] = {}
    matched_rule_ids: list[str] = []

    if recategorise_scope is not None:
        clear_changes = clear_recategorise_fields(working_row, headers)
        if clear_changes:
            stats.rows_with_cleared_fields += 1
            stats.fields_cleared += len(clear_changes)
            row_changes.update(clear_changes)

    rule_changed_this_row = False

    for rule in category_rules:
        if not rule_matches(rule, working_row, category_rule=True):
            continue
        rule_changes = apply_rule_to_row(rule, working_row, CATEGORY_SET_FIELDS)
        if rule_changes:
            merge_changes(row_changes, rule_changes)
            matched_rule_ids.append(rule.rule_id)
            rule_changed_this_row = True
            stats.fields_changed_by_rules += len(rule_changes)
        if as_yes(rule.raw.get("StopIfMatched"), default=True):
            break

    for rule in ownership_rules:
        if not rule_matches(rule, working_row, category_rule=False):
            continue
        rule_changes = apply_rule_to_row(rule, working_row, OWNERSHIP_SET_FIELDS)
        clear_owner = as_yes(rule.raw.get("ClearOwner"), default=False)
        if clear_owner and "Owner" in working_row and (
            is_blank(rule.raw.get("SetOwner"))
            or should_set_value(working_row.get("Owner"), "__CLEAR__", rule.raw.get("OverwriteMode"))
        ):
            old_owner = working_row.get("Owner")
            working_row["Owner"] = ""
            rule_changes["Owner"] = (display_blank(old_owner), "")

        if rule_changes:
            merge_changes(row_changes, rule_changes)
            matched_rule_ids.append(rule.rule_id)
            rule_changed_this_row = True
            stats.fields_changed_by_rules += len(rule_changes)
        if as_yes(rule.raw.get("StopIfMatched"), default=True):
            break

    if rule_changed_this_row:
        stats.rows_changed_by_rules += 1

    if not row_changes:
        return None, None

    stats.rows_changed_total += 1
    return working_row, RowChange(
        excel_row=excel_row,
        unified_id=as_text(original_row.get("UnifiedID")),
        raw_receiver=as_text(original_row.get("RawReceiver")),
        normalized_receiver=as_text(original_row.get("NormalizedReceiver")),
        rule_ids=matched_rule_ids,
        changes=row_changes,
    )


def categorise_workbook_fresh_rebuild(
    workbook_path: Path,
    rules_path: Path,
    *,
    dry_run: bool,
    recategorise_scope: str | None = None,
    stats: CategoriseStats | None = None,
) -> list[RowChange]:
    if not workbook_path.exists():
        raise FileNotFoundError(f"Parsed transactions workbook not found: {workbook_path}")

    if recategorise_scope is not None and recategorise_scope not in RECATEGORISE_SCOPES:
        raise ValueError(f"Unsupported recategorise scope: {recategorise_scope}")

    category_rules = load_rules(rules_path, CATEGORY_RULES_SHEET)
    ownership_rules = load_rules(rules_path, OWNERSHIP_RULES_SHEET)

    print(f"Read previous workbook values only: {workbook_path}")
    sheets = read_workbook_values_only(workbook_path)
    print(f"Loaded sheets: {len(sheets)}")
    print("No loaded workbook object will be saved.")

    if UNIFIED_SHEET not in sheets:
        raise ValueError(f"Workbook does not contain required sheet: {UNIFIED_SHEET}")

    headers_list, records = sheet_values_to_records(sheets[UNIFIED_SHEET])
    headers = {header: idx + 1 for idx, header in enumerate(headers_list) if header}
    comment_column = find_comment_column(list(headers.keys()))

    changes: list[RowChange] = []
    if stats is None:
        stats = CategoriseStats()

    for record_idx, original_row in enumerate(records, start=2):
        updated_row, row_change = update_row_with_rules(
            original_row,
            excel_row=record_idx,
            headers=headers,
            comment_column=comment_column,
            category_rules=category_rules,
            ownership_rules=ownership_rules,
            recategorise_scope=recategorise_scope,
            stats=stats,
        )
        if row_change:
            changes.append(row_change)
            if not dry_run and updated_row is not None:
                records[record_idx - 2] = updated_row

    if not dry_run and changes:
        if recategorise_scope:
            action = f"Recategorise UnifiedTransactions ({recategorise_scope})"
            details = (
                "Fresh-rebuild write mode. Cleared existing "
                "Supercategory/Category/Subcategory/Owner/Tag values "
                f"with scope={recategorise_scope}, then applied CategoryRules and OwnershipRules."
            )
        else:
            action = "Categorise UnifiedTransactions"
            details = "Fresh-rebuild write mode. Applied CategoryRules and OwnershipRules to rows with blank Review/Notes."

        sheets[UNIFIED_SHEET] = records_to_sheet_values(headers_list, records)
        append_changelog_row(
            sheets,
            script=SCRIPT_NAME,
            action=action,
            sheet=UNIFIED_SHEET,
            rows_updated=len(changes),
            status="Completed",
            details=details,
            backup_file=None,
        )

        print("Created fresh workbook from scratch.")
        backup_path, _, writer_stats = replace_with_fresh_workbook(
            workbook_path,
            sheets,
            backup_label="before_categorise_fresh_rebuild",
            basic_formatting=True,
            # Real workbook test showed Excel repairs generated table parts
            # (Removed Feature: AutoFilter/Table from xl/tables/table*.xml).
            # Keep basic formatting and sheet-level autofilters, but do not
            # create Excel Table objects in the production fresh-rebuild path.
            excel_tables=False,
        )
        print("Validated fresh workbook: OK")
        print(f"Backed up old workbook: {backup_path}")
        print(f"Sheets written: {writer_stats['sheets_written']}")
        print(f"Excel tables added: {writer_stats['tables_added']}")
        print("No loaded workbook object was saved.")

    return changes


def categorise_workbook(
    workbook_path: Path,
    rules_path: Path,
    *,
    dry_run: bool,
    recategorise_scope: str | None = None,
    stats: CategoriseStats | None = None,
    write_mode: str = "fresh-rebuild",
) -> list[RowChange]:
    if write_mode not in WRITE_MODES:
        raise ValueError(f"Unsupported write mode: {write_mode}")

    if write_mode == "openpyxl-mutating":
        print(
            "WARNING: --write-mode openpyxl-mutating is a legacy mode. "
            "It may trigger Excel repair/corruption warnings on existing workbooks. "
            "Use the default fresh-rebuild mode for normal use."
        )

    if write_mode == "fresh-rebuild":
        return categorise_workbook_fresh_rebuild(
            workbook_path,
            rules_path,
            dry_run=dry_run,
            recategorise_scope=recategorise_scope,
            stats=stats,
        )

    if not workbook_path.exists():
        raise FileNotFoundError(f"Parsed transactions workbook not found: {workbook_path}")

    if recategorise_scope is not None and recategorise_scope not in RECATEGORISE_SCOPES:
        raise ValueError(f"Unsupported recategorise scope: {recategorise_scope}")

    category_rules = load_rules(rules_path, CATEGORY_RULES_SHEET)
    ownership_rules = load_rules(rules_path, OWNERSHIP_RULES_SHEET)

    wb = load_workbook(workbook_path)
    if UNIFIED_SHEET not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Workbook does not contain required sheet: {UNIFIED_SHEET}")

    ws = wb[UNIFIED_SHEET]
    headers = worksheet_headers(ws)
    comment_column = find_comment_column(list(headers.keys()))

    changes: list[RowChange] = []
    if stats is None:
        stats = CategoriseStats()

    for row_idx in range(2, ws.max_row + 1):
        original_row = read_row(ws, row_idx, headers)
        updated_row, row_change = update_row_with_rules(
            original_row,
            excel_row=row_idx,
            headers=headers,
            comment_column=comment_column,
            category_rules=category_rules,
            ownership_rules=ownership_rules,
            recategorise_scope=recategorise_scope,
            stats=stats,
        )
        if row_change:
            changes.append(row_change)
            if not dry_run:
                write_row_changes(ws, row_idx, headers, row_change.changes)

    if not dry_run and changes:
        backup_path = save_workbook_safely(wb, workbook_path, recategorise_scope=recategorise_scope)

        if recategorise_scope:
            action = f"Recategorise UnifiedTransactions ({recategorise_scope})"
            details = (
                "Cleared existing Supercategory/Category/Subcategory/Owner/Tag values "
                f"with scope={recategorise_scope}, then applied CategoryRules and OwnershipRules."
            )
        else:
            action = "Categorise UnifiedTransactions"
            details = "Applied CategoryRules and OwnershipRules to rows with blank Review/Notes."

        append_change_log_entry(
            workbook_path,
            script=SCRIPT_NAME,
            action=action,
            sheet=UNIFIED_SHEET,
            rows_updated=len(changes),
            backup_file=backup_path,
            status="Completed",
            details=details,
        )
    else:
        wb.close()

    return changes


def print_changes(changes: list[RowChange], max_rows: int) -> None:
    if not changes:
        print("No UnifiedTransactions rows changed.")
        return

    print(f"Changed rows: {len(changes)}")
    shown = changes if max_rows <= 0 else changes[:max_rows]
    for change in shown:
        print("  " + change.describe())

    remaining = len(changes) - len(shown)
    if remaining > 0:
        print(f"  ... {remaining} more changed row(s) not printed. Use --max-print 0 to print all.")


def print_stats(stats: CategoriseStats, *, recategorise_scope: str | None, dry_run: bool = False) -> None:
    print("Categoriser complete." if not dry_run else "Categoriser dry run complete.")
    print("Categoriser statistics:")
    print(f"Rows scanned:                       {stats.rows_scanned}")
    print(f"Rows protected by Review/Notes:     {stats.rows_protected_by_comments}")
    print(f"Rows eligible:                      {stats.rows_eligible}")

    if recategorise_scope is not None:
        print(f"Rows with fields cleared:     {stats.rows_with_cleared_fields}")
        print(f"Fields cleared:               {stats.fields_cleared}")

    print(f"Rows changed by rules:        {stats.rows_changed_by_rules}")
    print(f"Fields changed by rules:      {stats.fields_changed_by_rules}")
    print(f"Rows changed total:           {stats.rows_changed_total}")


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Apply budgeting category and ownership rules to UnifiedTransactions. "
            "Only rows with blank Review/Notes/Comment are eligible unless --recategorise --all is used."
        )
    )
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK, help=f"Parsed transactions workbook. Default: {DEFAULT_WORKBOOK}")
    parser.add_argument("--rules", type=Path, default=DEFAULT_RULES, help=f"Transaction rules workbook. Default: {DEFAULT_RULES}")
    parser.add_argument("--dry-run", action="store_true", help="Show changes without modifying the workbook or ChangeLog.")
    parser.add_argument("--recategorise", action="store_true", help="Clear existing categorisation fields before applying rules. Requires --all or --only-automatic.")
    recategorise_scope = parser.add_mutually_exclusive_group()
    recategorise_scope.add_argument("--all", action="store_true", help="With --recategorise: wipe and recategorise all rows, including rows with non-blank Review/Notes.")
    recategorise_scope.add_argument("--only-automatic", action="store_true", help="With --recategorise: wipe and recategorise only rows where Review/Notes/Comment is blank.")
    parser.add_argument(
        "--write-mode",
        choices=sorted(WRITE_MODES),
        default="fresh-rebuild",
        help=(
            "Workbook write backend. fresh-rebuild is the default and writes a brand-new "
            "workbook package from values only. openpyxl-mutating is a legacy/debug mode "
            "and may trigger Excel repair warnings. Default: fresh-rebuild."
        ),
    )
    parser.add_argument("--max-print", type=int, default=50, help="Maximum changed rows to print. Use 0 to print all. Default: 50.")
    return parser


def main() -> None:
    parser = build_arg_parser()
    args = parser.parse_args()

    if (args.all or args.only_automatic) and not args.recategorise:
        parser.error("--all and --only-automatic can only be used together with --recategorise")

    if args.recategorise and not (args.all or args.only_automatic):
        parser.error("--recategorise requires either --all or --only-automatic")

    recategorise_scope = None
    if args.recategorise:
        recategorise_scope = "all" if args.all else "only_automatic"
        print_recategorise_warning(recategorise_scope)

        if not args.dry_run and not confirm_recategorise(recategorise_scope):
            return

    stats = CategoriseStats()
    changes = categorise_workbook(
        args.workbook,
        args.rules,
        dry_run=args.dry_run,
        recategorise_scope=recategorise_scope,
        stats=stats,
        write_mode=args.write_mode,
    )
    print_changes(changes, args.max_print)
    print_stats(stats, recategorise_scope=recategorise_scope, dry_run=args.dry_run)

    if args.dry_run:
        print()
        print("Dry run only: workbook was not modified.")
    elif changes:
        print()
        print(f"Output workbook: {args.workbook}")
        print("ChangeLog appended.")


if __name__ == "__main__":
    main()
