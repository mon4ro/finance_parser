"""
Sync CategoryTaxonomy from the real Master Budget workbook.

Master Budget (external to this repo, hand-built) is treated as the
canonical source of truth for the category taxonomy: its own budget summary
sheet lays out Supercategory/Category/Subcategory as real row headers. This
script reads that sheet read-only and regenerates a CategoryTaxonomy sheet
inside the rules workbook (TransactionRules.xlsx) - it never writes to
Master Budget itself.

CategoryTaxonomy is then used two ways, both deliberately non-destructive to
already-set data (see transaction_categoriser.py and fresh_workbook_writer.py):
- A CategoryRules row that sets SetSubcategory alone can have its parent
  Supercategory/Category auto-derived at rule-match time, for the ~95% of
  subcategories with exactly one real parent.
- A currently-blank Category/Supercategory cell in Excel gets a live XLOOKUP
  formula suggestion keyed on whatever the user types into Subcategory.
Neither path ever touches a cell that already has a value - once anything is
set (by a rule, by the formula, or typed by hand as a deliberate exception),
it's permanent. This script itself only ever touches the CategoryTaxonomy
reference sheet, never CategoryRules/OwnershipRules/UnifiedTransactions.
"""

from __future__ import annotations

import argparse
import re
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

from finance_parser.settings import get_settings
from finance_parser.utilities.fresh_workbook_writer import timestamped_backup_path


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_RULES_WORKBOOK = PROJECT_ROOT / "rules" / "budgeting" / "TransactionRules.xlsx"

MASTER_BUDGET_SHEET = "Master Budget"
CATEGORY_TAXONOMY_SHEET = "CategoryTaxonomy"
VALID_SUPERCATEGORIES = {"INCOME", "EXPENSES", "SAVINGS"}

_VERSION_RE = re.compile(r"v(\d+)\.(\d+)\.(\d+)", re.IGNORECASE)


def find_latest_master_budget_file(folder: Path) -> Path:
    """
    Auto-detect the current Master Budget workbook: highest "vMAJOR.MINOR.PATCH"
    in the filename among real "Master Budget *.xlsm" files, excluding any
    filename containing "test" (a scratch/draft copy, never authoritative).

    Raises FileNotFoundError with a clear message if none are found, rather
    than silently falling back to something wrong.
    """
    candidates = []
    for path in folder.glob("Master Budget*.xlsm"):
        if "test" in path.name.lower():
            continue
        match = _VERSION_RE.search(path.name)
        if not match:
            continue
        version = tuple(int(part) for part in match.groups())
        candidates.append((version, path))

    if not candidates:
        raise FileNotFoundError(
            f"No 'Master Budget*.xlsm' file with a vMAJOR.MINOR.PATCH version "
            f"found in {folder} (excluding filenames containing 'test')."
        )

    candidates.sort(key=lambda pair: pair[0])
    return candidates[-1][1]


def extract_taxonomy(path: Path, *, sheet_name: str = MASTER_BUDGET_SHEET) -> list[tuple[str, str, str]]:
    """
    Read-only extraction of (Supercategory, Category, Subcategory) leaves
    from Master Budget's own budget summary sheet, where they're laid out as
    real row headers in columns A/B/C.

    Walks top to bottom, forward-filling the current Supercategory (column A)
    and Category (column B) so a bare Subcategory row (column C) can be
    attributed to both. Two structural quirks handled deliberately:
    - The sheet has a one-line KPI overview block before the real detailed
      section (bare Supercategory names with no child rows directly under
      them) - a lookahead check (is the very next row a child row, with no
      column A of its own?) tells a real section start apart from one of
      those overview lines, so they're skipped rather than mistaken for the
      start of the taxonomy.
    - The detailed section ends where a column-A header appears that isn't a
      real Supercategory (e.g. "TOTAL ASSETS", the start of the unrelated
      balance-sheet section further down the same sheet) - extraction stops
      there rather than misattributing unrelated rows to the last-seen
      Supercategory.
    """
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        ws = wb[sheet_name]

        def cell(row: int, col: int) -> object:
            return ws.cell(row=row, column=col).value

        triples: list[tuple[str, str, str]] = []
        current_super: str | None = None
        current_cat: str | None = None
        started = False

        row = 1
        max_row = ws.max_row
        while row <= max_row:
            a, b, c = cell(row, 1), cell(row, 2), cell(row, 3)
            if a is not None:
                a = str(a).strip()
                if a in VALID_SUPERCATEGORIES:
                    next_a, next_b, next_c = cell(row + 1, 1), cell(row + 1, 2), cell(row + 1, 3)
                    is_real_section = next_a is None and (next_b is not None or next_c is not None)
                    if is_real_section:
                        current_super = a
                        current_cat = None
                        started = True
                    elif started:
                        break
                elif started:
                    break
            elif b is not None:
                current_cat = str(b).strip()
            elif c is not None and current_super is not None:
                triples.append((current_super, current_cat or "", str(c).strip()))
            row += 1

        return triples
    finally:
        wb.close()


def build_taxonomy_rows(triples: list[tuple[str, str, str]]) -> list[dict[str, object]]:
    """
    Group extracted triples by Subcategory to flag ambiguous ones (a
    Subcategory name with more than one real (Supercategory, Category)
    parent) - those need a rule/user to also supply Category explicitly,
    since the name alone doesn't determine the parent.
    """
    parents_by_sub: dict[str, set[tuple[str, str]]] = {}
    for supercategory, category, subcategory in triples:
        parents_by_sub.setdefault(subcategory, set()).add((supercategory, category))

    rows: list[dict[str, object]] = []
    seen: set[tuple[str, str, str]] = set()
    for supercategory, category, subcategory in triples:
        key = (supercategory, category, subcategory)
        if key in seen:
            continue
        seen.add(key)
        ambiguous = len(parents_by_sub[subcategory]) > 1
        rows.append(
            {
                "Supercategory": supercategory,
                "Category": category,
                "Subcategory": subcategory,
                "Ambiguous": "YES" if ambiguous else "NO",
            }
        )
    return rows


def write_category_taxonomy_sheet(rules_workbook: Path, rows: list[dict[str, object]]) -> Path:
    """
    In-place edit of the rules workbook: replace (or add) only the
    CategoryTaxonomy sheet, leaving every other sheet (TransactionRules,
    CategoryRules, OwnershipRules, Lookups, README - all hand-maintained by
    the user directly in Excel) completely untouched. No Excel Table objects
    exist in this workbook (checked), so in-place openpyxl mutation is safe
    here - unlike ParsedTransactions.xlsx, this is never fresh-rebuilt.

    Returns the backup path (a timestamped copy made before the edit).
    """
    backup_path = timestamped_backup_path(rules_workbook, label="before_taxonomy_sync")
    shutil.copy2(rules_workbook, backup_path)

    wb = load_workbook(rules_workbook)
    if CATEGORY_TAXONOMY_SHEET in wb.sheetnames:
        del wb[CATEGORY_TAXONOMY_SHEET]
    ws = wb.create_sheet(CATEGORY_TAXONOMY_SHEET)

    headers = ["Supercategory", "Category", "Subcategory", "Ambiguous"]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])

    tmp_path = rules_workbook.with_name(rules_workbook.stem + "_tmp_taxonomy_sync" + rules_workbook.suffix)
    wb.save(tmp_path)
    wb.close()

    verify_wb = load_workbook(tmp_path, read_only=True)
    verify_wb.close()

    tmp_path.replace(rules_workbook)
    return backup_path


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--master-budget", type=Path, help="Explicit Master Budget file path (overrides auto-detection).")
    parser.add_argument("--rules-workbook", type=Path, default=DEFAULT_RULES_WORKBOOK)
    parser.add_argument("--dry-run", action="store_true", help="Report what would change without writing anything.")
    return parser


def main() -> None:
    args = build_arg_parser().parse_args()

    if args.master_budget:
        master_budget_path = args.master_budget
    else:
        settings = get_settings()
        folder = settings.master_budget_folder()
        if folder is None:
            print("No paths.master_budget_folder configured in settings.yaml - nothing to sync.")
            sys.exit(1)
        master_budget_path = find_latest_master_budget_file(folder)

    print(f"Reading taxonomy from: {master_budget_path.name}")
    triples = extract_taxonomy(master_budget_path)
    rows = build_taxonomy_rows(triples)
    ambiguous_rows = [r for r in rows if r["Ambiguous"] == "YES"]

    print(f"Extracted {len(triples)} taxonomy leaves ({len(rows)} distinct rows).")
    if ambiguous_rows:
        print(f"Ambiguous subcategories ({len({r['Subcategory'] for r in ambiguous_rows})} distinct names):")
        for row in ambiguous_rows:
            print(f"  {row['Supercategory']} / {row['Category'] or '(none)'} / {row['Subcategory']}")

    if args.dry_run:
        print(f"\nDry-run: would write {len(rows)} rows to '{CATEGORY_TAXONOMY_SHEET}' in {args.rules_workbook}")
        print("No files were modified.")
        return

    backup_path = write_category_taxonomy_sheet(args.rules_workbook, rows)
    print(f"\nBackup written: {backup_path}")
    print(f"'{CATEGORY_TAXONOMY_SHEET}' sheet updated in {args.rules_workbook}")


if __name__ == "__main__":
    main()
