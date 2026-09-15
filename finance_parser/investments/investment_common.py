from __future__ import annotations

import os
import shutil
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

from finance_parser.common import (
    clean_for_excel,
    clean_dataframe_for_excel,
    format_date,
    normalise_amount,
    normalise_header,
    normalise_text,
)
from finance_parser.settings import get_settings


INVESTMENT_RAW_SHEET = "RawInvestmentTransactions"
INVESTMENT_TRANSACTIONS_SHEET = "InvestmentTransactions"
INVESTMENT_IMPORT_LOG_SHEET = "InvestmentImportLog"
INSTRUMENT_MASTER_SHEET = "InstrumentMaster"

INVESTMENT_RAW_COLUMNS = [
    "InvestmentRawID",
    "Broker",
    "Portfolio",
    "PortfolioOwner",
    "PortfolioType",
    "BookingDate",
    "TradeDate",
    "SettlementDate",
    "TransactionTypeRaw",
    "InstrumentName",
    "ISIN",
    "Quantity",
    "UnitPrice",
    "Interest",
    "TotalFees",
    "TotalFeesCurrency",
    "CashAmount",
    "CashCurrency",
    "AcquisitionValue",
    "AcquisitionCurrency",
    "Result",
    "ResultCurrency",
    "TotalQuantity",
    "CashBalance",
    "ExchangeRate",
    "Description",
    "CancellationDate",
    "CalculationID",
    "ConfirmationNumber",
    "BrokerageFee",
    "BrokerageFeeCurrency",
    "ReferenceExchangeRate",
    "OriginalInterest",
    "ExportDate",
    "ImportedAt",
    "SourceFile",
]

INVESTMENT_TRANSACTIONS_COLUMNS = [
    "InvestmentTransactionID",
    "InvestmentRawID",
    "Broker",
    "Portfolio",
    "PortfolioOwner",
    "PortfolioType",
    "TransactionType",
    "TradeDate",
    "SettlementDate",
    "InstrumentName",
    "NormalizedInstrument",
    "ISIN",
    "InstrumentType",
    "AssetClass",
    "InstrumentCurrency",
    "Quantity",
    "UnitPrice",
    "CashAmount",
    "CashCurrency",
    "AcquisitionValue",
    "AcquisitionCurrency",
    "TotalQuantity",
    "ExchangeRate",
    "Description",
    "BrokerageFee",
    "BrokerageFeeCurrency",
    "ExportDate",
    "ImportedAt",
    "SourceFile",
    "Comments",
]

INVESTMENT_IMPORT_LOG_COLUMNS = [
    "ImportRunID",
    "ImportedAt",
    "Broker",
    "SourceFile",
    "RowsRead",
    "RowsNew",
    "RowsDuplicate",
    "Status",
    "Error",
]

INSTRUMENT_MASTER_COLUMNS = [
    "Enabled",
    "RuleID",
    "Broker",
    "ISIN",
    "RawInstrumentName",
    "NormalizedInstrument",
    "InstrumentType",
    "AssetClass",
    "Currency",
    "Ticker",
    "Exchange",
    "PriceSource",
    "PriceSymbol",
    "Notes",
]


def normalise_key(value: object) -> str:
    return normalise_text(value).upper()


def normalise_investment_transaction_type(value: object) -> str:
    text = normalise_key(value)

    mapping = {
        "OSTO": "BUY",
        "MERKINTÄ": "BUY",
        "MERKINTA": "BUY",
        "SHARE PURCHASE": "BUY",
        "MYYNTI": "SELL",
        "LUNASTUS": "SELL",
        "SELL": "SELL",
        "SELL OF PURCHASED SHARE": "SELL",
        "OSINKO": "DIVIDEND",
        "DIVIDEND": "DIVIDEND",
        "OSINGONMAKSU": "DIVIDEND",
        "TALLETUS": "DEPOSIT",
        "SAVINGS": "DEPOSIT",
        "NOSTO": "WITHDRAWAL",
        "CASH TRANSFERRED": "WITHDRAWAL",
        "KORKO": "INTEREST",
        "KULU": "FEE",
        "PALKKIO": "FEE",
        "VERO": "TAX",
        "VERONPIDÄTYS": "TAX",
        "ENNAKONPIDÄTYS": "TAX",
        "DELIVERY": "DELIVERY",
        "MATCHING": "MATCHING",
        "REDEMPTION": "REDEMPTION",
    }

    return mapping.get(text, text)


def investment_transaction_id(raw_id: object) -> str:
    return "I-" + normalise_text(raw_id)


def read_sheet(path: Path, sheet_name: str, columns: list[str]) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame(columns=columns)

    try:
        df = pd.read_excel(path, sheet_name=sheet_name, dtype=object, engine="openpyxl")
    except ValueError:
        return pd.DataFrame(columns=columns)

    df.columns = [normalise_header(c) for c in df.columns]

    for col in columns:
        if col not in df.columns:
            df[col] = ""

    return df[columns]


def read_dynamic_sheet(path: Path, sheet_name: str) -> pd.DataFrame:
    """
    Read an optional broker-specific raw-export sheet with its existing columns.

    Broker raw sheets intentionally retain a closer-to-export shape, so they do
    not use a fixed global column list like InvestmentTransactions.
    """
    if not path.exists():
        return pd.DataFrame()

    try:
        df = pd.read_excel(path, sheet_name=sheet_name, dtype=object, engine="openpyxl")
    except ValueError:
        return pd.DataFrame()

    df.columns = [normalise_header(c) for c in df.columns]
    return df


def load_instrument_master(path: Path | None) -> pd.DataFrame:
    """
    Load rules/investments/InstrumentMaster.xlsx.

    Missing file is not an error: the parser still works, but
    NormalizedInstrument falls back to InstrumentName and InstrumentType remains
    blank unless already present.
    """
    if path is None or not path.exists():
        return pd.DataFrame(columns=INSTRUMENT_MASTER_COLUMNS)

    try:
        df = pd.read_excel(path, sheet_name=INSTRUMENT_MASTER_SHEET, dtype=object, engine="openpyxl")
    except ValueError:
        # Fallback for a workbook where the first sheet is the master table.
        df = pd.read_excel(path, sheet_name=0, dtype=object, engine="openpyxl")

    df.columns = [normalise_header(c) for c in df.columns]

    for col in INSTRUMENT_MASTER_COLUMNS:
        if col not in df.columns:
            df[col] = ""

    df = df[INSTRUMENT_MASTER_COLUMNS].copy()

    # Ignore disabled rows. Blank Enabled is treated as YES to make manual editing easier.
    enabled = df["Enabled"].map(lambda v: normalise_key(v) not in {"NO", "N", "FALSE", "0", "DISABLED"})
    df = df.loc[enabled].copy()

    # Drop rows that cannot match anything.
    has_key = (
        df["ISIN"].map(normalise_text).ne("")
        | df["RawInstrumentName"].map(normalise_text).ne("")
        | df["NormalizedInstrument"].map(normalise_text).ne("")
    )
    df = df.loc[has_key].copy()

    # Make matching deterministic.
    df["_ISIN_KEY"] = df["ISIN"].map(normalise_key)
    df["_RAW_NAME_KEY"] = df["RawInstrumentName"].map(normalise_key)
    df["_NORMALIZED_KEY"] = df["NormalizedInstrument"].map(normalise_key)
    df["_BROKER_KEY"] = df["Broker"].map(normalise_key)

    return df


def _master_value(rule: pd.Series, column: str) -> str:
    return normalise_text(rule.get(column, ""))


def _find_instrument_rule(row: pd.Series, master: pd.DataFrame) -> tuple[pd.Series | None, str]:
    """
    Matching priority:

    1. Broker + ISIN
    2. ISIN
    3. Broker + RawInstrumentName
    4. RawInstrumentName
    5. Broker + NormalizedInstrument
    6. NormalizedInstrument

    RawInstrumentName matching also lets special names such as LEHTOU0120 map to
    the same normalized instrument as the ordinary stock.

    A broker-pinned master row (Broker not blank) is preferred when one exists,
    but is not a hard filter: a row pinned to one broker (e.g. NOKIA entered
    from a Nordnet export) must still match the same real instrument held at a
    different broker (e.g. EVLI). Real data confirmed this is safe - a scan of
    the whole InstrumentMaster sheet found zero cases of the same
    RawInstrumentName legitimately referring to two different instruments
    across brokers (share classes like Kesko A/B already have distinct raw
    names, so they're never at risk of this fallback conflating them).
    """
    if master.empty:
        return None, ""

    broker_key = normalise_key(row.get("Broker", ""))
    isin_key = normalise_key(row.get("ISIN", ""))
    raw_name_key = normalise_key(row.get("InstrumentName", ""))
    normalized_key = normalise_key(row.get("NormalizedInstrument", ""))

    def _best_match(mask, own_broker_label: str, any_broker_label: str):
        matches = master.loc[mask].copy()
        if len(matches) == 0:
            return None, ""
        matches["_BROKER_SPECIFIC"] = matches["_BROKER_KEY"].map(lambda v: 1 if v == broker_key else 0)
        matches = matches.sort_values("_BROKER_SPECIFIC", ascending=False)
        best = matches.iloc[0]
        label = own_broker_label if best["_BROKER_KEY"] == broker_key else any_broker_label
        return best, label

    if isin_key:
        rule, label = _best_match(master["_ISIN_KEY"] == isin_key, "BROKER_ISIN", "ISIN")
        if rule is not None:
            return rule, label

    if raw_name_key:
        rule, label = _best_match(master["_RAW_NAME_KEY"] == raw_name_key, "BROKER_RAW_NAME", "RAW_NAME")
        if rule is not None:
            return rule, label

    if normalized_key:
        rule, label = _best_match(master["_NORMALIZED_KEY"] == normalized_key, "BROKER_NORMALIZED", "NORMALIZED")
        if rule is not None:
            return rule, label

    return None, ""


def _is_blank(value: object) -> bool:
    return normalise_text(value) == ""


def _set_if_blank_or_default(df: pd.DataFrame, idx, column: str, value: object, default_value: object = "") -> None:
    """
    Safe update helper.

    This is deliberately conservative: it fills blanks and also replaces the
    default NormalizedInstrument=InstrumentName fallback, but avoids overwriting
    a user's clearly manual value.
    """
    new_value = normalise_text(value)
    if not new_value:
        return

    current = normalise_text(df.at[idx, column])
    default_text = normalise_text(default_value)

    if current == "" or (default_text and current == default_text):
        df.at[idx, column] = new_value


def apply_portfolio_ownership(transactions: pd.DataFrame) -> pd.DataFrame:
    """
    Backfill PortfolioOwner/PortfolioType from settings.yaml across the whole
    transaction history, not just newly-parsed rows.

    raw_to_investment_transactions() already applies this default the moment
    a row is first created, but a settings.yaml mapping added or changed
    later (e.g. after importing more history) should retroactively fill
    existing rows too - same reasoning as the ISIN backfill in
    apply_instrument_master(), and why apply_instrument_master() itself is
    re-run against the full combined transaction set on every pipeline run,
    not just against new rows.
    """
    out = transactions.copy()
    settings = get_settings()

    for idx, row in out.iterrows():
        broker = row.get("Broker", "")
        portfolio = row.get("Portfolio", "")

        if _is_blank(row.get("PortfolioOwner", "")):
            owner = settings.portfolio_owner(broker, portfolio)
            if owner:
                out.at[idx, "PortfolioOwner"] = owner

        if _is_blank(row.get("PortfolioType", "")):
            portfolio_type = settings.portfolio_type(broker, portfolio)
            if portfolio_type:
                out.at[idx, "PortfolioType"] = portfolio_type

    return out


def apply_instrument_master(transactions: pd.DataFrame, master: pd.DataFrame) -> pd.DataFrame:
    out = transactions.copy()

    for col in INVESTMENT_TRANSACTIONS_COLUMNS:
        if col not in out.columns:
            out[col] = ""

    if len(out) == 0:
        return out[INVESTMENT_TRANSACTIONS_COLUMNS]

    # Basic safe fallback: keep raw instrument visible even without a master match.
    for idx, row in out.iterrows():
        if _is_blank(row.get("NormalizedInstrument", "")):
            out.at[idx, "NormalizedInstrument"] = normalise_text(row.get("InstrumentName", ""))

    if master.empty:
        return out[INVESTMENT_TRANSACTIONS_COLUMNS]

    for idx, row in out.iterrows():
        rule, matched_by = _find_instrument_rule(row, master)
        if rule is None:
            continue

        raw_name = normalise_text(row.get("InstrumentName", ""))

        _set_if_blank_or_default(
            out,
            idx,
            "NormalizedInstrument",
            _master_value(rule, "NormalizedInstrument"),
            default_value=raw_name,
        )
        _set_if_blank_or_default(out, idx, "InstrumentType", _master_value(rule, "InstrumentType"))
        _set_if_blank_or_default(out, idx, "AssetClass", _master_value(rule, "AssetClass"))
        _set_if_blank_or_default(out, idx, "InstrumentCurrency", _master_value(rule, "Currency"))
        # ISIN was only ever populated from whatever the raw broker export
        # happened to include natively (raw_to_investment_transactions()) -
        # never backfilled from the master sheet, even when it has one. Real
        # data check: several real instruments (NOKIA, TELIA, NESTE...) have
        # an ISIN in InstrumentMaster that never made it into transaction rows.
        _set_if_blank_or_default(out, idx, "ISIN", _master_value(rule, "ISIN"))


    return out[INVESTMENT_TRANSACTIONS_COLUMNS]


def raw_to_investment_transactions(raw_new: pd.DataFrame) -> pd.DataFrame:
    rows = []
    settings = get_settings()

    for _, row in raw_new.iterrows():
        instrument_name = normalise_text(row.get("InstrumentName", ""))
        broker = row.get("Broker", "")
        portfolio = row.get("Portfolio", "")

        # PortfolioOwner/PortfolioType are account-level facts (which person,
        # which account wrapper), not content-based rules, so - same reasoning
        # as Owner on the budgeting side (see CLAUDE.md) - they're defaulted
        # from settings.yaml here, not hardcoded in parser code or looked up
        # from InstrumentMaster (which is content-matching, like CategoryRules).
        portfolio_owner = normalise_text(row.get("PortfolioOwner", "")) or settings.portfolio_owner(broker, portfolio)
        portfolio_type = normalise_text(row.get("PortfolioType", "")) or settings.portfolio_type(broker, portfolio)

        rows.append({
            "InvestmentTransactionID": investment_transaction_id(row.get("InvestmentRawID", "")),
            "InvestmentRawID": row.get("InvestmentRawID", ""),
            "Broker": broker,
            "Portfolio": portfolio,
            "PortfolioOwner": portfolio_owner,
            "PortfolioType": portfolio_type,
            "TransactionType": normalise_investment_transaction_type(row.get("TransactionTypeRaw", "")),
            "TradeDate": format_date(row.get("TradeDate", "")),
            "SettlementDate": format_date(row.get("SettlementDate", "")),
            "InstrumentName": instrument_name,
            "NormalizedInstrument": instrument_name,
            "ISIN": normalise_text(row.get("ISIN", "")),
            "InstrumentType": "",
            "AssetClass": "",
            "InstrumentCurrency": "",
            "Quantity": normalise_amount(row.get("Quantity", 0)),
            "UnitPrice": normalise_amount(row.get("UnitPrice", 0)),
            "CashAmount": normalise_amount(row.get("CashAmount", 0)),
            "CashCurrency": normalise_text(row.get("CashCurrency", "")) or "EUR",
            "AcquisitionValue": normalise_amount(row.get("AcquisitionValue", 0)),
            "AcquisitionCurrency": normalise_text(row.get("AcquisitionCurrency", "")) or normalise_text(row.get("CashCurrency", "")) or "EUR",
            "TotalQuantity": normalise_amount(row.get("TotalQuantity", 0)),
            "ExchangeRate": row.get("ExchangeRate", ""),
            "Description": normalise_text(row.get("Description", "")),
            "BrokerageFee": normalise_amount(row.get("BrokerageFee", 0)),
            "BrokerageFeeCurrency": normalise_text(row.get("BrokerageFeeCurrency", "")) or normalise_text(row.get("CashCurrency", "")) or "EUR",
            "ExportDate": row.get("ExportDate", ""),
            "ImportedAt": row.get("ImportedAt", ""),
            "SourceFile": row.get("SourceFile", ""),
            "Comments": "",
        })

    return pd.DataFrame(rows, columns=INVESTMENT_TRANSACTIONS_COLUMNS)


def append_dataframe_to_worksheet(wb: Workbook, sheet_name: str, df: pd.DataFrame) -> None:
    df = clean_dataframe_for_excel(df)

    if sheet_name in wb.sheetnames:
        del wb[sheet_name]

    ws = wb.create_sheet(sheet_name)
    ws.append([clean_for_excel(c) for c in df.columns])

    for row in df.itertuples(index=False, name=None):
        ws.append([clean_for_excel(v) for v in row])

    ws.freeze_panes = "A2"

    if len(df) > 0 and len(df.columns) > 0:
        end_col = ws.cell(row=1, column=len(df.columns)).column_letter
        end_row = len(df) + 1
        ws.auto_filter.ref = f"A1:{end_col}{end_row}"


def write_investment_output_workbook(
    workbook_path: Path,
    raw_df: pd.DataFrame,
    transactions_df: pd.DataFrame,
    import_log_df: pd.DataFrame,
    extra_sheets: dict[str, pd.DataFrame] | None = None,
) -> None:
    workbook_path.parent.mkdir(parents=True, exist_ok=True)

    tmp_path = workbook_path.with_name(workbook_path.stem + "_tmp_save" + workbook_path.suffix)
    backup_path = workbook_path.with_name(workbook_path.stem + "_backup_before_last_run" + workbook_path.suffix)

    if workbook_path.exists():
        shutil.copy2(workbook_path, backup_path)

    wb = Workbook()
    default = wb.active
    wb.remove(default)

    append_dataframe_to_worksheet(wb, INVESTMENT_RAW_SHEET, raw_df)
    append_dataframe_to_worksheet(wb, INVESTMENT_TRANSACTIONS_SHEET, transactions_df)
    append_dataframe_to_worksheet(wb, INVESTMENT_IMPORT_LOG_SHEET, import_log_df)

    for sheet_name, df in (extra_sheets or {}).items():
        append_dataframe_to_worksheet(wb, sheet_name, df)

    wb.save(tmp_path)

    test_wb = load_workbook(tmp_path, read_only=True)
    test_wb.close()

    os.replace(tmp_path, workbook_path)
