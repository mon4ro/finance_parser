from __future__ import annotations

import argparse
from pathlib import Path

import openpyxl
import pandas as pd

from finance_parser.common import normalise_text
from finance_parser.investments.fetch_instrument_prices import (
    DEFAULT_INSTRUMENT_MASTER,
    DEFAULT_PRICES_WORKBOOK,
    PRICES_COLUMNS,
    imported_at_now,
    load_existing_prices,
    load_instrument_master_price_config,
    write_prices_workbook,
)

# One-time import of the user's own manually-tracked monthly fund prices
# (input/investments/investment_price_backfill.xlsm, gitignored - real
# personal data, never committed). Hardcoded to this specific file's layout
# (Investments sheet, PRICES > FUNDS block, rows 12-25, month columns from
# row 2) rather than a generalised importer - deliberately not meant to be
# rerun against a different file shape.
#
# Rows in the source file NOT included here:
#   - Seligson: Global Top 25 Brands, Spiltan Aktiefond Investmentbolag:
#     already fully covered by real sources (seligson.fi, Avanza) - importing
#     would only ever be a no-op under the gap-fill-only rule.
#   - Nordea savings 25: belongs to another family member, not yet imported
#     into this pipeline - nothing to attach a price history to yet.
#
# Verified before building: compared each instrument's manual value nearest
# the real-data transition date against the real first value. Ratios were
# all in the 0.84-1.16 range (normal week-to-week/multi-year market drift),
# not the order-of-magnitude mismatch a wrong currency would produce (as
# caught earlier with Avanza's SEK-converted Franklin Technology values).
MANUAL_PRICE_ROWS: dict[str, str] = {
    "OP-Eurooppa Pienyhtiöt A": "OP-EUROOPPA PIENYHTIÖT A",
    "OP-Suomi pienyhtiöt A": "OP-SUOMI PIENYHTIÖT A",
    "OP-Aasia Indeksi A": "OP-AASIA INDEKSI A",
    "Franklin Technology Fund A (Acc) ": "FRANKLIN TECHNOLOGY A ACC",
    "BlackRock Global Funds - World Technology Fund A2": "BGF WORLD TECHNOLOGY",
    "Nordnet Suomi-superrahasto": "NORDNET SUOMI INDEKSI",
    "Handelsbank USA indeksi": "HANDELSBANKEN USA INDEX",
    "OP-Maltillinen A": "OP-MALTILLINEN A",
    "OP-Amerikka Arvoyhtiöt A": "OP-AMERIKKA A",
    "OP-Ilmasto A": "OP-ILMASTO A",
    "OP-Suomi A": "OP-SUOMI A",
}

SOURCE_SHEET = "Investments"
SOURCE_ROW_RANGE = range(12, 26)  # PRICES > FUNDS block
HEADER_ROW = 2
NAME_COLUMN = 3
PRICE_SOURCE = "MANUAL"


def _month_end(ts: pd.Timestamp) -> pd.Timestamp:
    """
    The source file labels each monthly column inconsistently (e.g. a
    reference column dated 2017-12-31, but the regular monthly series dated
    2018-01-01, 2018-02-01, ...) - confirmed these are genuinely distinct
    months' values, not duplicates, by comparing adjacent columns' values.
    Normalise every column to that month's real last day, matching this
    project's month-end-anchored data model.
    """
    return pd.Timestamp(ts.year, ts.month, 1) + pd.offsets.MonthEnd(0)


def read_manual_prices(source_path: Path) -> pd.DataFrame:
    wb = openpyxl.load_workbook(source_path, data_only=True)
    ws = wb[SOURCE_SHEET]

    col_dates: dict[int, pd.Timestamp] = {}
    for col in range(1, ws.max_column + 1):
        value = ws.cell(row=HEADER_ROW, column=col).value
        if hasattr(value, "year"):
            col_dates[col] = _month_end(pd.Timestamp(value))

    rows: list[dict[str, object]] = []
    for row_idx in SOURCE_ROW_RANGE:
        raw_name = normalise_text(ws.cell(row=row_idx, column=NAME_COLUMN).value)
        instrument = MANUAL_PRICE_ROWS.get(raw_name)
        if instrument is None:
            continue

        for col, month_end in col_dates.items():
            value = ws.cell(row=row_idx, column=col).value
            # Real bug caught on the first real run: some cells (mostly
            # future/not-yet-reached months) evaluate to a literal 0 rather
            # than being blank - a fund price of exactly EUR0.00 is never a
            # real value, so treat <= 0 the same as blank/missing.
            if not isinstance(value, (int, float)) or value <= 0:
                continue
            rows.append({
                "NormalizedInstrument": instrument,
                "Date": month_end,
                "Close": float(value),
            })

    return pd.DataFrame(rows, columns=["NormalizedInstrument", "Date", "Close"])


def _instrument_metadata(instrument_master_path: Path) -> dict[str, dict[str, str]]:
    config = load_instrument_master_price_config(instrument_master_path)
    metadata = {}
    for _, row in config.iterrows():
        metadata[row["NormalizedInstrument"]] = {
            "ISIN": normalise_text(row.get("ISIN", "")),
            "Currency": normalise_text(row.get("Currency", "")),
        }
    return metadata


def backfill_gaps(
    source_path: Path,
    prices_workbook: Path,
    instrument_master_path: Path,
    *,
    verbose: bool = True,
) -> tuple[pd.DataFrame, dict[str, object]]:
    """Fill gaps only - an existing row for a (instrument, date) already on
    file is never touched or replaced. Mirrors fetch_seligson_history.py /
    fetch_avanza_history.py."""
    manual = read_manual_prices(source_path)
    existing = load_existing_prices(prices_workbook)
    metadata = _instrument_metadata(instrument_master_path)

    stats: dict[str, object] = {"instruments_backfilled": 0, "rows_added": 0, "rows_skipped_already_covered": 0}
    new_rows: list[dict[str, object]] = []
    fetched_at = imported_at_now()

    for instrument, group in manual.groupby("NormalizedInstrument"):
        existing_dates = set(
            pd.to_datetime(existing[existing["NormalizedInstrument"] == instrument]["Date"]).dt.date
        )
        meta = metadata.get(instrument, {})

        added = 0
        for _, row in group.iterrows():
            day = row["Date"].date()
            if day in existing_dates:
                stats["rows_skipped_already_covered"] += 1
                continue
            new_rows.append({
                "NormalizedInstrument": instrument,
                "ISIN": meta.get("ISIN", ""),
                "PriceSymbol": "",
                "Date": day.isoformat(),
                "Close": row["Close"],
                "Currency": meta.get("Currency", ""),
                "PriceSource": PRICE_SOURCE,
                "FetchedAt": fetched_at,
            })
            added += 1

        if added:
            stats["instruments_backfilled"] += 1
        stats["rows_added"] += added
        if verbose:
            print(f"  {instrument}: {added} new row(s) added, {len(group) - added} already covered")

    new_prices = pd.DataFrame(new_rows, columns=PRICES_COLUMNS)
    if len(existing) == 0:
        combined = new_prices
    elif len(new_prices) == 0:
        combined = existing
    else:
        combined = pd.concat([existing, new_prices], ignore_index=True)
    combined = combined.drop_duplicates(subset=["NormalizedInstrument", "Date"], keep="last")
    combined = combined.sort_values(["NormalizedInstrument", "Date"]).reset_index(drop=True)

    return combined[PRICES_COLUMNS], stats


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="One-time import of investment_price_backfill.xlsm's manually-tracked monthly fund prices. "
                    "Fills InstrumentPrices.xlsx gaps only - never touches an existing (instrument, date) row."
    )
    parser.add_argument("--source", default=None, help="Path to the manual backfill .xlsm. Defaults to input/investments/investment_price_backfill.xlsm.")
    parser.add_argument("--prices-workbook", default=str(DEFAULT_PRICES_WORKBOOK))
    parser.add_argument("--instrument-master", default=str(DEFAULT_INSTRUMENT_MASTER))
    parser.add_argument("--dry-run", action="store_true")
    return parser


def main() -> None:
    args = build_arg_parser().parse_args()

    project_root = Path(__file__).resolve().parents[2]
    source_path = Path(args.source).expanduser().resolve() if args.source else project_root / "input" / "investments" / "investment_price_backfill.xlsm"
    prices_workbook = Path(args.prices_workbook).expanduser().resolve()
    instrument_master_path = Path(args.instrument_master).expanduser().resolve()

    print(f"Reading manual prices from: {source_path}")
    combined, stats = backfill_gaps(source_path, prices_workbook, instrument_master_path)

    print()
    print("Manual price backfill complete." if not args.dry_run else "Manual price backfill dry run complete.")
    print(f"Instruments backfilled: {stats['instruments_backfilled']}")
    print(f"Rows added:             {stats['rows_added']}")
    print(f"Rows already covered:   {stats['rows_skipped_already_covered']}")

    if args.dry_run:
        print()
        print("Dry run only: workbook was not modified.")
        return

    write_prices_workbook(prices_workbook, combined)
    print()
    print(f"Output workbook: {prices_workbook}")


if __name__ == "__main__":
    main()
